import io
import struct
import zipfile
from pathlib import Path

import pymupdf
from fontTools.ttLib import TTFont
from lxml import etree
from pptx import Presentation

from vivepdf.ops.convert_slides import PptxParams, to_pptx
from vivepdf.rpc.progress import silent_progress

SIDECAR = Path(__file__).resolve().parent.parent
FONTS = SIDECAR / "vivepdf" / "assets" / "fonts"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
EOT_MAGIC = 0x504C
PREVIEW_AND_PRINT = 0x0004


def _convert(document: pymupdf.Document, tmp_path: Path, mode: str = "editable") -> Path:
    source = tmp_path / "kaynak.pdf"
    document.save(source)
    document.close()
    result = to_pptx(
        PptxParams(path=str(source), output=str(tmp_path / "sunum.pptx"), mode=mode),
        silent_progress(),
    )
    return Path(result.output)


def _font_parts(output: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(output) as archive:
        return {
            name: archive.read(name) for name in archive.namelist() if name.startswith("ppt/fonts/")
        }


def _program(data: bytes) -> TTFont:
    size, font_size = struct.unpack_from("<2L", data)
    assert size == len(data)
    assert struct.unpack_from("<H", data, 34)[0] == EOT_MAGIC
    return TTFont(io.BytesIO(data[-font_size:]))


def test_an_uninstalled_font_travels_inside_the_presentation(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_text(
        (40, 60), "Gömülü yazı", fontsize=14, fontname="dv", fontfile=str(FONTS / "DejaVuSans.ttf")
    )
    page.insert_text(
        (40, 100),
        "Kalın satır",
        fontsize=14,
        fontname="dvb",
        fontfile=str(FONTS / "DejaVuSans-Bold.ttf"),
    )

    output = _convert(document, tmp_path)

    with zipfile.ZipFile(output) as archive:
        listed = etree.fromstring(archive.read("ppt/presentation.xml"))
    (entry,) = listed.iter(f"{{{P}}}embeddedFont")
    assert entry.find(f"{{{P}}}font").get("typeface") == "DejaVu Sans"
    assert entry.find(f"{{{P}}}regular") is not None
    assert entry.find(f"{{{P}}}bold") is not None
    programs = [_program(data) for data in _font_parts(output).values()]
    assert {font["name"].getBestFamilyName() for font in programs} == {"DejaVu Sans"}
    assert all(ord("ö") in font.getBestCmap() for font in programs)
    assert Presentation(str(output)).slides[0].shapes


def test_a_font_that_forbids_editing_is_not_embedded(tmp_path: Path) -> None:
    font = TTFont(FONTS / "DejaVuSans.ttf")
    font["OS/2"].fsType = PREVIEW_AND_PRINT
    buffer = io.BytesIO()
    font.save(buffer)
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_font(fontname="kisitli", fontbuffer=buffer.getvalue())
    page.insert_text((40, 60), "Kısıtlı", fontsize=14, fontname="kisitli")

    output = _convert(document, tmp_path)

    texts = [
        shape.text_frame.text
        for shape in Presentation(str(output)).slides[0].shapes
        if shape.has_text_frame
    ]
    assert texts == ["Kısıtlı"]
    assert _font_parts(output) == {}


def test_standard_fonts_and_picture_slides_embed_nothing(tmp_path: Path) -> None:
    for mode in ("editable", "image"):
        document = pymupdf.open()
        page = document.new_page(width=400, height=300)
        page.insert_text((40, 60), "Standart", fontsize=14, fontname="helv")
        page.insert_text(
            (40, 100), "Gömülü", fontsize=14, fontname="dv", fontfile=str(FONTS / "DejaVuSans.ttf")
        )
        folder = tmp_path / mode
        folder.mkdir()

        output = _convert(document, folder, mode)

        expected = 1 if mode == "editable" else 0
        assert len(_font_parts(output)) == expected
