from pathlib import Path

import pymupdf
import pytest
from pptx import Presentation

from vivepdf.ops.convert_slides import PptxParams, to_pptx
from vivepdf.rpc.progress import silent_progress

SIDECAR = Path(__file__).resolve().parent.parent
FONT = str(SIDECAR / "vivepdf" / "assets" / "fonts" / "DejaVuSans.ttf")
EMU_PER_POINT = 12700
TEXT = "Eğik yazı satırı burada"


def _slide_box(tmp_path: Path, angle: float, rotation: int = 0):
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    pivot = pymupdf.Point(60, 200)
    page.insert_text(
        pivot,
        TEXT,
        fontsize=18,
        fontname="dejavu",
        fontfile=FONT,
        morph=(pivot, pymupdf.Matrix(-angle)),
    )
    page.set_rotation(rotation)
    page = document.reload_page(page)
    (line,) = [line for block in page.get_text("rawdict")["blocks"] for line in block["lines"]]
    quad = pymupdf.recover_line_quad(line) * page.rotation_matrix
    source = tmp_path / "kaynak.pdf"
    document.save(source)
    document.close()
    result = to_pptx(
        PptxParams(path=str(source), output=str(tmp_path / "sunum.pptx")), silent_progress()
    )
    presentation = Presentation(result.output)
    width = 300 if rotation % 180 else 400
    ratio = presentation.slide_width / (width * EMU_PER_POINT)
    (shape,) = [shape for shape in presentation.slides[0].shapes if shape.has_text_frame]
    return shape, ratio, quad


def _points(value: int, ratio: float) -> float:
    return value / EMU_PER_POINT / ratio


@pytest.mark.parametrize("angle", [10, 30])
def test_a_slanted_line_keeps_its_size_and_place(tmp_path: Path, angle: float) -> None:
    shape, ratio, quad = _slide_box(tmp_path, angle)

    center = (quad.ul + quad.lr) / 2
    assert abs(_points(shape.width, ratio) - abs(quad.ur - quad.ul)) < 0.5
    assert abs(_points(shape.height, ratio) - abs(quad.ll - quad.ul)) < 0.5
    assert abs(_points(shape.left + shape.width // 2, ratio) - center.x) < 0.5
    assert abs(_points(shape.top + shape.height // 2, ratio) - center.y) < 0.5
    assert shape.text_frame.text == TEXT


def test_a_slanted_line_on_a_turned_page_keeps_its_size(tmp_path: Path) -> None:
    shape, ratio, quad = _slide_box(tmp_path, 10, rotation=90)

    center = (quad.ul + quad.lr) / 2
    assert abs(_points(shape.width, ratio) - abs(quad.ur - quad.ul)) < 0.5
    assert abs(_points(shape.height, ratio) - abs(quad.ll - quad.ul)) < 0.5
    assert abs(_points(shape.left + shape.width // 2, ratio) - center.x) < 0.5
    assert abs(_points(shape.top + shape.height // 2, ratio) - center.y) < 0.5


def test_a_level_line_keeps_its_bounding_box(tmp_path: Path) -> None:
    shape, ratio, quad = _slide_box(tmp_path, 0)

    assert shape.rotation == 0
    assert abs(_points(shape.width, ratio) - quad.rect.width) < 0.5
    assert abs(_points(shape.height, ratio) - quad.rect.height) < 1
