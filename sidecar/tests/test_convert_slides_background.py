from pathlib import Path

import pymupdf
from pptx import Presentation

from vivepdf.ops.convert_slides import PptxParams, to_pptx
from vivepdf.rpc.progress import silent_progress

SIDECAR = Path(__file__).resolve().parent.parent
FONT = str(SIDECAR / "vivepdf" / "assets" / "fonts" / "DejaVuSans.ttf")
EMU_PER_POINT = 12700
SHADING = (
    "<</ShadingType 2/ColorSpace/DeviceRGB/Coords[40 0 300 0]"
    "/Function<</FunctionType 2/Domain[0 1]/C0[1 0 0]/C1[0 0 1]/N 1>>/Extend[true true]>>"
)


def _presentation(document: pymupdf.Document, tmp_path: Path):
    source = tmp_path / "kaynak.pdf"
    document.save(source)
    document.close()
    result = to_pptx(
        PptxParams(path=str(source), output=str(tmp_path / "sunum.pptx")), silent_progress()
    )
    return Presentation(result.output)


def _background(slide) -> pymupdf.Pixmap:
    (picture,) = [shape for shape in slide.shapes if shape.shape_type == 13]
    return pymupdf.Pixmap(picture.image.blob)


def _pixel(background: pymupdf.Pixmap, point: tuple[float, float], width: float):
    scale = background.width / width
    return background.pixel(round(point[0] * scale), round(point[1] * scale))


def _gradient_text_page() -> pymupdf.Document:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_text((40, 100), "Gradyan", fontsize=60, fontname="helv")
    xref = page.get_contents()[0]
    stream = document.xref_stream(xref).decode()
    stream = stream.replace("[<4772616479616e>]TJ", "7 Tr [<4772616479616e>]TJ", 1)
    document.update_stream(xref, stream.replace("ET\nQ", "ET\n/Sh0 sh\nQ", 1).encode())
    shading = document.get_new_xref()
    document.update_object(shading, SHADING)
    kind, value = document.xref_get_key(page.xref, "Resources")
    entry = f"<</Sh0 {shading} 0 R>>"
    if kind == "xref":
        document.xref_set_key(int(value.split()[0]), "Shading", entry)
    else:
        document.xref_set_key(page.xref, "Resources/Shading", entry)
    return document


def test_gradient_filled_lettering_stays_in_the_slide_background(tmp_path: Path) -> None:
    document = _gradient_text_page()
    painted = [
        point
        for point in ((48, 80), (52, 70), (46, 85), (50, 90))
        if document[0].get_pixmap(clip=pymupdf.Rect(point, point[0] + 1, point[1] + 1)).pixel(0, 0)
        != (255, 255, 255)
    ]

    background = _background(_presentation(document, tmp_path).slides[0])

    assert painted
    assert all(_pixel(background, point, 400) != (255, 255, 255) for point in painted)


def test_invisible_text_does_not_widen_the_text_box(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    style = {"fontsize": 14, "fontname": "dejavu", "fontfile": FONT}
    page.insert_text((40, 80), "1", render_mode=3, **style)
    width = pymupdf.Font(fontfile=FONT).text_length("1", fontsize=14)
    page.insert_text((40 + width, 80), "Görünen satır", **style)

    presentation = _presentation(document, tmp_path)

    ratio = presentation.slide_width / (400 * EMU_PER_POINT)
    (box,) = [shape for shape in presentation.slides[0].shapes if shape.has_text_frame]
    assert box.text_frame.text == "Görünen satır"
    assert abs(box.left / EMU_PER_POINT / ratio - (40 + width)) < 1


def test_text_placed_on_the_slide_leaves_the_background(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_text((40, 100), "Yazı", fontsize=60, fontname="dejavu", fontfile=FONT)
    inked = [
        (round(x), round(y))
        for x, y in ((48, 80), (55, 70), (60, 90), (75, 85), (90, 80))
        if page.get_pixmap(clip=pymupdf.Rect(x, y, x + 1, y + 1)).pixel(0, 0) != (255, 255, 255)
    ]

    background = _background(_presentation(document, tmp_path).slides[0])

    assert inked
    assert all(_pixel(background, point, 400) == (255, 255, 255) for point in inked)
