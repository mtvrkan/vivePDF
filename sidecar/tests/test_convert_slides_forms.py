from pathlib import Path

import pymupdf
from pptx import Presentation

from vivepdf.ops._hidden_text import widget_origins
from vivepdf.ops.convert_slides import PptxParams, to_pptx
from vivepdf.rpc.progress import silent_progress

SIDECAR = Path(__file__).resolve().parent.parent
FONT = str(SIDECAR / "vivepdf" / "assets" / "fonts" / "DejaVuSans.ttf")
FIELD = pymupdf.Rect(50, 100, 300, 130)
BOX = pymupdf.Rect(50, 160, 66, 176)


def _form(value: str = "Nesibe Şafak", checked: bool = True) -> pymupdf.Document:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_text((50, 90), "Ad soyad", fontsize=12, fontname="dejavu", fontfile=FONT)
    field = pymupdf.Widget()
    field.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
    field.field_name = "fullName"
    field.rect = FIELD
    field.field_value = value
    field.text_fontsize = 14
    page.add_widget(field)
    box = pymupdf.Widget()
    box.field_type = pymupdf.PDF_WIDGET_TYPE_CHECKBOX
    box.field_name = "news"
    box.rect = BOX
    box.field_value = checked
    page.add_widget(box)
    return document


def _slide(document: pymupdf.Document, tmp_path: Path):
    source = tmp_path / "kaynak.pdf"
    document.save(source)
    document.close()
    result = to_pptx(
        PptxParams(path=str(source), output=str(tmp_path / "sunum.pptx")), silent_progress()
    )
    return Presentation(result.output).slides[0]


def _texts(slide) -> list[str]:
    return [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]


def _inked(slide, area: pymupdf.Rect) -> int:
    (picture,) = [shape for shape in slide.shapes if shape.shape_type == 13]
    background = pymupdf.Pixmap(picture.image.blob)
    scale = background.width / 400
    return sum(
        1
        for y in range(round(area.y0 * scale), round(area.y1 * scale))
        for x in range(round(area.x0 * scale), round(area.x1 * scale))
        if sum(background.pixel(x, y)) < 300
    )


def test_a_filled_field_shows_its_value_once(tmp_path: Path) -> None:
    slide = _slide(_form(), tmp_path)

    assert _texts(slide) == ["Ad soyad"]
    assert _inked(slide, FIELD) > 0


def test_a_ticked_box_is_not_turned_into_a_letter(tmp_path: Path) -> None:
    slide = _slide(_form(value=""), tmp_path)

    assert _texts(slide) == ["Ad soyad"]
    assert _inked(slide, BOX) > 0


def test_a_page_without_fields_has_no_field_text() -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_text((50, 90), "Ad soyad", fontsize=12)

    assert widget_origins(page) == set()


def test_printed_text_under_an_empty_field_stays_editable(tmp_path: Path) -> None:
    document = _form(value="", checked=False)
    document[0].insert_text((60, 120), "Örnek değer", fontsize=12, fontname="dejavu", fontfile=FONT)

    slide = _slide(document, tmp_path)

    assert sorted(_texts(slide)) == ["Ad soyad", "Örnek değer"]
