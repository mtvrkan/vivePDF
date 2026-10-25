import csv
import datetime as dt
from pathlib import Path

import numpy as np
import pymupdf
from openpyxl import load_workbook

from vivepdf.ops._cell_text import cell_text, page_fragments
from vivepdf.ops._ocr_layer import without_rules
from vivepdf.ops.convert_tables import XlsxParams, to_xlsx
from vivepdf.rpc.progress import silent_progress

FONT = str(
    Path(__file__).resolve().parent.parent / "vivepdf" / "assets" / "fonts" / "DejaVuSans.ttf"
)
TEXT = {"fontname": "dejavu", "fontfile": FONT}
HEADER = ("Tarih", "Saat", "Kalem", "Tutar")
COLUMNS = (50, 150, 210, 410, 500)
ROW = 22


def _grid(page: pymupdf.Page, top: float, rows: list[tuple[str, ...]]) -> None:
    bottom = top + ROW * len(rows)
    for row in range(len(rows) + 1):
        page.draw_line((COLUMNS[0], top + row * ROW), (COLUMNS[-1], top + row * ROW))
    for x in COLUMNS:
        page.draw_line((x, top), (x, bottom))
    for row, values in enumerate(rows):
        for column, value in enumerate(values):
            page.insert_text(
                (COLUMNS[column] + 4, top + row * ROW + 15), value, fontsize=10, **TEXT
            )


def _body(first: int, last: int) -> list[tuple[str, ...]]:
    return [
        (f"{day:02d}.10.2026", f"{day % 10:02d}:15", f"Kalem {day}", f"{day},50 TL")
        for day in range(first, last + 1)
    ]


def _split_table(path: Path) -> Path:
    document = pymupdf.open()
    first = document.new_page()
    first.insert_text((50, 60), "Harcama listesi", fontsize=14, **TEXT)
    _grid(first, 80, [HEADER, *_body(1, 30)])
    second = document.new_page()
    _grid(second, 50, [HEADER, *_body(31, 33)])
    second.insert_text((50, 800), "2", fontsize=9, **TEXT)
    document.save(path)
    document.close()
    return path


def _convert(source: Path, tmp_path: Path, name: str, **options) -> object:
    return to_xlsx(
        XlsxParams(path=str(source), output=str(tmp_path / name), **options), silent_progress()
    )


def test_a_table_split_across_pages_becomes_one_sheet_with_one_header(tmp_path: Path) -> None:
    source = _split_table(tmp_path / "harcama.pdf")

    result = _convert(source, tmp_path, "o.xlsx")

    workbook = load_workbook(result.output)
    assert result.table_count == 1
    assert workbook.sheetnames == ["S1-S2-T1", "Text"]
    sheet = workbook["S1-S2-T1"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == HEADER
    assert len(rows) == 1 + 33
    assert rows[1][:2] == (dt.datetime(2026, 10, 1), dt.time(1, 15))
    assert rows[-1][2] == "Kalem 33"


def test_the_header_row_is_bold_and_frozen(tmp_path: Path) -> None:
    source = _split_table(tmp_path / "harcama.pdf")

    sheet = load_workbook(_convert(source, tmp_path, "o.xlsx").output)["S1-S2-T1"]

    assert sheet.freeze_panes == "A2"
    assert all(cell.font.b for cell in sheet[1])
    assert not any(cell.font.b for cell in sheet[2])
    assert sheet["B2"].number_format == "hh:mm"


def test_per_page_sheets_keep_each_page_part_on_its_own_sheet(tmp_path: Path) -> None:
    source = _split_table(tmp_path / "harcama.pdf")

    result = _convert(source, tmp_path, "o.xlsx", sheets="page")

    assert load_workbook(result.output).sheetnames == ["S1", "S2", "Text"]
    assert result.table_count == 2


def test_a_text_only_table_without_a_bold_first_row_gets_no_header(tmp_path: Path) -> None:
    document = pymupdf.open()
    _grid(document.new_page(), 80, [("Ali", "Veli", "Ayşe", "Can"), ("Ece", "Efe", "Su", "Ata")])
    source = tmp_path / "isimler.pdf"
    document.save(source)

    sheet = load_workbook(_convert(source, tmp_path, "o.xlsx").output).active

    assert sheet.freeze_panes is None
    assert not any(cell.font.b for cell in sheet[1])


def test_csv_uses_the_chosen_separator(tmp_path: Path) -> None:
    source = _split_table(tmp_path / "harcama.pdf")

    result = _convert(source, tmp_path, "o.csv", format="csv", delimiter=";")

    with Path(result.output).open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle, delimiter=";"))
    assert rows[0] == list(HEADER)
    assert rows[1] == ["01.10.2026", "01:15", "Kalem 1", "1,50 TL"]
    assert len(rows) == 1 + 33


def test_text_wrapped_inside_a_cell_comes_out_as_one_line(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_font(fontname="dejavu", fontfile=FONT)
    edges, tops = (72, 172, 372), (100, 124, 190, 240)
    for y in tops:
        page.draw_line((edges[0], y), (edges[-1], y))
    for x in edges:
        page.draw_line((x, tops[0]), (x, tops[-1]))
    cells = (
        ("Kalem", "Açıklama"),
        ("Tasarım", "Arayüz tasarımı, kullanıcı görüşmeleri ve üç tur düzeltme dahil edilmiştir."),
        ("Teslim", "• Kaynak dosyalar\n• Kılavuz"),
    )
    for row, values in enumerate(cells):
        for column, value in enumerate(values):
            box = pymupdf.Rect(
                edges[column] + 4, tops[row] + 3, edges[column + 1] - 4, tops[row + 1]
            )
            assert page.insert_textbox(box, value, fontsize=10, fontname="dejavu") >= 0
    source = tmp_path / "teklif.pdf"
    document.save(source)

    rows = list(load_workbook(_convert(source, tmp_path, "o.xlsx").output).active.values)

    assert rows[1][1] == (
        "Arayüz tasarımı, kullanıcı görüşmeleri ve üç tur düzeltme dahil edilmiştir."
    )
    assert rows[2][1] == "• Kaynak dosyalar\n• Kılavuz"


def test_lines_far_apart_in_a_cell_stay_apart() -> None:
    document = pymupdf.open()
    page = document.new_page(width=300, height=300)
    page.insert_text((10, 40), "Unicode karakterleri", fontsize=12, **TEXT)
    page.insert_text((10, 140), "listenin", fontsize=12, **TEXT)

    text = cell_text(page_fragments(page), (5, 20, 140, 160))

    assert text == "Unicode karakterleri\nlistenin"


def test_long_rules_are_erased_before_recognition_and_letters_stay() -> None:
    document = pymupdf.open()
    page = document.new_page(width=300, height=200)
    page.draw_line((20, 50), (280, 50), width=1.5)
    page.draw_line((150, 20), (150, 180), width=1.5)
    page.insert_text((30, 100), "Tutar", fontsize=20, **TEXT)
    pixmap = page.get_pixmap(dpi=150, colorspace=pymupdf.csRGB, alpha=False)

    cleared, lines = without_rules(pixmap)

    pixels = np.frombuffer(cleared.samples, np.uint8).reshape(cleared.height, cleared.width, 3)
    scale = 150 / 72
    assert pixels[round(50 * scale)].min() > 200
    assert pixels[:, round(150 * scale)].min() > 200
    letters = pixels[round(80 * scale) : round(105 * scale), round(30 * scale) : round(100 * scale)]
    assert letters.min() < 100
    found = sorted((round(a.x), round(a.y), round(b.x), round(b.y)) for a, b in lines)
    assert found == [(20, 50, 280, 50), (150, 20, 150, 180)]


def test_a_scanned_table_is_read_with_ocr(tmp_path: Path) -> None:
    items = ("Kitap", "Yemek", "Taksi", "Otel")
    body = [
        (day, time, item, cost)
        for (day, time, _item, cost), item in zip(_body(1, 4), items, strict=True)
    ]
    drawn = pymupdf.open()
    _grid(drawn.new_page(), 80, [HEADER, *body])
    scan = drawn[0].get_pixmap(dpi=200)
    document = pymupdf.open()
    page = document.new_page()
    page.insert_image(page.rect, pixmap=scan)
    source = tmp_path / "tarama.pdf"
    document.save(source)

    plain = _convert(source, tmp_path, "plain.xlsx")
    result = _convert(source, tmp_path, "o.xlsx", ocr=True)

    assert plain.textless_pages == [1] and plain.table_count == 0
    assert result.ocr_pages == [1]
    assert result.table_count == 1
    rows = list(load_workbook(result.output).active.values)
    assert rows[0] == HEADER
    assert rows[1][2:] == ("Kitap", 1.5)
    assert rows[4][2] == "Otel"


def test_a_table_on_a_turned_page_whose_text_is_upright_fills_its_sheet(tmp_path: Path) -> None:
    document = pymupdf.open()
    rows = [HEADER, *_body(1, 2)]
    for rotation in (90, 180, 270):
        page = document.new_page(width=595, height=842)
        page.set_rotation(rotation)
        turn = page.derotation_matrix
        bottom = 100 + ROW * len(rows)
        for row in range(len(rows) + 1):
            y = 100 + row * ROW
            page.draw_line(
                pymupdf.Point(COLUMNS[0], y) * turn, pymupdf.Point(COLUMNS[-1], y) * turn
            )
        for x in COLUMNS:
            page.draw_line(pymupdf.Point(x, 100) * turn, pymupdf.Point(x, bottom) * turn)
        for row, values in enumerate(rows):
            for column, value in enumerate(values):
                page.insert_text(
                    pymupdf.Point(COLUMNS[column] + 4, 100 + row * ROW + 15) * turn,
                    value,
                    fontsize=10,
                    rotate=rotation,
                    **TEXT,
                )
    source = tmp_path / "dönük.pdf"
    document.save(source)
    document.close()

    to_xlsx(
        XlsxParams(path=str(source), output=str(tmp_path / "dönük.xlsx"), sheets="page"),
        silent_progress(),
    )

    book = load_workbook(tmp_path / "dönük.xlsx")
    tables = [sheet for sheet in book.worksheets if sheet.title != "Text"]
    assert len(tables) == 3
    for sheet in tables:
        assert [cell.value for cell in sheet[1]] == list(HEADER)
        assert sheet.cell(3, 3).value == "Kalem 2"
    assert "Text" not in book.sheetnames
