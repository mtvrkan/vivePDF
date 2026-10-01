import csv
import gzip
import io
import subprocess
import sys
import time
import zipfile
from pathlib import Path

import pymupdf
import pytest
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from vivepdf.external import _process, libreoffice
from vivepdf.ops import _story, web
from vivepdf.ops.convert import TextSourceParams
from vivepdf.ops.convert_images import ImagesParams, ImagesToPdfParams, images_to_pdf, to_images
from vivepdf.ops.convert_slides import PptxParams, slide_size, to_pptx
from vivepdf.ops.convert_tables import XlsxParams, to_xlsx
from vivepdf.ops.convert_text import TextParams, to_html, to_text
from vivepdf.ops.convert_to_pdf import FileToPdfParams, csv_filter_options, file_to_pdf
from vivepdf.ops.epub import EpubParams, to_epub
from vivepdf.ops.web import WebPageParams, from_url
from vivepdf.rpc.errors import ErrorCode, OpError
from vivepdf.rpc.progress import silent_progress

FONT = "vivepdf/assets/fonts/DejaVuSans.ttf"


def two_colour_photo(path: Path, orientation: int) -> Path:
    picture = Image.new("RGB", (400, 200), (0, 0, 255))
    picture.paste((255, 0, 0), (0, 0, 200, 200))
    exif = picture.getexif()
    exif[0x0112] = orientation
    picture.save(path, exif=exif)
    return path


def rendered(path: str) -> pymupdf.Pixmap:
    with pymupdf.open(path) as document:
        return document[0].get_pixmap()


def is_red(pixel: tuple[int, ...]) -> bool:
    return pixel[0] > 180 and pixel[2] < 90


def is_blue(pixel: tuple[int, ...]) -> bool:
    return pixel[2] > 180 and pixel[0] < 90


def table_document(path: Path, tables_per_page: list[int], extra: str = "") -> Path:
    document = pymupdf.open()
    for count in tables_per_page:
        page = document.new_page()
        top = 80
        for table in range(count):
            for row in range(3):
                page.draw_rect(pymupdf.Rect(72, top, 300, top + 24), color=(0, 0, 0))
                page.draw_line((186, top), (186, top + 24), color=(0, 0, 0))
                page.insert_text(
                    (80, top + 16),
                    f"T{table}R{row}{extra}",
                    fontsize=10,
                    fontfile=FONT,
                    fontname="dv",
                )
                page.insert_text((194, top + 16), str(row), fontsize=10)
                top += 24
            top += 60
    document.save(path)
    document.close()
    return path


def test_photo_is_turned_upright_by_its_exif_orientation(tmp_path: Path) -> None:
    photo = two_colour_photo(tmp_path / "photo.jpg", 6)
    result = images_to_pdf(
        ImagesToPdfParams(images=[str(photo)], page_size="image", output=str(tmp_path / "o.pdf")),
        silent_progress(),
    )
    pixmap = rendered(result.output)
    assert pixmap.height > pixmap.width
    assert is_red(pixmap.pixel(pixmap.width // 2, 5))
    assert is_blue(pixmap.pixel(pixmap.width // 2, pixmap.height - 5))


def test_mirrored_exif_orientation_is_applied_too(tmp_path: Path) -> None:
    photo = two_colour_photo(tmp_path / "mirror.jpg", 2)
    result = images_to_pdf(
        ImagesToPdfParams(images=[str(photo)], page_size="image", output=str(tmp_path / "o.pdf")),
        silent_progress(),
    )
    pixmap = rendered(result.output)
    assert is_blue(pixmap.pixel(5, pixmap.height // 2))
    assert is_red(pixmap.pixel(pixmap.width - 5, pixmap.height // 2))


def test_webp_and_every_tiff_frame_become_pages(tmp_path: Path) -> None:
    Image.new("RGB", (320, 240), (0, 200, 0)).save(tmp_path / "a.webp")
    frames = [Image.new("RGB", (200, 300), colour) for colour in ((255, 0, 0), (0, 0, 255))]
    frames[0].save(tmp_path / "b.tiff", save_all=True, append_images=frames[1:])
    result = images_to_pdf(
        ImagesToPdfParams(
            images=[str(tmp_path / "a.webp"), str(tmp_path / "b.tiff")],
            output=str(tmp_path / "o.pdf"),
        ),
        silent_progress(),
    )
    assert result.page_count == 3


def test_fill_covers_the_area_inside_the_margin_and_leaves_the_margin_white(
    tmp_path: Path,
) -> None:
    photo = two_colour_photo(tmp_path / "p.jpg", 1)
    result = images_to_pdf(
        ImagesToPdfParams(
            images=[str(photo)],
            orientation="portrait",
            fit="fill",
            margin=40,
            output=str(tmp_path / "o.pdf"),
        ),
        silent_progress(),
    )
    pixmap = rendered(result.output)
    assert pixmap.pixel(10, 10)[:3] == (255, 255, 255)
    assert is_red(pixmap.pixel(pixmap.width // 2 - 20, 60))
    assert is_blue(pixmap.pixel(pixmap.width // 2 + 20, pixmap.height - 60))


def test_an_unreadable_file_in_a_folder_is_skipped_but_a_chosen_one_is_refused(
    tmp_path: Path,
) -> None:
    folder = tmp_path / "photos"
    folder.mkdir()
    Image.new("RGB", (50, 50), (1, 2, 3)).save(folder / "good.png")
    (folder / "broken.jpg").write_bytes(b"not an image")
    result = images_to_pdf(
        ImagesToPdfParams(folders=[str(folder)], output=str(tmp_path / "o.pdf")), silent_progress()
    )
    assert result.page_count == 1
    assert result.skipped == ["broken.jpg"]
    with pytest.raises(OpError) as caught:
        images_to_pdf(
            ImagesToPdfParams(images=[str(folder / "broken.jpg")], output=str(tmp_path / "x.pdf")),
            silent_progress(),
        )
    assert caught.value.code == ErrorCode.INVALID_PARAMS


def test_xlsx_drops_characters_a_worksheet_cannot_hold(tmp_path: Path) -> None:
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "abc\x01\x02\x1bdef", fontname="helv")
    source = tmp_path / "control.pdf"
    document.save(source)
    document.close()
    result = to_xlsx(
        XlsxParams(path=str(source), output=str(tmp_path / "o.xlsx")), silent_progress()
    )
    sheet = load_workbook(result.output).active
    assert sheet.cell(1, 2).value == "abcdef"


def test_xlsx_can_put_tables_per_page_or_on_one_sheet(tmp_path: Path) -> None:
    source = table_document(tmp_path / "t.pdf", [2, 1])
    per_table = to_xlsx(
        XlsxParams(path=str(source), output=str(tmp_path / "a.xlsx")), silent_progress()
    )
    assert load_workbook(per_table.output).sheetnames == ["S1-T1", "S1-T2", "S2-T1"]
    per_page = to_xlsx(
        XlsxParams(path=str(source), output=str(tmp_path / "b.xlsx"), sheets="page"),
        silent_progress(),
    )
    assert load_workbook(per_page.output).sheetnames == ["S1", "S2"]
    single = to_xlsx(
        XlsxParams(path=str(source), output=str(tmp_path / "c.xlsx"), sheets="single"),
        silent_progress(),
    )
    workbook = load_workbook(single.output)
    assert workbook.sheetnames == ["S1-S2"]
    assert single.table_count == 3
    assert workbook.active.max_row == 3 * 3 + 2


def test_csv_holds_every_table_with_a_blank_row_between(tmp_path: Path) -> None:
    source = table_document(tmp_path / "t.pdf", [1, 1], extra="ş")
    result = to_xlsx(
        XlsxParams(path=str(source), output=str(tmp_path / "tables.xlsx"), format="csv"),
        silent_progress(),
    )
    assert result.output.endswith(".csv")
    raw = Path(result.output).read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    assert len(rows) == 7
    assert rows[3] == []
    assert rows[0][0] == "T0R0ş"


def test_pptx_keeps_the_proportions_of_a_page_unlike_the_first(tmp_path: Path) -> None:
    document = pymupdf.open()
    document.new_page(width=595, height=842)
    document.new_page(width=842, height=595)
    source = tmp_path / "mixed.pdf"
    document.save(source)
    document.close()
    result = to_pptx(
        PptxParams(path=str(source), output=str(tmp_path / "o.pptx"), dpi=50), silent_progress()
    )
    slides = list(Presentation(result.output).slides)
    landscape = next(iter(slides[1].shapes))
    assert landscape.width > landscape.height
    assert abs(landscape.width / landscape.height - 842 / 595) < 0.01


def test_slide_size_stays_inside_the_limits_powerpoint_accepts() -> None:
    width, height = slide_size(pymupdf.Rect(0, 0, 14400, 7200))
    assert width <= 51206400
    assert abs(width / height - 2) < 0.001
    small_width, small_height = slide_size(pymupdf.Rect(0, 0, 36, 72))
    assert min(small_width, small_height) >= 914400


def test_text_pages_do_not_run_into_each_other(tmp_path: Path) -> None:
    document = pymupdf.open()
    for word in ("ilk", "ikinci"):
        document.new_page().insert_text((72, 72), word)
    source = tmp_path / "two.pdf"
    document.save(source)
    document.close()
    result = to_text(
        TextParams(path=str(source), output=str(tmp_path / "o.txt")), silent_progress()
    )
    assert Path(result.output).read_text(encoding="utf-8") == "ilk\n\fikinci\n"


def test_html_export_escapes_the_document_title(tmp_path: Path) -> None:
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "x")
    document.set_metadata({"title": "</title><script>alert(1)</script>"})
    source = tmp_path / "t.pdf"
    document.save(source)
    document.close()
    result = to_html(
        TextSourceParams(path=str(source), output=str(tmp_path / "o.html")), silent_progress()
    )
    markup = Path(result.output).read_text(encoding="utf-8")
    assert "<script>" not in markup
    assert "&lt;/title&gt;" in markup


def test_page_images_can_be_joined_into_one_tall_picture(sample_pdf: Path, tmp_path: Path) -> None:
    result = to_images(
        ImagesParams(path=str(sample_pdf), output_dir=str(tmp_path), dpi=36, single=True),
        silent_progress(),
    )
    assert len(result.outputs) == 1
    with Image.open(result.outputs[0]) as picture:
        assert picture.height > 2.9 * picture.width * 842 / 595 * 0.99


def test_png_pages_can_keep_a_transparent_background(sample_pdf: Path, tmp_path: Path) -> None:
    result = to_images(
        ImagesParams(path=str(sample_pdf), output_dir=str(tmp_path), dpi=36, transparent=True),
        silent_progress(),
    )
    with Image.open(result.outputs[0]) as picture:
        assert picture.mode == "RGBA"
        assert picture.getpixel((2, 2))[3] == 0


def test_an_existing_page_image_stops_the_run_before_anything_is_written(
    sample_pdf: Path, tmp_path: Path
) -> None:
    (tmp_path / "sample-3.png").write_bytes(b"old")
    with pytest.raises(OpError):
        to_images(ImagesParams(path=str(sample_pdf), output_dir=str(tmp_path)), silent_progress())
    assert not (tmp_path / "sample-1.png").exists()


def test_a_joined_picture_beyond_the_format_limit_is_refused(tmp_path: Path) -> None:
    document = pymupdf.open()
    for _ in range(4):
        document.new_page(width=595, height=842)
    source = tmp_path / "long.pdf"
    document.save(source)
    document.close()
    with pytest.raises(OpError) as caught:
        to_images(
            ImagesParams(
                path=str(source), output_dir=str(tmp_path), format="webp", dpi=400, single=True
            ),
            silent_progress(),
        )
    assert caught.value.data["reason"] == "imageTooLarge"


def test_text_in_the_local_code_page_is_read_as_such(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(_story.locale, "getpreferredencoding", lambda _do_setlocale=True: "cp1254")
    source = tmp_path / "düz metin.txt"
    source.write_bytes("Türkçe ğüşıöç İ".encode("cp1254"))
    result = file_to_pdf(
        FileToPdfParams(path=str(source), output=str(tmp_path / "o.pdf")), silent_progress()
    )
    with pymupdf.open(result.output) as document:
        assert "Türkçe ğüşıöç İ" in document[0].get_text()


def test_decode_text_follows_byte_order_marks_and_declared_charsets() -> None:
    assert _story.decode_text("ğüş".encode("utf-16")) == "ğüş"
    assert _story.decode_text("﻿BOM".encode()) == "BOM"
    page = '<meta charset="windows-1254"><p>ş</p>'.encode("cp1254")
    assert "ş" in _story.decode_text(page, _story.declared_charset(page))
    assert _story.decode_text(b"abc", "no-such-codec") == "abc"


def test_csv_separator_is_detected_and_the_copy_is_utf8(tmp_path: Path) -> None:
    source = tmp_path / "veri.csv"
    source.write_bytes("ad;yaş\nAyşe;30\n".encode("cp1254"))
    staging = tmp_path / "staging"
    staging.mkdir()
    normalized, infilter = csv_filter_options(source, staging)
    assert infilter.endswith(":59,34,76,1")
    assert normalized.read_text(encoding="utf-8").startswith("ad;ya")


def test_password_protected_office_file_is_named_as_such(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "şifreli.docx"
    source.write_bytes(
        libreoffice.OLE_SIGNATURE + b"\x00" * 512 + libreoffice.ENCRYPTED_PACKAGE_NAME
    )
    monkeypatch.setattr(libreoffice, "find_soffice", lambda: tmp_path / "soffice.exe")
    with pytest.raises(OpError) as caught:
        libreoffice.convert_to_pdf(source, tmp_path / "out")
    assert caught.value.code == ErrorCode.ENCRYPTED
    assert caught.value.data == {"reason": "officeEncrypted"}


def test_an_old_word_file_renamed_to_docx_is_not_called_password_protected(
    tmp_path: Path,
) -> None:
    renamed = tmp_path / "eski.docx"
    renamed.write_bytes(
        libreoffice.OLE_SIGNATURE + b"\x00" * 4096 + "WordDocument".encode("utf-16-le")
    )
    assert libreoffice.is_encrypted_office(renamed) is False


def test_a_password_protected_odf_file_is_recognised(tmp_path: Path) -> None:
    locked = tmp_path / "gizli.odt"
    with zipfile.ZipFile(locked, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr(
            "META-INF/manifest.xml",
            '<manifest:manifest><manifest:file-entry manifest:full-path="content.xml">'
            "<manifest:encryption-data/></manifest:file-entry></manifest:manifest>",
        )
    plain = tmp_path / "acik.odt"
    with zipfile.ZipFile(plain, "w") as archive:
        archive.writestr("META-INF/manifest.xml", "<manifest:manifest/>")
    assert libreoffice.is_encrypted_office(locked) is True
    assert libreoffice.is_encrypted_office(plain) is False


def test_a_load_failure_reported_with_exit_code_zero_gets_its_own_reason(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "bozuk.docx"
    source.write_bytes(b"PK\x03\x04garbage")

    def quiet_failure(arguments: list[str], **_kwargs: object) -> subprocess.CompletedProcess:
        return subprocess.CompletedProcess(
            arguments, 0, "Error: source file could not be loaded\n", ""
        )

    monkeypatch.setattr(libreoffice, "find_soffice", lambda: tmp_path / "soffice.exe")
    monkeypatch.setattr(libreoffice, "run_tool", quiet_failure)
    with pytest.raises(OpError) as caught:
        libreoffice.convert_to_pdf(source, tmp_path / "out")
    assert caught.value.data["reason"] == "officeUnreadable"


def test_a_cancelled_tool_run_is_stopped_quickly() -> None:
    started = time.monotonic()
    calls = {"count": 0}

    def cancel_on_third_check() -> None:
        calls["count"] += 1
        if calls["count"] >= 3:
            raise OpError(ErrorCode.CANCELLED, "operation cancelled")

    with pytest.raises(OpError) as caught:
        _process.run_tool(
            [sys.executable, "-c", "import time; time.sleep(60)"],
            timeout=60,
            tool="Probe",
            check_cancelled=cancel_on_third_check,
        )
    assert caught.value.code == ErrorCode.CANCELLED
    assert time.monotonic() - started < 20


def test_a_file_libreoffice_cannot_load_gets_its_own_reason(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "bozuk.docx"
    source.write_bytes(b"PK\x03\x04garbage")

    def failing_tool(*_args: object, **_kwargs: object) -> None:
        raise OpError(
            ErrorCode.EXTERNAL_TOOL_FAILED,
            "LibreOffice failed (exit 1)",
            {"tool": "LibreOffice", "detail": "Error: source file could not be loaded"},
        )

    monkeypatch.setattr(libreoffice, "find_soffice", lambda: tmp_path / "soffice.exe")
    monkeypatch.setattr(libreoffice, "run_tool", failing_tool)
    with pytest.raises(OpError) as caught:
        libreoffice.convert_to_pdf(source, tmp_path / "out")
    assert caught.value.data["reason"] == "officeUnreadable"


def test_a_tool_that_hangs_is_stopped_with_its_children() -> None:
    started = time.monotonic()
    script = (
        "import subprocess, sys, time; "
        "subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)']); "
        "time.sleep(60)"
    )
    with pytest.raises(OpError) as caught:
        _process.run_tool([sys.executable, "-c", script], timeout=1.5, tool="Probe")
    assert caught.value.data["reason"] == "timeout"
    assert time.monotonic() - started < 20


def test_a_web_page_with_an_unknown_charset_still_converts(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    body = "<html><body><p>" + "Türkçe metin. " * 60 + "</p></body></html>"
    monkeypatch.setattr(
        web,
        "fetch",
        lambda url, limit, allow_internal=False, **_options: (
            body.encode(),
            "text/html; charset=bogus-xyz",
            url,
        ),
    )
    result = from_url(
        WebPageParams(url="https://example.com/", output=str(tmp_path / "o.pdf")),
        silent_progress(),
    )
    with pymupdf.open(result.output) as document:
        assert "Türkçe metin." in document[0].get_text()


def test_a_plain_text_address_is_laid_out_as_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        web,
        "fetch",
        lambda url, limit, allow_internal=False, **_options: (
            "Düz metin satırı".encode(),
            "text/plain; charset=utf-8",
            url,
        ),
    )
    result = from_url(
        WebPageParams(url="https://example.com/notes.txt", output=str(tmp_path / "o.pdf")),
        silent_progress(),
    )
    with pymupdf.open(result.output) as document:
        assert "Düz metin satırı" in document[0].get_text()


def test_a_compressed_response_is_not_inflated_past_the_limit() -> None:
    bomb = gzip.compress(b"a" * 5000)
    with pytest.raises(OpError) as caught:
        web._decompress(bomb, "gzip", 1000)
    assert caught.value.data["reason"] == "pageTooLarge"
    assert web._decompress(gzip.compress(b"small"), "gzip", 1000) == b"small"


def test_epub_carries_a_cover_picture_of_the_first_page(sample_pdf: Path, tmp_path: Path) -> None:
    result = to_epub(
        EpubParams(path=str(sample_pdf), output=str(tmp_path / "b.epub")), silent_progress()
    )
    assert result.cover
    with zipfile.ZipFile(result.output) as archive:
        opf = archive.read("OEBPS/content.opf").decode()
        assert 'properties="cover-image"' in opf
        assert '<meta name="cover" content="cover-image"/>' in opf
        assert archive.read("OEBPS/cover.jpg")[:2] == b"\xff\xd8"
    without = to_epub(
        EpubParams(path=str(sample_pdf), output=str(tmp_path / "c.epub"), cover=False),
        silent_progress(),
    )
    with zipfile.ZipFile(without.output) as archive:
        assert "OEBPS/cover.jpg" not in archive.namelist()


def test_a_single_joined_page_is_exactly_the_page_picture(sample_pdf: Path, tmp_path: Path) -> None:
    result = to_images(
        ImagesParams(
            path=str(sample_pdf), output_dir=str(tmp_path), dpi=72, single=True, pages="1"
        ),
        silent_progress(),
    )
    with pymupdf.open(sample_pdf) as document:
        expected = document[0].get_pixmap(dpi=72)
    with Image.open(result.outputs[0]) as picture:
        assert picture.size == (expected.width, expected.height)
