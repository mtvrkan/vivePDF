import os
import re
import subprocess
import sys
import threading
import time
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from pdf2docx.table.TableBlock import TableBlock

from vivepdf.ops import _docx_progress as docx_progress
from vivepdf.ops._pptx_layout import font_family, slide_text
from vivepdf.ops.convert import DocxParams, to_docx
from vivepdf.ops.convert_slides import PptxParams, to_pptx
from vivepdf.ops.convert_text import MarkdownParams, _markdown_pages, to_markdown
from vivepdf.rpc.errors import ErrorCode, OpError
from vivepdf.rpc.progress import Progress, silent_progress

SIDECAR = Path(__file__).resolve().parent.parent
FONT = str(SIDECAR / "vivepdf" / "assets" / "fonts" / "DejaVuSans.ttf")
EMU_PER_POINT = 12700


def _write(page: pymupdf.Page, point: tuple[float, float], text: str, size: float = 12) -> None:
    page.insert_text(point, text, fontsize=size, fontname="dejavu", fontfile=FONT)


def _cancelling(predicate) -> tuple[Progress, list]:
    events: list[tuple[float, str | None, dict | None]] = []
    cancel = threading.Event()

    def sink(value, message, detail):
        events.append((value, message, detail))
        if predicate(events):
            cancel.set()

    return Progress(sink, cancel), events


@pytest.fixture
def outlined_pdf(tmp_path: Path) -> Path:
    document = pymupdf.open()
    first = document.new_page(width=595, height=842)
    _write(first, (72, 100), "Doğrusal cebir özeti", 20)
    _write(first, (72, 140), "Matrisler ve vektörler üzerine kısa bir giriş.")
    outlined = document.new_page(width=595, height=842)
    shape = outlined.new_shape()
    for index in range(2200):
        x = 40 + (index % 100) * 5
        y = 60 + (index // 100) * 12
        shape.draw_line((x, y), (x + 3, y + 8))
    shape.finish(color=(0, 0, 0), width=0.5)
    shape.commit()
    path = tmp_path / "çizgisel metin ş.pdf"
    document.save(path)
    document.close()
    return path


@pytest.fixture
def headed_pdf(tmp_path: Path) -> Path:
    document = pymupdf.open()
    for number in range(4):
        page = document.new_page(width=595, height=842)
        _write(page, (72, 90), f"Bölüm {number + 1}", 26 if number == 0 else 18)
        _write(page, (72, 130), f"Alt başlık {number + 1}", 14)
        for line in range(6):
            _write(page, (72, 170 + line * 18), f"Satır {line + 1} gövde metni ğüşiöç", 10)
    path = tmp_path / "başlıklı.pdf"
    document.save(path)
    document.close()
    return path


def test_markdown_matches_the_whole_document_conversion(headed_pdf: Path) -> None:
    import pymupdf4llm

    with pymupdf.open(headed_pdf) as whole:
        expected = pymupdf4llm.to_markdown(
            whole, pages=list(range(4)), write_images=False, show_progress=False
        )
    with pymupdf.open(headed_pdf) as paged:
        markdown, textless = _markdown_pages(paged, [2, 0, 1, 3, 1], silent_progress())
    assert markdown == expected
    assert textless == []
    assert "Bölüm 1" in markdown


def test_markdown_skips_outlined_pages_and_names_them(outlined_pdf: Path, tmp_path: Path) -> None:
    progress, events = _cancelling(lambda _events: False)
    started = time.monotonic()
    result = to_markdown(
        MarkdownParams(path=str(outlined_pdf), output=str(tmp_path / "out.md")), progress
    )
    assert time.monotonic() - started < 60
    assert result.textless_pages == [2]
    text = Path(result.output).read_text(encoding="utf-8")
    assert "Doğrusal cebir özeti" in text
    pages = [event[2] for event in events if event[1] == "progress.convertingPages"]
    assert pages == [{"current": 1, "total": 2}, {"current": 2, "total": 2}]


def test_markdown_skips_the_layout_model_on_blank_pages(
    headed_pdf: Path, tmp_path: Path, monkeypatch
) -> None:
    from pymupdf4llm.helpers import document_layout

    with pymupdf.open(headed_pdf) as document:
        for _ in range(3):
            document.new_page(width=595, height=842)
        padded = tmp_path / "boş sayfalı.pdf"
        document.save(padded)
    with pymupdf.open(headed_pdf) as whole:
        expected, _ = _markdown_pages(whole, list(range(4)), silent_progress())
    parsed: list[list[int]] = []
    original = document_layout.parse_document

    def counted(document, pages, **options):
        parsed.append(pages)
        return original(document, pages=pages, **options)

    monkeypatch.setattr(document_layout, "parse_document", counted)
    monkeypatch.setattr(document_layout, "select_ocr_function", lambda: None)
    result = to_markdown(
        MarkdownParams(path=str(padded), output=str(tmp_path / "out.md")), silent_progress()
    )
    assert parsed == [[0], [1], [2], [3]]
    assert result.textless_pages == [5, 6, 7]
    assert Path(result.output).read_text(encoding="utf-8") == expected


def test_markdown_cancel_between_pages_writes_nothing(headed_pdf: Path, tmp_path: Path) -> None:
    progress, events = _cancelling(
        lambda events: sum(1 for event in events if event[1] == "progress.convertingPages") >= 2
    )
    target = tmp_path / "iptal.md"
    with pytest.raises(OpError) as caught:
        to_markdown(MarkdownParams(path=str(headed_pdf), output=str(target)), progress)
    assert caught.value.code == ErrorCode.CANCELLED
    assert not target.exists()
    assert sum(1 for event in events if event[1] == "progress.convertingPages") == 2


def test_markdown_rejects_an_out_of_range_page(headed_pdf: Path, tmp_path: Path) -> None:
    with pytest.raises(OpError) as caught:
        to_markdown(
            MarkdownParams(path=str(headed_pdf), output=str(tmp_path / "x.md"), pages="9"),
            silent_progress(),
        )
    assert caught.value.code == ErrorCode.INVALID_PARAMS


@pytest.fixture
def slide_pdf(tmp_path: Path) -> Path:
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    page.draw_rect(pymupdf.Rect(60, 60, 535, 200), color=(0, 0, 1), fill=(0.9, 0.9, 1))
    _write(page, (72, 100), "Sunum başlığı Çağ", 24)
    page.insert_text((72, 300), "Kalin satir", fontsize=14, fontname="hebo", color=(1, 0, 0))
    page.insert_text((72, 400), "Gizli OCR katmanı", fontsize=12, render_mode=3)
    rotated = document.new_page(width=595, height=842)
    _write(rotated, (72, 100), "Döndürülmüş sayfa", 16)
    rotated.set_rotation(90)
    path = tmp_path / "sunum ğ.pdf"
    document.save(path)
    document.close()
    return path


def _text_shapes(slide) -> list:
    return [shape for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text]


def test_pptx_editable_slides_carry_real_text_boxes(slide_pdf: Path, tmp_path: Path) -> None:
    from pptx import Presentation

    result = to_pptx(
        PptxParams(path=str(slide_pdf), output=str(tmp_path / "deck.pptx"), dpi=60),
        silent_progress(),
    )
    presentation = Presentation(result.output)
    first = presentation.slides[0]
    shapes = {shape.text_frame.text: shape for shape in _text_shapes(first)}
    assert set(shapes) == {"Sunum başlığı Çağ", "Kalin satir"}
    assert any(shape.shape_type == 13 for shape in first.shapes)
    assert not first.has_notes_slide
    ratio = presentation.slide_width / 595
    with pymupdf.open(slide_pdf) as document:
        spans = {
            span["text"]: span
            for block in document[0].get_text("dict")["blocks"]
            for line in block.get("lines", [])
            for span in line["spans"]
        }
    for text, shape in shapes.items():
        left, top = spans[text]["bbox"][:2]
        assert abs(shape.left / ratio - left) < 3
        assert abs(shape.top / ratio - top) < 3
        run = shape.text_frame.paragraphs[0].runs[0]
        assert abs(run.font.size / EMU_PER_POINT - spans[text]["size"] * ratio / EMU_PER_POINT) < 1
    bold = shapes["Kalin satir"].text_frame.paragraphs[0].runs[0].font
    assert bold.bold is True
    assert str(bold.color.rgb) == "FF0000"
    assert bold.name == "Helvetica" or bold.name == "Arial"


def test_pptx_rotated_page_keeps_its_text_where_it_is_shown(
    slide_pdf: Path, tmp_path: Path
) -> None:
    from pptx import Presentation

    result = to_pptx(
        PptxParams(path=str(slide_pdf), output=str(tmp_path / "deck.pptx"), dpi=60, pages="2"),
        silent_progress(),
    )
    presentation = Presentation(result.output)
    (shape,) = _text_shapes(presentation.slides[0])
    assert shape.text_frame.text == "Döndürülmüş sayfa"
    assert shape.rotation == 90
    with pymupdf.open(slide_pdf) as document:
        page = document[1]
        line = page.get_text("dict")["blocks"][0]["lines"][0]
        shown = pymupdf.Rect(line["bbox"]) * page.rotation_matrix
        ratio = presentation.slide_width / page.rect.width
    center_x = (shape.left + shape.width / 2) / ratio
    center_y = (shape.top + shape.height / 2) / ratio
    assert abs(center_x - (shown.x0 + shown.x1) / 2) < 3
    assert abs(center_y - (shown.y0 + shown.y1) / 2) < 3


def test_pptx_image_mode_keeps_the_picture_and_notes(slide_pdf: Path, tmp_path: Path) -> None:
    from pptx import Presentation

    result = to_pptx(
        PptxParams(path=str(slide_pdf), output=str(tmp_path / "deck.pptx"), dpi=60, mode="image"),
        silent_progress(),
    )
    slide = Presentation(result.output).slides[0]
    assert _text_shapes(slide) == []
    assert "Sunum başlığı" in slide.notes_slide.notes_text_frame.text


def test_pptx_editable_reads_an_encrypted_document(encrypted_pdf: Path, tmp_path: Path) -> None:
    from pptx import Presentation

    result = to_pptx(
        PptxParams(
            path=str(encrypted_pdf), password="secret", output=str(tmp_path / "e.pptx"), dpi=50
        ),
        silent_progress(),
    )
    presentation = Presentation(result.output)
    assert [shape.text_frame.text for shape in _text_shapes(presentation.slides[2])] == ["Page 3"]


def test_pptx_rejects_an_unknown_mode(slide_pdf: Path, tmp_path: Path) -> None:
    with pytest.raises(ValueError):
        PptxParams(path=str(slide_pdf), output=str(tmp_path / "x.pptx"), mode="outline")


def test_font_family_and_slide_text_are_normalised() -> None:
    assert font_family("ABCDEF+TimesNewRomanPSMT") == "Times New Roman"
    assert font_family("ArialMT") == "Arial"
    assert font_family("Calibri-Bold") == "Calibri"
    assert font_family("") == "Arial"
    assert slide_text("a\x00b\x0bc\tç") == "abc\tç"


@pytest.fixture
def merged_table_pdf(tmp_path: Path) -> Path:
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    left, top, width, height = 72, 100, 90, 24
    rows, columns = 6, 5
    for row in range(rows + 1):
        page.draw_line((left, top + row * height), (left + columns * width, top + row * height))
    for column in range(columns + 1):
        x = left + column * width
        page.draw_line((x, top), (x, top + height))
        page.draw_line((x, top + 3 * height), (x, top + rows * height))
        if column in (0, 1, columns):
            page.draw_line((x, top + height), (x, top + 3 * height))
    for row in range(rows):
        for column in range(columns):
            origin = (left + column * width + 6, top + row * height + 16)
            _write(page, origin, f"h{row}{column}ş", 9)
    path = tmp_path / "birleşik tablo.pdf"
    document.save(path)
    document.close()
    return path


def _document_xml(path: str) -> str:
    with zipfile.ZipFile(path) as archive:
        return re.sub(r'r:(id|embed)="rId\d+"', "", archive.read("word/document.xml").decode())


def test_docx_cached_table_cells_write_the_same_document(
    merged_table_pdf: Path, tmp_path: Path, monkeypatch
) -> None:
    fast = to_docx(
        DocxParams(path=str(merged_table_pdf), output=str(tmp_path / "fast.docx")),
        silent_progress(),
    )
    monkeypatch.setattr(docx_progress, "stabilise_pdf2docx", lambda: None)
    monkeypatch.setattr(TableBlock, "make_docx", docx_progress._plain_table_make_docx)
    plain = to_docx(
        DocxParams(path=str(merged_table_pdf), output=str(tmp_path / "plain.docx")),
        silent_progress(),
    )
    body = _document_xml(fast.output)
    assert "gridSpan" in body
    assert body == _document_xml(plain.output)
    assert Document(fast.output).tables


def test_docx_overlapping_links_always_pick_the_first(tmp_path: Path) -> None:
    source = tmp_path / "bağlantı.pdf"
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    _write(page, (72, 100), "Bağlantı metni burada", 14)
    area = pymupdf.Rect(70, 86, 260, 104)
    for uri in ("https://ilk.example/", "https://ikinci.example/", "https://ucuncu.example/"):
        page.insert_link({"kind": pymupdf.LINK_URI, "from": area, "uri": uri})
    document.save(source)
    document.close()
    targets = set()
    for attempt in range(4):
        result = to_docx(
            DocxParams(path=str(source), output=str(tmp_path / f"l{attempt}.docx")),
            silent_progress(),
        )
        with zipfile.ZipFile(result.output) as archive:
            relations = archive.read("word/_rels/document.xml.rels").decode()
        targets.add(tuple(re.findall(r'Target="(https://[^"]+)"', relations)))
    assert targets == {("https://ilk.example/",)}


def test_docx_cancel_during_page_parsing_stops_before_the_next_page(
    tmp_path: Path,
) -> None:
    source = tmp_path / "parça.pdf"
    document = pymupdf.open()
    for number in range(5):
        _write(document.new_page(), (72, 100), f"Sayfa {number + 1}")
    document.save(source)
    document.close()
    parse_start = docx_progress.PARSING_BAND[0]
    progress, events = _cancelling(
        lambda events: events[-1][1] == "progress.converting" and events[-1][0] >= parse_start
    )
    target = tmp_path / "iptal.docx"
    with pytest.raises(OpError) as caught:
        to_docx(DocxParams(path=str(source), output=str(target)), progress)
    assert caught.value.code == ErrorCode.CANCELLED
    assert not target.exists()
    parsing = [event for event in events if event[1] == "progress.converting" and event[0] >= 0.3]
    assert len(parsing) == 1


WORKER_PARENT = """
import multiprocessing
import sys
import time

from vivepdf.ops import _docx_parallel as parallel

if __name__ == "__main__":
    context = multiprocessing.get_context("spawn")
    events = context.Queue()
    pool = context.Pool(2, initializer=parallel._start_worker, initargs=(sys.argv[1], None, events))
    jobs = [pool.apply_async(time.sleep, (120,)) for _ in range(2)]
    print(" ".join(str(process.pid) for process in pool._pool), flush=True)
    time.sleep(600)
"""


def test_pool_workers_exit_when_the_engine_is_killed(sample_pdf: Path, tmp_path: Path) -> None:
    psutil = pytest.importorskip("psutil")
    script = tmp_path / "engine.py"
    script.write_text(WORKER_PARENT, encoding="utf-8")
    environment = {**os.environ, "PYTHONPATH": str(SIDECAR)}
    parent = subprocess.Popen(
        [sys.executable, str(script), str(sample_pdf)],
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        cwd=str(SIDECAR),
        env=environment,
        text=True,
    )
    try:
        pids = [int(pid) for pid in parent.stdout.readline().split()]
        assert len(pids) == 2
        time.sleep(8)
        assert all(psutil.pid_exists(pid) for pid in pids)
    finally:
        parent.kill()
        parent.wait(timeout=30)
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline and any(psutil.pid_exists(pid) for pid in pids):
        time.sleep(0.25)
    survivors = [pid for pid in pids if psutil.pid_exists(pid)]
    for pid in survivors:
        psutil.Process(pid).kill()
    assert survivors == []
