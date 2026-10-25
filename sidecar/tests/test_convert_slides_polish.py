import io
import zipfile
from pathlib import Path

import numpy as np
import pymupdf
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from vivepdf.ops.convert_slides import PptxParams, to_pptx
from vivepdf.rpc.progress import silent_progress

SIDECAR = Path(__file__).resolve().parent.parent
FONT = str(SIDECAR / "vivepdf" / "assets" / "fonts" / "DejaVuSans.ttf")
INK = (40, 40, 140)
LINES = [((60, 100), "Fatura no: 4821"), ((60, 140), "Toplam 15.000,00 TL")]


def _presentation(document: pymupdf.Document, tmp_path: Path):
    tmp_path.mkdir(parents=True, exist_ok=True)
    source = tmp_path / "kaynak.pdf"
    document.save(source)
    document.close()
    result = to_pptx(
        PptxParams(path=str(source), output=str(tmp_path / "sunum.pptx")), silent_progress()
    )
    return Presentation(result.output), Path(result.output)


def _texts(slide) -> list[str]:
    return [
        shape.text_frame.text
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and shape.text_frame.text.strip()
    ]


def _background(slide) -> Image.Image:
    (picture,) = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    return Image.open(io.BytesIO(picture.image.blob)).convert("RGB")


def _scan(rotation: int = 0) -> pymupdf.Document:
    printed = pymupdf.open()
    paper = printed.new_page(width=400, height=300)
    for origin, text in LINES:
        paper.insert_text(origin, text, fontsize=16, fontname="hebo", color=[c / 255 for c in INK])
    picture = paper.get_pixmap(dpi=150).tobytes("png")
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_image(page.rect, stream=picture)
    for origin, text in LINES:
        page.insert_text(origin, text, fontsize=12, fontname="helv", render_mode=3)
    page.set_rotation(rotation)
    return document


def test_a_recognised_scan_gets_editable_text_over_a_scan_without_those_words(
    tmp_path: Path,
) -> None:
    presentation, _ = _presentation(_scan(), tmp_path)
    slide = presentation.slides[0]

    assert _texts(slide) == [text for _origin, text in LINES]
    background = np.asarray(_background(slide).resize((400, 300)))
    words = background[86:146, 55:260]
    assert (np.abs(words.astype(int) - 255).sum(axis=2) > 120).sum() == 0
    runs = slide.shapes[1].text_frame.paragraphs[0].runs
    colour = runs[0].font.color.rgb
    assert abs(colour[2] - INK[2]) < 40 and colour[0] < 90
    assert runs[0].font.name == "Arial"


def test_the_words_of_a_recognised_scan_fill_the_width_of_the_scanned_words(
    tmp_path: Path,
) -> None:
    presentation, _ = _presentation(_scan(), tmp_path)
    slide = presentation.slides[0]

    printed = pymupdf.get_text_length(LINES[1][1], fontname="hebo", fontsize=16)
    (box,) = [
        shape
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text == LINES[1][1]
    ]
    assert abs(box.width / presentation.slide_width * 400 - printed) < printed * 0.08


def test_a_turned_scan_page_is_cleaned_as_well(tmp_path: Path) -> None:
    presentation, _ = _presentation(_scan(rotation=90), tmp_path)
    slide = presentation.slides[0]

    assert _texts(slide) == [text for _origin, text in LINES]
    background = np.asarray(_background(slide))
    assert (np.abs(background.astype(int) - 255).sum(axis=2) > 120).sum() == 0


def test_a_page_whose_hidden_text_has_no_scan_under_it_keeps_it_hidden(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_text((60, 100), "Gizli katman", fontsize=12, render_mode=3)

    presentation, _ = _presentation(document, tmp_path)

    assert _texts(presentation.slides[0]) == []


def test_editable_slides_carry_the_page_text_in_their_notes(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_text((60, 100), "Sunum metni", fontsize=14)

    presentation, _ = _presentation(document, tmp_path)

    assert presentation.slides[0].notes_slide.notes_text_frame.text == "Sunum metni"


def test_a_photo_page_becomes_a_jpeg_and_a_text_page_stays_a_png(tmp_path: Path) -> None:
    generator = np.random.default_rng(3)
    pixels = generator.integers(0, 60, (600, 800, 3), dtype=np.uint8)
    pixels[:, :400, 0] += 190
    pixels[:, 400:, 2] += 190
    photo = Image.fromarray(pixels)
    stream = io.BytesIO()
    photo.save(stream, format="PNG")
    document = pymupdf.open()
    document.new_page(width=400, height=300).insert_image(
        pymupdf.Rect(0, 0, 400, 300), stream=stream.getvalue()
    )
    document.new_page(width=400, height=300).insert_text((60, 100), "Yazı", fontsize=14)

    _, output = _presentation(document, tmp_path)

    with zipfile.ZipFile(output) as archive:
        media = sorted(name for name in archive.namelist() if name.startswith("ppt/media/"))
    assert [name.rsplit(".", 1)[1] for name in media] == ["jpg", "png"]


def _languages(slide) -> set[str | None]:
    return {
        run._r.find(qn("a:rPr")).get("lang")
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    }


def test_turkish_text_is_marked_turkish_for_spelling(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_text((60, 100), "Başlık ışığı", fontsize=14, fontname="dejavu", fontfile=FONT)

    presentation, _ = _presentation(document, tmp_path)

    assert _languages(presentation.slides[0]) == {"tr-TR"}


def test_the_language_the_pdf_declares_wins_and_unknown_text_gets_none(tmp_path: Path) -> None:
    declared = pymupdf.open()
    declared.new_page(width=400, height=300).insert_text(
        (60, 100), "Başlık ışığı", fontsize=14, fontname="dejavu", fontfile=FONT
    )
    declared.xref_set_key(declared.pdf_catalog(), "Lang", "(en-GB)")
    english = pymupdf.open()
    english.new_page(width=400, height=300).insert_text((60, 100), "Plain title", fontsize=14)

    first, _ = _presentation(declared, tmp_path / "a")
    second, _ = _presentation(english, tmp_path / "b")

    assert _languages(first.slides[0]) == {"en-GB"}
    assert _languages(second.slides[0]) == {None}


def _justified() -> pymupdf.Document:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_text((40, 100), "bir iki uc dort", fontsize=12, fontname="helv")
    page.insert_text((40, 140), "bes alti yedi", fontsize=12, fontname="helv")
    xref = page.get_contents()[0]
    stream = document.xref_stream(xref).decode()
    document.update_stream(xref, stream.replace("BT", "BT 6 Tw", 1).encode())
    return document


def _spacing(shape) -> list[str | None]:
    return [run._r.find(qn("a:rPr")).get("spc") for run in shape.text_frame.paragraphs[0].runs]


def test_a_stretched_line_keeps_its_word_gaps_as_space_spacing(tmp_path: Path) -> None:
    presentation, _ = _presentation(_justified(), tmp_path)

    first, second = [shape for shape in presentation.slides[0].shapes if shape.has_text_frame]
    extra = str(round(6 * presentation.slide_width / 400 / 12700 * 100))
    assert [run.text for run in first.text_frame.paragraphs[0].runs] == [
        "bir",
        " ",
        "iki",
        " ",
        "uc",
        " ",
        "dort",
    ]
    assert _spacing(first) == [None, extra, None, extra, None, extra, None]
    assert _spacing(second) == [None]
