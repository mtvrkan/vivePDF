import threading
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from openpyxl import load_workbook

from vivepdf.ops._cell_text import cell_text, page_fragments
from vivepdf.ops.convert import DocxParams, to_docx
from vivepdf.ops.convert_tables import XlsxParams, to_xlsx
from vivepdf.ops.convert_text import MarkdownParams, to_markdown
from vivepdf.rpc.errors import ErrorCode, OpError
from vivepdf.rpc.progress import Progress, silent_progress

FONT = str(
    Path(__file__).resolve().parent.parent / "vivepdf" / "assets" / "fonts" / "DejaVuSans.ttf"
)


@pytest.fixture
def paged(tmp_path: Path) -> Path:
    document = pymupdf.open()
    for number in range(6):
        page = document.new_page(width=595, height=842)
        page.insert_text(
            (72, 100), f"Bölüm {number + 1} gövde metni", fontname="dejavu", fontfile=FONT
        )
    path = tmp_path / "sayfalı belge.pdf"
    document.save(path)
    document.close()
    return path


def _recording(cancel_after: int | None = None):
    events: list[tuple[float, str | None, dict | None]] = []
    cancel = threading.Event()

    def sink(value, message, detail):
        events.append((value, message, detail))
        pages = [event for event in events if event[1] == "progress.convertingPages"]
        if cancel_after is not None and len(pages) >= cancel_after:
            cancel.set()

    return Progress(sink, cancel), events


def test_docx_reports_every_page_and_keeps_the_text(paged: Path, tmp_path: Path) -> None:
    progress, events = _recording()
    result = to_docx(DocxParams(path=str(paged), output=str(tmp_path / "out.docx")), progress)
    page_events = [event for event in events if event[1] == "progress.convertingPages"]
    assert len(page_events) == 12
    assert page_events[0][2] == {"current": 1, "total": 6}
    values = [event[0] for event in events]
    assert values == sorted(values)
    text = "\n".join(paragraph.text for paragraph in Document(result.output).paragraphs)
    assert "Bölüm 1 gövde metni" in text
    assert "Bölüm 6 gövde metni" in text


def test_docx_page_range_counts_only_the_chosen_pages(paged: Path, tmp_path: Path) -> None:
    progress, events = _recording()
    result = to_docx(
        DocxParams(path=str(paged), output=str(tmp_path / "out.docx"), pages="2-3"),
        progress,
    )
    totals = {event[2]["total"] for event in events if event[1] == "progress.convertingPages"}
    assert totals == {2}
    text = "\n".join(paragraph.text for paragraph in Document(result.output).paragraphs)
    assert "Bölüm 2" in text
    assert "Bölüm 5" not in text


@pytest.mark.parametrize("cancel_after", [2, 8])
def test_docx_cancel_stops_mid_conversion_and_writes_nothing(
    paged: Path, tmp_path: Path, cancel_after: int
) -> None:
    progress, events = _recording(cancel_after)
    target = tmp_path / "out.docx"
    with pytest.raises(OpError) as caught:
        to_docx(DocxParams(path=str(paged), output=str(target)), progress)
    assert caught.value.code == ErrorCode.CANCELLED
    assert not target.exists()
    page_events = [event for event in events if event[1] == "progress.convertingPages"]
    assert len(page_events) == cancel_after


def test_xlsx_fallback_sheet_is_named_after_the_pages(paged: Path, tmp_path: Path) -> None:
    result = to_xlsx(
        XlsxParams(path=str(paged), output=str(tmp_path / "out.xlsx"), pages="2-4"),
        silent_progress(),
    )
    assert result.table_count == 0
    assert load_workbook(result.output).sheetnames == ["S2-S4"]


def _split_word_page() -> pymupdf.Page:
    document = pymupdf.open()
    page = document.new_page(width=400, height=200)
    page.insert_text((50, 60), "Ö", fontname="dejavu", fontfile=FONT, fontsize=12)
    width = pymupdf.Font(fontfile=FONT).text_length("Ö", fontsize=12)
    page.insert_text(
        (50 + width + 0.6, 60), "zellik", fontname="dejavu", fontfile=FONT, fontsize=12
    )
    page.insert_text((200, 60), "Avantaj", fontname="dejavu", fontfile=FONT, fontsize=12)
    page.insert_text((50, 90), "Hiç bitmeyen", fontname="dejavu", fontfile=FONT, fontsize=12)
    page.insert_text((50, 105), "döngüler", fontname="dejavu", fontfile=FONT, fontsize=12)
    return page


def test_cell_text_joins_letters_drawn_apart_and_keeps_word_gaps() -> None:
    page = _split_word_page()
    fragments = page_fragments(page)
    assert cell_text(fragments, (40, 45, 190, 70)) == "Özellik"
    assert cell_text(fragments, (190, 45, 360, 70)) == "Avantaj"
    assert cell_text(fragments, (40, 75, 190, 115)) == "Hiç bitmeyen\ndöngüler"
    assert cell_text(fragments, (40, 150, 190, 190)) == ""


def test_xlsx_cells_come_out_whole(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=200)
    for x in (40, 190, 360):
        page.draw_line((x, 40), (x, 120))
    for y in (40, 75, 120):
        page.draw_line((40, y), (360, y))
    page.insert_text((50, 60), "Ö", fontname="dejavu", fontfile=FONT, fontsize=12)
    width = pymupdf.Font(fontfile=FONT).text_length("Ö", fontsize=12)
    page.insert_text(
        (50 + width + 0.6, 60), "zellik", fontname="dejavu", fontfile=FONT, fontsize=12
    )
    page.insert_text((200, 60), "Avantaj", fontname="dejavu", fontfile=FONT, fontsize=12)
    page.insert_text((50, 95), "Hiç bitmeyen", fontname="dejavu", fontfile=FONT, fontsize=12)
    page.insert_text((200, 95), "Dinamik", fontname="dejavu", fontfile=FONT, fontsize=12)
    source = tmp_path / "tablo.pdf"
    document.save(source)
    result = to_xlsx(
        XlsxParams(path=str(source), output=str(tmp_path / "tablo.xlsx")), silent_progress()
    )
    assert result.table_count == 1
    rows = [list(row) for row in load_workbook(result.output).active.iter_rows(values_only=True)]
    assert rows == [["Özellik", "Avantaj"], ["Hiç bitmeyen", "Dinamik"]]


def test_xlsx_finds_the_same_tables_after_a_markdown_conversion(
    paged: Path, tmp_path: Path
) -> None:
    to_markdown(MarkdownParams(path=str(paged), output=str(tmp_path / "out.md")), silent_progress())
    test_xlsx_cells_come_out_whole(tmp_path)


def _unruled_table(tmp_path):
    import pymupdf as _pymupdf

    document = _pymupdf.open()
    page = document.new_page(width=595, height=842)
    rows = [
        ("Ürün", "Adet", "Fiyat"),
        ("Kalem", "12", "4,50"),
        ("Defter", "3", "18,00"),
        ("Silgi", "7", "2,25"),
    ]
    for index, row in enumerate(rows):
        for column, value in enumerate(row):
            page.insert_text(
                (72 + column * 150, 120 + index * 22), value, fontsize=11, fontname="helv"
            )
    path = tmp_path / "çizgisiz.pdf"
    document.save(path)
    document.close()
    return path


def test_borderless_tables_are_found_only_when_asked(tmp_path):
    from openpyxl import load_workbook

    from vivepdf.ops.convert_tables import XlsxParams, to_xlsx
    from vivepdf.rpc.progress import silent_progress

    source = _unruled_table(tmp_path)
    plain = to_xlsx(
        XlsxParams(path=str(source), output=str(tmp_path / "plain.xlsx")), silent_progress()
    )
    assert plain.table_count == 0
    found = to_xlsx(
        XlsxParams(path=str(source), output=str(tmp_path / "found.xlsx"), borderless=True),
        silent_progress(),
    )
    assert found.table_count == 1
    sheet = load_workbook(found.output).active
    values = [[cell for cell in row if cell] for row in sheet.iter_rows(values_only=True)]
    assert ["Ürün", "Adet", "Fiyat"] in values
    assert ["Defter", 3, 18.0] in values
    assert sheet.cell(3, 3).number_format == "0.00"
