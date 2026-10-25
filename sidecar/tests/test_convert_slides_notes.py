from pathlib import Path

import pymupdf
from pptx import Presentation

from vivepdf.ops.convert_slides import PptxParams, to_pptx
from vivepdf.rpc.progress import silent_progress

NOTE = pymupdf.Rect(60, 100, 300, 160)
YELLOW = (1, 0.9, 0.4)


def _slide(document: pymupdf.Document, tmp_path: Path):
    source = tmp_path / "kaynak.pdf"
    document.save(source)
    document.close()
    result = to_pptx(
        PptxParams(path=str(source), output=str(tmp_path / "sunum.pptx")), silent_progress()
    )
    return Presentation(result.output).slides[0]


def _texts(slide) -> list[str]:
    return [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]


def _colour(slide, point: tuple[float, float], width: float) -> tuple[int, int, int]:
    (picture,) = [shape for shape in slide.shapes if shape.shape_type == 13]
    background = pymupdf.Pixmap(picture.image.blob)
    scale = background.width / width
    return background.pixel(round(point[0] * scale), round(point[1] * scale))


def _yellowish(colour: tuple[int, int, int]) -> bool:
    red, green, blue = colour
    return red > 230 and green > 210 and blue < 140


def _noted(rotation: int = 0) -> pymupdf.Document:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_text((40, 60), "Sayfa yazisi", fontsize=14)
    page.set_rotation(rotation)
    note = page.add_freetext_annot(
        NOTE * page.derotation_matrix if rotation else NOTE,
        "Not kutusu yazisi",
        fontsize=14,
        text_color=(0, 0, 0.6),
        fill_color=YELLOW,
        border_width=2,
        rotate=rotation,
    )
    note.update()
    return document


def test_a_text_note_keeps_its_fill_behind_its_text(tmp_path: Path) -> None:
    slide = _slide(_noted(), tmp_path)

    assert sorted(_texts(slide)) == ["Not kutusu yazisi", "Sayfa yazisi"]
    assert _yellowish(_colour(slide, (250, 145), 400))


def test_a_text_note_on_a_turned_page_keeps_its_fill(tmp_path: Path) -> None:
    document = _noted(rotation=90)
    page = document[0]
    shown = pymupdf.Rect(page.first_annot.rect) * page.rotation_matrix
    del page

    slide = _slide(document, tmp_path)

    assert "Not kutusu yazisi" in _texts(slide)
    assert _yellowish(_colour(slide, (shown.x1 - 6, shown.y1 - 6), 300))


def test_a_highlight_stays_behind_the_slide_text(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_text((40, 100), "Vurgulu satir", fontsize=20)
    area = page.search_for("Vurgulu")[0]
    page.add_highlight_annot(area)

    slide = _slide(document, tmp_path)

    assert _texts(slide) == ["Vurgulu satir"]
    assert _yellowish(_colour(slide, (area.x0 + 2, area.y0 + 2), 400))
