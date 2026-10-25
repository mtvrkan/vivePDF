from pathlib import Path

import pymupdf
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from vivepdf.ops.convert_slides import PptxParams, to_pptx
from vivepdf.rpc.progress import silent_progress

WHITE = (255, 255, 255)
BAR = (0.2, 0.4, 0.8)
AXIS = (0.3, 0.3, 0.3)
BARS = [pymupdf.Rect(60, 120, 100, 240), pymupdf.Rect(130, 80, 170, 240)]


def _presentation(document: pymupdf.Document, tmp_path: Path):
    source = tmp_path / "kaynak.pdf"
    document.save(source)
    document.close()
    result = to_pptx(
        PptxParams(path=str(source), output=str(tmp_path / "sunum.pptx")), silent_progress()
    )
    return Presentation(result.output)


def _chart(rotation: int = 0) -> pymupdf.Document:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    for bar in BARS:
        page.draw_rect(bar, color=None, fill=BAR, width=0)
    page.draw_line((50, 240), (200, 240), color=AXIS, width=1)
    page.set_rotation(rotation)
    return document


def _kinds(slide) -> list[MSO_SHAPE_TYPE]:
    return [shape.shape_type for shape in slide.shapes]


def _rectangles(slide) -> list:
    return [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]


def _picture(slide):
    (picture,) = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    return picture


def _pixel(slide, point: tuple[float, float], width: float = 400) -> tuple[int, ...]:
    background = pymupdf.Pixmap(_picture(slide).image.blob)
    scale = background.width / width
    return background.pixel(round(point[0] * scale), round(point[1] * scale))


def _points(slide, shape) -> tuple[float, float, float, float]:
    picture = _picture(slide)
    scale = picture.width / 400
    return tuple(
        round(value / scale, 1)
        for value in (
            shape.left - picture.left,
            shape.top - picture.top,
            shape.left - picture.left + shape.width,
            shape.top - picture.top + shape.height,
        )
    )


def _fill(shape) -> str:
    return shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr")).get("val")


def test_bars_and_an_axis_become_editable_shapes_and_leave_the_background(
    tmp_path: Path,
) -> None:
    slide = _presentation(_chart(), tmp_path).slides[0]

    rectangles = _rectangles(slide)
    lines = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.LINE]
    assert len(rectangles) == 2 and len(lines) == 1
    assert [_points(slide, shape) for shape in rectangles] == [tuple(bar) for bar in BARS]
    assert {_fill(shape) for shape in rectangles} == {"3366CC"}
    assert _points(slide, lines[0]) == (50, 240, 200, 240)
    assert _pixel(slide, (80, 180)) == WHITE
    assert _pixel(slide, (150, 240)) == WHITE
    assert _pixel(slide, (190, 240)) == WHITE


def test_shapes_sit_above_the_background_and_below_the_text(tmp_path: Path) -> None:
    document = _chart()
    document[0].insert_text((60, 60), "Satışlar", fontsize=14, fontname="helv")

    kinds = _kinds(_presentation(document, tmp_path).slides[0])

    assert kinds[0] == MSO_SHAPE_TYPE.PICTURE
    assert kinds[-1] == MSO_SHAPE_TYPE.TEXT_BOX
    assert kinds.index(MSO_SHAPE_TYPE.TEXT_BOX) > kinds.index(MSO_SHAPE_TYPE.LINE)


@pytest.mark.parametrize("rotation", [90, 270])
def test_shapes_on_a_turned_page_land_where_they_are_seen(tmp_path: Path, rotation: int) -> None:
    document = _chart(rotation)
    page = document[0]
    seen = [bar * page.rotation_matrix for bar in BARS]

    slide = _presentation(document, tmp_path).slides[0]

    rectangles = _rectangles(slide)
    picture = _picture(slide)
    scale = picture.width / 300
    for shape, bar in zip(rectangles, seen, strict=True):
        assert shape.left - picture.left == pytest.approx(bar.x0 * scale, abs=2)
        assert shape.top - picture.top == pytest.approx(bar.y0 * scale, abs=2)
        assert shape.width == pytest.approx(bar.width * scale, abs=2)
    for bar in seen:
        assert _pixel(slide, ((bar.x0 + bar.x1) / 2, (bar.y0 + bar.y1) / 2), 300) == WHITE


def test_a_bar_painted_over_by_a_later_drawing_stays_in_the_background(tmp_path: Path) -> None:
    document = _chart()
    document[0].draw_circle((80, 130), 15, color=None, fill=(1, 0, 0))

    slide = _presentation(document, tmp_path).slides[0]

    rectangles = _rectangles(slide)
    assert len(rectangles) == 1
    assert _pixel(slide, (80, 200)) != WHITE


def test_curves_dashes_and_white_fills_stay_in_the_background(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.draw_rect((0, 0, 400, 300), color=None, fill=(1, 1, 1))
    page.draw_circle((100, 100), 30, color=None, fill=BAR)
    page.draw_line((50, 200), (350, 200), color=AXIS, width=2, dashes="[6] 0")

    slide = _presentation(document, tmp_path).slides[0]

    assert _kinds(slide) == [MSO_SHAPE_TYPE.PICTURE]
    assert _pixel(slide, (100, 100)) != WHITE


def test_a_see_through_fill_keeps_its_opacity(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.draw_rect((50, 50, 150, 150), color=None, fill=BAR, fill_opacity=0.4)

    slide = _presentation(document, tmp_path).slides[0]

    (rectangle,) = _rectangles(slide)
    colour = rectangle._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    assert colour.find(qn("a:alpha")).get("val") == "40000"


def test_a_box_drawn_over_text_stays_in_the_background_with_the_text_under_it(
    tmp_path: Path,
) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_text((60, 100), "Gizli Kişi", fontsize=20, fontname="helv")
    page.draw_rect((55, 80, 170, 106), color=None, fill=(0, 0, 0))
    page.draw_rect((220, 80, 300, 106), color=None, fill=BAR)

    slide = _presentation(document, tmp_path).slides[0]

    assert len(_rectangles(slide)) == 1
    assert _pixel(slide, (110, 93)) == (0, 0, 0)


def test_a_label_written_on_a_bar_keeps_the_bar_editable(tmp_path: Path) -> None:
    document = _chart()
    document[0].insert_text((66, 200), "42", fontsize=14, fontname="helv")

    slide = _presentation(document, tmp_path).slides[0]

    assert len(_rectangles(slide)) == 2
    assert [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame][-1] == "42"
