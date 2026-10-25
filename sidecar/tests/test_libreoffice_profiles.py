import shutil
import sys
import tempfile
import threading
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from vivepdf.external import libreoffice
from vivepdf.rpc.errors import OpError


@pytest.fixture
def short_data_dir(monkeypatch: pytest.MonkeyPatch):
    directory = Path(tempfile.mkdtemp(prefix="vp"))
    monkeypatch.setenv("VIVEPDF_DATA_DIR", str(directory))
    yield directory
    shutil.rmtree(directory, ignore_errors=True)


def test_concurrent_slots_get_their_own_profile(short_data_dir: Path) -> None:
    with libreoffice.profile_slot() as first, libreoffice.profile_slot() as second:
        assert first != second
        assert first.parent == second.parent == short_data_dir / "office-profiles"
    with libreoffice.profile_slot() as again:
        assert again == first


def test_overflow_slots_are_removed_afterwards(short_data_dir: Path) -> None:
    with (
        libreoffice.profile_slot() as first,
        libreoffice.profile_slot() as second,
        libreoffice.profile_slot() as third,
    ):
        for slot in (first, second, third):
            (slot / "user").mkdir(parents=True)
    root = short_data_dir / "office-profiles"
    assert first.is_dir() and second.is_dir()
    assert not third.exists()
    assert not (root / f"{third.name}.lock").exists()


def test_arguments_keep_every_write_inside_the_profile(short_data_dir: Path) -> None:
    with libreoffice.profile_slot() as slot:
        arguments = libreoffice.profile_arguments(slot)
    assert arguments[0] == f"-env:UserInstallation={slot.resolve().as_uri()}"
    assert arguments[1].startswith("-env:UNO_SHARED_PACKAGES_CACHE=")
    assert arguments[1].endswith((slot / "shared").resolve().as_uri())
    environment = libreoffice.conversion_environment()
    assert environment["SAL_DISABLE_OPENCL"] == "1"
    assert not any(key.upper().startswith("PYTHON") for key in environment)


def _letter(folder: Path, number: int) -> Path:
    document = Document()
    document.add_paragraph(f"Mektup {number}: çalışma ağı şöyle")
    path = folder / f"mektup {number}.docx"
    document.save(path)
    return path


def test_three_conversions_run_side_by_side(short_data_dir: Path, tmp_path: Path) -> None:
    if libreoffice.find_soffice() is None:
        pytest.skip("LibreOffice is not available")
    sources = [_letter(tmp_path, number) for number in range(3)]
    produced: dict[int, Path] = {}
    failures: list[BaseException] = []

    def convert(number: int) -> None:
        try:
            produced[number] = libreoffice.convert_to_pdf(
                sources[number], tmp_path / f"out{number}"
            )
        except BaseException as error:  # noqa: BLE001
            failures.append(error)

    threads = [threading.Thread(target=convert, args=(number,)) for number in range(3)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert not failures
    for number in range(3):
        with pymupdf.open(produced[number]) as document:
            assert f"Mektup {number}" in document[0].get_text()
    profiles = short_data_dir / "office-profiles"
    assert sorted(item.name for item in profiles.iterdir() if item.is_dir()) == ["p0", "p1"]


@pytest.mark.skipif(sys.platform != "win32", reason="the path limit is a Windows rule")
def test_a_deep_data_folder_falls_back_to_a_short_profile_root(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    deep = tmp_path / ("uzun-klasör-" * 12)
    monkeypatch.setenv("VIVEPDF_DATA_DIR", str(deep))
    root = libreoffice.profiles_root()
    assert len(str(root)) <= libreoffice.PROFILE_ROOT_LIMIT
    assert root.name == "vivepdf-office"


@pytest.mark.skipif(sys.platform != "win32", reason="the path limit is a Windows rule")
def test_no_short_root_is_a_clear_error(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("VIVEPDF_DATA_DIR", str(tmp_path / ("a" * 120)))
    monkeypatch.setattr(libreoffice.tempfile, "gettempdir", lambda: str(tmp_path / ("b" * 120)))
    with pytest.raises(OpError) as refused:
        libreoffice.profiles_root()
    assert refused.value.data["reason"] == "officeProfilePath"


def test_the_font_configuration_points_at_mac_font_folders_and_the_profile(tmp_path: Path) -> None:
    home = tmp_path / "Ali & Ayşe"

    text = libreoffice.font_configuration(tmp_path / "p0", home)

    assert "<dir>/System/Library/Fonts</dir>" in text
    assert f"<dir>{home / 'Library' / 'Fonts'}</dir>".replace("&", "&amp;") in text
    assert f"<cachedir>{(tmp_path / 'p0' / 'fonts' / 'cache').resolve()}</cachedir>" in text
    assert "<family>Aptos</family><prefer><family>Carlito</family>" in text
    assert "<family>sans-serif</family><prefer><family>Helvetica</family>" in text


def test_only_macos_gets_its_own_font_configuration(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(libreoffice.sys, "platform", "win32")

    environment = libreoffice.font_environment(tmp_path)

    assert environment == {}
    assert not (tmp_path / "fonts").exists()


def test_an_empty_office_file_is_unreadable_without_starting_libreoffice(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "bos.docx"
    source.write_bytes(b"")
    monkeypatch.setattr(libreoffice, "find_soffice", lambda: tmp_path / "soffice")

    def no_run(*_args: object, **_kwargs: object) -> None:
        raise AssertionError("LibreOffice must not start for an empty file")

    monkeypatch.setattr(libreoffice, "run_tool", no_run)

    with pytest.raises(OpError) as refused:
        libreoffice.convert_to_pdf(source, tmp_path / "out")

    assert refused.value.data["reason"] == "officeUnreadable"


@pytest.mark.skipif(sys.platform != "darwin", reason="macOS font folders")
def test_mac_system_fonts_and_missing_sans_fonts_keep_their_look(
    short_data_dir: Path, tmp_path: Path
) -> None:
    if libreoffice.find_soffice() is None:
        pytest.skip("LibreOffice is not available")
    document = Document()
    for name in ("Tahoma", "Geist", "Aptos"):
        document.add_paragraph().add_run(f"{name} çalışma").font.name = name
    source = tmp_path / "fontlar.docx"
    document.save(source)

    produced = libreoffice.convert_to_pdf(source, tmp_path / "out")

    with pymupdf.open(produced) as converted:
        fonts = {
            span["text"].split()[0]: span["font"]
            for block in converted[0].get_text("dict")["blocks"]
            for line in block.get("lines", [])
            for span in line["spans"]
        }
    assert fonts == {"Tahoma": "Tahoma", "Geist": "Helvetica", "Aptos": "Carlito"}
