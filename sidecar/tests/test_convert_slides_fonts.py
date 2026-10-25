from pathlib import Path

import pymupdf
from pptx import Presentation

from vivepdf.ops._pptx_fonts import slide_font_names
from vivepdf.ops.convert_slides import PptxParams, to_pptx
from vivepdf.rpc.progress import silent_progress

SIDECAR = Path(__file__).resolve().parent.parent
FONT = str(SIDECAR / "vivepdf" / "assets" / "fonts" / "DejaVuSans.ttf")
BASE_FONTS = {"helv": "Helvetica", "tiro": "Times-Roman", "cour": "Courier"}


def _renamed(names: dict[str, str]) -> pymupdf.Document:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    for number, (builtin, text) in enumerate(names.items()):
        page.insert_text((40, 60 + 40 * number), text, fontsize=14, fontname=builtin)
    for xref, _ext, _kind, base_font, *_rest in page.get_fonts():
        builtin = next(key for key, value in BASE_FONTS.items() if value == base_font)
        document.xref_set_key(xref, "BaseFont", f"/{names[builtin]}")
    return document


def _run_fonts(document: pymupdf.Document, tmp_path: Path) -> dict[str, str]:
    source = tmp_path / "kaynak.pdf"
    document.save(source)
    document.close()
    result = to_pptx(
        PptxParams(path=str(source), output=str(tmp_path / "sunum.pptx")), silent_progress()
    )
    return {
        run.text: run.font.name
        for shape in Presentation(result.output).slides[0].shapes
        if shape.has_text_frame
        for run in shape.text_frame.paragraphs[0].runs
    }


def test_run_together_font_names_get_their_spaces(tmp_path: Path) -> None:
    document = _renamed({"helv": "TrebuchetMS-Bold", "tiro": "SegoeUI", "cour": "CenturyGothic"})

    fonts = _run_fonts(document, tmp_path)

    assert fonts == {
        "TrebuchetMS-Bold": "Trebuchet MS",
        "SegoeUI": "Segoe UI",
        "CenturyGothic": "Century Gothic",
    }


def test_an_embedded_font_uses_the_family_name_inside_it(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_text((40, 60), "Gömülü", fontsize=14, fontname="dejavu", fontfile=FONT)

    assert _run_fonts(document, tmp_path) == {"Gömülü": "DejaVu Sans"}


def test_unknown_and_missing_names_fall_back_safely() -> None:
    document = _renamed({"helv": "OpenSymbol", "tiro": "ArialMT", "cour": "Courier"})

    names = slide_font_names(document)

    assert names("OpenSymbol") == "OpenSymbol"
    assert names("ArialMT") == "Arial"
    assert names("TimesNewRomanPS-BoldItal") == "Times New Roman"
    assert names("") == "Arial"
