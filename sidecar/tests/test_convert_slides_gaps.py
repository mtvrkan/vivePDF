from pathlib import Path

import pymupdf
from pptx import Presentation

from vivepdf.ops._line_gaps import wide_gap_starts
from vivepdf.ops.convert_slides import PptxParams, to_pptx
from vivepdf.rpc.progress import silent_progress

SIDECAR = Path(__file__).resolve().parent.parent
FONT = str(SIDECAR / "vivepdf" / "assets" / "fonts" / "DejaVuSans.ttf")
EMU_PER_POINT = 12700
SIZE = 10.0


def _line(words: list[tuple[str, float]], direction: tuple[float, float] = (1.0, 0.0)) -> dict:
    chars = []
    for word, start in words:
        if chars:
            end = chars[-1]["bbox"][2]
            chars.append({"c": " ", "bbox": (end, 0, end + 0.25 * SIZE, SIZE)})
        for offset, letter in enumerate(word):
            left = start + offset * 0.5 * SIZE
            chars.append({"c": letter, "bbox": (left, 0, left + 0.5 * SIZE, SIZE)})
    return {"dir": direction, "spans": [{"size": SIZE, "chars": chars}]}


def _starts(words: list[tuple[str, float]]) -> list[int]:
    positions, index = [], 0
    for number, (word, _start) in enumerate(words):
        if number:
            index += 1
            positions.append(index)
        index += len(word)
    return positions


def _spaced(words: list[str], gap: float) -> list[tuple[str, float]]:
    placed, left = [], 0.0
    for word in words:
        placed.append((word, left))
        left += len(word) * 0.5 * SIZE + gap * SIZE
    return placed


def test_a_wide_gap_after_a_bullet_starts_a_new_box() -> None:
    words = [("•", 0.0), ("Madde", 0.5 * SIZE + 0.7 * SIZE), ("metni", 0.0)]
    words[2] = ("metni", words[1][1] + 5 * 0.5 * SIZE + 0.3 * SIZE)

    starts = wide_gap_starts(_line(words))

    assert starts == {_starts(words)[0]}


def test_an_even_justified_line_stays_in_one_box() -> None:
    words = _spaced(["Bu", "satır", "iki", "yana", "yaslanmış", "durumda"], 0.85)

    assert wide_gap_starts(_line(words)) == set()


def test_a_short_first_word_of_a_justified_line_is_not_a_marker() -> None:
    words = _spaced(["Bu", "kuramı", "savunanlara", "göre"], 0.85)

    assert wide_gap_starts(_line(words)) == set()


def test_a_column_gap_inside_a_row_starts_a_new_box() -> None:
    words = [("Node", 0.0), ("Her", 4 * 0.5 * SIZE + 2.0 * SIZE)]
    words.append(("eleman", words[1][1] + 3 * 0.5 * SIZE + 0.3 * SIZE))

    assert wide_gap_starts(_line(words)) == {_starts(words)[0]}


def test_extra_spaces_typed_in_the_text_need_no_new_box() -> None:
    line = _line(_spaced(["sahiptir.", "Şekil"], 0.3))
    first = line["spans"][0]["chars"]
    end = first[8]["bbox"][2]
    gap = [{"c": " ", "bbox": (end + n * 2.5, 0, end + (n + 1) * 2.5, SIZE)} for n in range(3)]
    tail = [
        {"c": char["c"], "bbox": (char["bbox"][0] + 7.5, 0, char["bbox"][2] + 7.5, SIZE)}
        for char in first[10:]
    ]
    line["spans"][0]["chars"] = first[:9] + gap + tail

    assert wide_gap_starts(line) == set()


def test_a_turned_line_is_left_whole() -> None:
    words = [("•", 0.0), ("Madde", 2.0 * SIZE)]

    assert wide_gap_starts(_line(words, direction=(0.0, -1.0))) == set()


def test_slide_keeps_the_bullet_indent_and_column_start(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    font = pymupdf.Font(fontfile=FONT)
    column = 50 + font.text_length("Node", fontsize=14) + 0.75 * 14
    writer = pymupdf.TextWriter(page.rect)
    for point, text in (((50, 80), "•"), ((69, 80), "Madde metni")):
        writer.append(point, text, font=font, fontsize=14)
    for point, text in (((50, 140), "Node"), ((column, 140), "Her eleman")):
        writer.append(point, text, font=font, fontsize=14)
    writer.write_text(page)
    source = tmp_path / "kaynak.pdf"
    document.save(source)

    result = to_pptx(
        PptxParams(path=str(source), output=str(tmp_path / "sunum.pptx")), silent_progress()
    )

    presentation = Presentation(result.output)
    ratio = presentation.slide_width / (400 * EMU_PER_POINT)
    lefts = {
        shape.text_frame.text.strip(): shape.left / EMU_PER_POINT / ratio
        for shape in presentation.slides[0].shapes
        if shape.has_text_frame
    }
    assert abs(lefts["Madde metni"] - 69) < 1.5
    assert abs(lefts["Her eleman"] - column) < 1.5
