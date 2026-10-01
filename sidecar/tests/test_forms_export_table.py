import csv
import datetime as dt
from pathlib import Path

import pymupdf
import pytest
from openpyxl import load_workbook

from vivepdf.ops.forms_export import ExportParams, export_forms
from vivepdf.rpc.errors import OpError
from vivepdf.rpc.progress import silent_progress


def _radio_kid(xref: int, state: str, x: float, on: bool) -> str:
    return (
        f"{xref} 0 obj << /Type /Annot /Subtype /Widget /Parent 5 0 R "
        f"/Rect [{x} 700 {x + 14} 714] /F 4 /P 3 0 R /MK << /CA (l) >> "
        f"/AS /{state if on else 'Off'} /AP << /N << /{state} 20 0 R /Off 21 0 R >> >> >> endobj\n"
    )


def _text(xref: int, name: str, value: str, y: int) -> str:
    return (
        f"{xref} 0 obj << /Type /Annot /Subtype /Widget /FT /Tx /T ({name}) "
        f"/Rect [72 {y} 250 {y + 20}] /P 3 0 R /F 4 /V ({value}) /DA (/Helv 10 Tf 0 g) >> endobj\n"
    )


def _form(path: Path, note: str = "hello", amount: str = "42") -> Path:
    objects = [
        "1 0 obj << /Type /Catalog /Pages 2 0 R /AcroForm << "
        "/Fields [5 0 R 9 0 R 10 0 R 11 0 R 12 0 R 13 0 R 14 0 R 15 0 R] "
        "/DA (/Helv 0 Tf 0 g) /DR << /Font << /Helv 30 0 R >> >> >> >> endobj\n",
        "2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n",
        "3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        "/Annots [6 0 R 7 0 R 8 0 R 9 0 R 10 0 R 11 0 R 12 0 R 13 0 R 14 0 R 15 0 R] >> endobj\n",
        "5 0 obj << /FT /Btn /Ff 49152 /T (size) /V /L /Opt [(Small) (Medium) (Large)] "
        "/Kids [6 0 R 7 0 R 8 0 R] >> endobj\n",
        _radio_kid(6, "S", 72, False),
        _radio_kid(7, "M", 100, False),
        _radio_kid(8, "L", 128, True),
        "9 0 obj << /Type /Annot /Subtype /Widget /FT /Ch /Ff 131072 /T (country) "
        "/Rect [72 650 250 670] /P 3 0 R /F 4 /Opt [[(tr) (Turkiye)] [(de) (Almanya)]] "
        "/V (tr) /DA (/Helv 10 Tf 0 g) >> endobj\n",
        _text(10, "note", note, 600),
        _text(11, "=cmd", "x", 560),
        _text(12, "file", "inside", 520),
        _text(13, "amount", amount, 480),
        _text(14, "when", "05.09.2026", 440),
        "15 0 obj << /Type /Annot /Subtype /Widget /FT /Btn /Ff 65536 /T (send) "
        "/Rect [72 400 150 420] /P 3 0 R /F 4 >> endobj\n",
        "20 0 obj << /Type /XObject /Subtype /Form /BBox [0 0 14 14] /Length 0 >> stream\n\n"
        "endstream endobj\n",
        "21 0 obj << /Type /XObject /Subtype /Form /BBox [0 0 14 14] /Length 0 >> stream\n\n"
        "endstream endobj\n",
        "30 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
    ]
    raw = "%PDF-1.7\n" + "".join(objects) + "trailer << /Root 1 0 R /Size 31 >>\n%%EOF\n"
    document = pymupdf.open("pdf", raw.encode("latin-1"))
    document.save(path)
    document.close()
    return path


@pytest.fixture
def forms(tmp_path: Path) -> list[str]:
    return [
        str(_form(tmp_path / "a.pdf")),
        str(_form(tmp_path / "b.pdf", note="bell\\007ring", amount="7")),
    ]


def _export(paths: list[str], output: Path, **extra):
    return export_forms(ExportParams(paths=paths, output=str(output), **extra), silent_progress())


def test_spreadsheet_keeps_formula_like_headers_as_text_and_types_values(
    forms: list[str], tmp_path: Path
) -> None:
    result = _export(forms, tmp_path / "table.xlsx")
    assert result.files == 2 and result.failed == []
    sheet = load_workbook(result.output).active
    header = [cell.value for cell in sheet[1]]
    assert header[0] == "file"
    formula = sheet.cell(row=1, column=header.index("=cmd") + 1)
    assert formula.data_type == "s" and formula.quotePrefix
    first = {name: cell.value for name, cell in zip(header, sheet[2], strict=True)}
    assert sheet.cell(row=2, column=1).value == "a.pdf"
    assert first["file"] == "inside"
    assert first["amount"] == 42
    assert first["when"] == dt.datetime(2026, 9, 5)
    assert first["size"] == "Large"
    assert first["country"] == "Turkiye"
    assert header.count("file") == 2
    assert "send" not in header
    second = {name: cell.value for name, cell in zip(header, sheet[3], strict=True)}
    assert second["note"] == "bellring"


def test_csv_uses_the_chosen_separator_and_quotes_formula_headers(
    forms: list[str], tmp_path: Path
) -> None:
    result = _export(forms, tmp_path / "table.csv", format="csv", delimiter=",")
    with Path(result.output).open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    assert rows[0][0] == "file"
    assert "'=cmd" in rows[0]
    assert rows[1][0] == "a.pdf"


def test_unreadable_file_is_listed_and_the_rest_exported(forms: list[str], tmp_path: Path) -> None:
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.7\nnot really a pdf")
    result = _export([*forms, str(broken)], tmp_path / "table.xlsx")
    assert result.files == 2
    assert result.failed == [str(broken)]


def test_export_fails_when_no_file_can_be_read(tmp_path: Path) -> None:
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"nothing")
    with pytest.raises(OpError) as caught:
        _export([str(broken)], tmp_path / "table.xlsx")
    assert caught.value.data["reason"] == "noReadableFiles"
    assert not (tmp_path / "table.xlsx").exists()
