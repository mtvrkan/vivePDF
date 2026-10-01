import csv
from pathlib import Path

import numpy as np
import pymupdf
import pytest
from PIL import Image

from vivepdf.ops.omr import (
    OmrExportParams,
    OmrReadParams,
    OmrReviewParams,
    OmrSheetParams,
    classify,
    export,
    read,
    review,
    sheet,
)
from vivepdf.ops.omr_layout import SheetSpec, option_letters, parse_payload, sheet_layout
from vivepdf.rpc.errors import OpError
from vivepdf.rpc.progress import silent_progress

PENCIL = (0.28, 0.28, 0.28)
LIGHT = (0.7, 0.7, 0.7)
ANSWERS = {0: 1, 1: 0, 2: 4, 3: 3, 5: 2, 9: 0, 18: 1, 19: 3}


def _sheet(tmp_path: Path, name: str = "sheet.pdf", **extra) -> Path:
    params = {
        "output": str(tmp_path / name),
        "questions": 20,
        "options": 5,
        "idDigits": 4,
        "booklets": 2,
        "title": "Fizik sınavı",
        "labels": {"name": "Ad Soyad", "studentId": "Öğrenci no", "booklet": "Kitapçık"},
        **extra,
    }
    return Path(sheet(OmrSheetParams.model_validate(params), silent_progress()).output)


def _fill(page: pymupdf.Page, point: tuple[float, float], color=PENCIL, share: float = 0.9) -> None:
    page.draw_circle(point, 5.0 * share, color=None, fill=color)


def _filled(tmp_path: Path, extra_marks=None) -> Path:
    source = _sheet(tmp_path)
    layout = sheet_layout(SheetSpec(questions=20, options=5, digits=4, booklets=2))
    document = pymupdf.open(source)
    page = document[0]
    for column, digit in enumerate([2, 0, 2, 6]):
        _fill(page, layout.digit_bubbles[column][digit])
    _fill(page, layout.booklet_bubbles[1])
    for question, option in ANSWERS.items():
        _fill(page, layout.question_bubbles[question][option])
    for question, option, color in extra_marks or []:
        _fill(page, layout.question_bubbles[question][option], color)
    target = tmp_path / "filled.pdf"
    document.save(target)
    document.close()
    return target


def _photo(pdf: Path, target: Path, angle: float = 0.0, dpi: int = 150, noise: float = 0.0) -> Path:
    with pymupdf.open(pdf) as document:
        pixmap = document[0].get_pixmap(dpi=dpi)
    image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
    if angle:
        image = image.rotate(
            angle, expand=True, fillcolor=(245, 245, 240), resample=Image.Resampling.BICUBIC
        )
    if noise:
        values = np.asarray(image, dtype=float)
        values += np.random.default_rng(7).normal(0, noise, values.shape)
        image = Image.fromarray(np.clip(values, 0, 255).astype(np.uint8))
    image.save(target)
    return target


def _read(*paths: Path, threshold: float | None = None):
    params: dict = {"paths": [str(path) for path in paths]}
    if threshold is not None:
        params["threshold"] = threshold
    return read(OmrReadParams.model_validate(params), silent_progress())


def _answers(scan) -> dict[int, int]:
    return {
        index: mark.chosen[0] for index, mark in enumerate(scan.marks) if mark.state == "single"
    }


def test_the_sheet_code_describes_only_the_layout_and_round_trips() -> None:
    spec = SheetSpec(
        questions=40, options=4, digits=8, booklets=3, letter_case="lower", paper="letter"
    )
    assert parse_payload(spec.payload()) == spec
    assert parse_payload("VPOMR1;q=0;o=5") is None
    assert parse_payload("something else") is None
    assert option_letters(3, "lower") == ["a", "b", "c"]


def test_too_many_questions_for_one_sheet_are_refused_with_the_capacity() -> None:
    with pytest.raises(OpError) as caught:
        sheet_layout(SheetSpec(questions=200, options=8, digits=12, booklets=4))
    assert caught.value.data["reason"] == "omrTooManyQuestions"
    assert 0 < caught.value.data["capacity"] < 200


def test_a_generated_sheet_has_real_text_and_the_requested_copies(tmp_path: Path) -> None:
    path = _sheet(tmp_path, copies=3)
    with pymupdf.open(path) as document:
        assert document.page_count == 3
        text = document[2].get_text()
        for word in ("Fizik sınavı", "Ad Soyad", "Öğrenci no", "Kitapçık", "20"):
            assert word in text


def test_labels_in_scripts_without_latin_glyphs_are_printed_and_still_read(tmp_path: Path) -> None:
    labels = {
        "name": "氏名",
        "studentId": "受験番号",
        "booklet": "الكراسة",
        "hint": "濃い鉛筆で塗りつぶしてください",
    }
    path = _sheet(tmp_path, title="数学 第1回", labels=labels)
    with pymupdf.open(path) as document:
        text = document[0].get_text()
    for word in ("数学", "氏名", "受験番号", "塗りつぶしてください"):
        assert word in text
    result = _read(path)
    assert [scan.questions for scan in result.sheets] == [20]


def test_a_tilted_noisy_photo_of_a_filled_sheet_is_read_exactly(tmp_path: Path) -> None:
    photo = _photo(_filled(tmp_path), tmp_path / "tilted.png", angle=2.5, dpi=150, noise=8)
    result = _read(photo)
    assert result.failures == []
    [scan] = result.sheets
    assert scan.page is None
    assert scan.student_id == "2026"
    assert scan.booklet is not None and scan.booklet.chosen == [1]
    assert _answers(scan) == ANSWERS
    assert all(
        mark.state == "blank" for index, mark in enumerate(scan.marks) if index not in ANSWERS
    )


def test_an_upside_down_scan_inside_a_pdf_is_read_with_its_page_number(tmp_path: Path) -> None:
    upside_down = _photo(_filled(tmp_path), tmp_path / "turned.png", angle=180, dpi=200)
    scanned = tmp_path / "scans.pdf"
    document = pymupdf.open()
    document.new_page(width=595, height=842).insert_text((72, 72), "cover page without a sheet")
    page = document.new_page(width=595, height=842)
    page.insert_image(page.rect, filename=str(upside_down))
    document.save(scanned)
    document.close()
    result = _read(scanned)
    assert [(failure.page, failure.reason) for failure in result.failures] == [(1, "noSheetCode")]
    [scan] = result.sheets
    assert scan.page == 2
    assert scan.student_id == "2026"
    assert _answers(scan) == ANSWERS


def test_double_and_faint_marks_are_flagged_for_review(tmp_path: Path) -> None:
    filled = _filled(tmp_path, extra_marks=[(0, 3, PENCIL), (5, 4, LIGHT), (7, 1, LIGHT)])
    [scan] = _read(_photo(filled, tmp_path / "marks.png", dpi=150)).sheets
    assert scan.marks[0].state == "multiple"
    assert sorted(scan.marks[0].chosen) == [1, 3]
    assert scan.marks[5].state == "unclear"
    assert scan.marks[5].chosen[0] == 2
    assert scan.marks[7].state == "unclear"
    assert scan.marks[7].chosen == [1]


def test_pages_without_a_readable_sheet_are_reported_not_raised(tmp_path: Path) -> None:
    filled = _photo(_filled(tmp_path), tmp_path / "whole.png", dpi=150)
    with Image.open(filled) as image:
        image.crop((0, 0, image.width, image.height // 2)).save(tmp_path / "half.png")
    Image.new("RGB", (800, 1100), "white").save(tmp_path / "blank.png")
    (tmp_path / "broken.png").write_bytes(b"not an image")
    result = _read(tmp_path / "half.png", tmp_path / "blank.png", tmp_path / "broken.png")
    assert result.sheets == []
    assert [failure.reason for failure in result.failures] == [
        "cornersNotFound",
        "noSheetCode",
        "unreadable",
    ]


def test_classify_separates_single_multiple_unclear_and_blank() -> None:
    assert classify([0.05, 0.8, 0.04], 0.4).state == "single"
    assert classify([0.05, 0.8, 0.7], 0.4).state == "multiple"
    assert classify([0.05, 0.8, 0.3], 0.4).model_dump() == {
        "state": "unclear",
        "chosen": [1, 2],
        "fills": [0.05, 0.8, 0.3],
    }
    assert classify([0.02, 0.1, 0.05], 0.4).state == "blank"


def test_the_review_copy_circles_the_marks_on_each_scan(tmp_path: Path) -> None:
    photo = _photo(_filled(tmp_path), tmp_path / "checked.png", angle=-1.5, dpi=150)
    [scan] = _read(photo).sheets
    questions = [
        {
            "chosen": mark.chosen,
            "key": [1],
            "verdict": "correct"
            if mark.chosen == [1]
            else ("blank" if not mark.chosen else "wrong"),
        }
        for mark in scan.marks
    ]
    params = {
        "output": str(tmp_path / "review.pdf"),
        "items": [
            {
                "source": scan.source,
                "page": scan.page,
                "layout": scan.layout,
                "transform": scan.transform,
                "pixelScale": scan.pixel_scale,
                "header": "2026 · 3/20",
                "questions": questions,
            }
        ],
    }
    result = review(OmrReviewParams.model_validate(params), silent_progress())
    with pymupdf.open(result.output) as document:
        page = document[0]
        assert "2026 · 3/20" in page.get_text()
        rings = [
            drawing
            for drawing in page.get_drawings()
            if drawing.get("color") and drawing.get("fill") is None
        ]
        assert len(rings) >= 20


def test_results_export_to_csv_and_to_a_workbook_with_one_sheet_per_table(tmp_path: Path) -> None:
    tables = [
        {
            "title": "Students",
            "header": ["Number", "Net"],
            "rows": [["=2026", 17.25], ["0042", 12.0]],
        },
        {"title": "Questions", "header": ["Question", "Correct %"], "rows": [[1.0, 50.0]]},
    ]
    csv_result = export(
        OmrExportParams.model_validate(
            {"output": str(tmp_path / "out.csv"), "format": "csv", "tables": tables}
        ),
        silent_progress(),
    )
    with open(csv_result.output, encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle, delimiter=";"))
    assert rows[0] == ["Number", "Net"]
    assert rows[1][0] != "=2026" and rows[1][1] == "17.25"
    assert rows[2] == ["0042", "12"]

    from openpyxl import load_workbook

    xlsx_result = export(
        OmrExportParams.model_validate({"output": str(tmp_path / "out.xlsx"), "tables": tables}),
        silent_progress(),
    )
    workbook = load_workbook(xlsx_result.output)
    assert workbook.sheetnames == ["Students", "Questions"]
    assert workbook["Students"]["B2"].value == 17.25
    assert workbook["Students"]["A3"].value == "0042"
    assert workbook["Questions"]["A2"].value == 1
