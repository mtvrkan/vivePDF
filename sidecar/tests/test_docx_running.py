import tempfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn

from vivepdf.ops._docx_running import find_running_text
from vivepdf.ops.convert import DocxParams, to_docx
from vivepdf.rpc.progress import silent_progress

WIDTH, HEIGHT = 595, 842


def _page(
    document: pymupdf.Document,
    header: str | None,
    footer: str | None,
    body: str,
    footer_right: bool = True,
) -> None:
    page = document.new_page(width=WIDTH, height=HEIGHT)
    if header:
        width = pymupdf.get_text_length(header, fontsize=10)
        page.insert_text(((WIDTH - width) / 2, 40), header, fontsize=10)
    if footer:
        width = pymupdf.get_text_length(footer, fontsize=9)
        x = WIDTH - 72 - width if footer_right else 72
        page.insert_text((x, 810), footer, fontsize=9)
    page.insert_text((72, 100), body, fontsize=16)
    for row in range(6):
        page.insert_text((72, 140 + row * 18), f"Ordinary body line {row} for {body}.", fontsize=11)


def _report(tmp_path: Path, pages: int = 5, first: int = 1, name: str = "report.pdf") -> Path:
    document = pymupdf.open()
    for number in range(first, first + pages):
        page_count = first + pages - 1
        _page(document, "ACME Annual Report", f"Page {number} of {page_count}", f"Part {number}")
        document[-1].insert_text((72, 810), "Confidential", fontsize=9)
    target = tmp_path / name
    document.save(target)
    return target


def _fields(part) -> list[str]:
    return [node.get(qn("w:instr")) for node in part._element.iter(qn("w:fldSimple"))]


def _body(word: Document) -> str:
    return "\n".join(paragraph.text for paragraph in word.paragraphs)


def test_repeated_header_and_footer_go_into_word(tmp_path: Path) -> None:
    source = _report(tmp_path)
    result = to_docx(
        DocxParams(path=str(source), output=str(tmp_path / "out.docx")), silent_progress()
    )
    assert (result.header_lines, result.footer_lines) == (1, 2)
    word = Document(result.output)
    header = word.sections[0].header
    footer = word.sections[0].footer
    assert [paragraph.text for paragraph in header.paragraphs] == ["ACME Annual Report"]
    assert "Confidential" in footer.paragraphs[0].text
    assert "Page " in footer.paragraphs[0].text and " of 5" in footer.paragraphs[0].text
    assert _fields(footer) == ["PAGE"]
    body = _body(word)
    assert "ACME" not in body and "Confidential" not in body and "Page 3 of 5" not in body
    assert "Part 3" in body and "Ordinary body line 5 for Part 5." in body
    assert all(section.header.is_linked_to_previous for section in word.sections[1:])


def test_header_and_footer_can_stay_in_the_body(tmp_path: Path) -> None:
    source = _report(tmp_path)
    result = to_docx(
        DocxParams(path=str(source), output=str(tmp_path / "plain.docx"), header_footer=False),
        silent_progress(),
    )
    assert (result.header_lines, result.footer_lines) == (0, 0)
    word = Document(result.output)
    assert "ACME Annual Report" in _body(word)
    assert _fields(word.sections[0].footer) == []


def test_numbering_starts_where_the_pdf_starts(tmp_path: Path) -> None:
    source = _report(tmp_path, first=7)
    result = to_docx(
        DocxParams(path=str(source), output=str(tmp_path / "late.docx")), silent_progress()
    )
    numbering = Document(result.output).sections[0]._sectPr.find(qn("w:pgNumType"))
    assert numbering is not None and numbering.get(qn("w:start")) == "7"


def test_encrypted_input_keeps_working(tmp_path: Path) -> None:
    source = _report(tmp_path)
    locked = tmp_path / "locked.pdf"
    with pymupdf.open(source) as document:
        document.save(
            locked, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="sahip", user_pw="gizli"
        )
    result = to_docx(
        DocxParams(path=str(locked), password="gizli", output=str(tmp_path / "locked.docx")),
        silent_progress(),
    )
    assert result.header_lines == 1
    assert "ACME" not in _body(Document(result.output))


def test_short_documents_and_changing_numbers_stay_in_the_body(tmp_path: Path) -> None:
    document = pymupdf.open()
    for chapter in (2, 2, 5, 5, 5):
        _page(document, f"Chapter {chapter}", None, "Text")
    assert find_running_text(document, list(range(5))).lines == []
    assert find_running_text(document, [0, 1]).lines == []


def test_page_numbers_that_alternate_sides_still_count(tmp_path: Path) -> None:
    document = pymupdf.open()
    for number in range(1, 7):
        _page(document, None, str(number), "Text", footer_right=number % 2 == 0)
    running = find_running_text(document, list(range(6)))
    assert [(line.band, line.parts) for line in running.lines] == [("footer", [None])]
    assert sorted(running.areas) == list(range(6))


@pytest.mark.parametrize("missing", [0, 2])
def test_a_page_without_the_header_keeps_its_own_text(tmp_path: Path, missing: int) -> None:
    document = pymupdf.open()
    for number in range(5):
        _page(document, None if number == missing else "Running Title", None, f"Body {number}")
    running = find_running_text(document, list(range(5)))
    assert len(running.lines) == 1
    assert missing not in running.areas


def test_the_temporary_copy_is_removed(tmp_path: Path, monkeypatch) -> None:
    source = _report(tmp_path)
    made: list[Path] = []
    original = tempfile.mkdtemp

    def tracked(prefix: str) -> str:
        folder = original(prefix=prefix, dir=tmp_path)
        made.append(Path(folder))
        return folder

    monkeypatch.setattr(tempfile, "mkdtemp", tracked)
    to_docx(DocxParams(path=str(source), output=str(tmp_path / "t.docx")), silent_progress())
    assert len(made) == 1 and not made[0].exists()


def test_a_line_that_is_the_whole_page_is_not_a_header(tmp_path: Path) -> None:
    document = pymupdf.open()
    for number in range(1, 5):
        document.new_page(width=WIDTH, height=HEIGHT).insert_text((72, 60), f"Slide {number}")
    assert find_running_text(document, list(range(4))).lines == []


def test_two_running_lines_sharing_a_corner_both_reach_the_header(tmp_path: Path) -> None:
    document = pymupdf.open()
    for number in range(1, 5):
        page = document.new_page(width=WIDTH, height=HEIGHT)
        page.insert_text((40, 40), "Draft", fontsize=10)
        page.insert_text((90, 40), "Internal", fontsize=10)
        page.insert_text((WIDTH - 120, 40), "Rev 7", fontsize=10)
        page.insert_text((72, 100), f"Part {number}", fontsize=16)
        for row in range(6):
            page.insert_text(
                (72, 140 + row * 18), f"Body line {row} of part {number}.", fontsize=11
            )
    source = tmp_path / "corners.pdf"
    document.save(source)

    result = to_docx(
        DocxParams(path=str(source), output=str(tmp_path / "corners.docx")), silent_progress()
    )

    header = Document(result.output).sections[0].header
    assert [paragraph.text for paragraph in header.paragraphs] == ["Draft Internal\t\tRev 7"]
