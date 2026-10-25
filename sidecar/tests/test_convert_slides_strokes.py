from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.oxml.ns import qn

from vivepdf.ops.convert_slides import PptxParams, to_pptx
from vivepdf.rpc.progress import silent_progress

SIDECAR = Path(__file__).resolve().parent.parent
FONT = str(SIDECAR / "vivepdf" / "assets" / "fonts" / "DejaVuSans.ttf")
EMU_PER_POINT = 12700
SIZE = 20


def _write(page: pymupdf.Page, point: tuple[float, float], text: str, **style) -> None:
    page.insert_text(point, text, fontsize=SIZE, fontname="dejavu", fontfile=FONT, **style)


def _presentation(document: pymupdf.Document, tmp_path: Path):
    source = tmp_path / "kaynak.pdf"
    document.save(source)
    document.close()
    result = to_pptx(
        PptxParams(path=str(source), output=str(tmp_path / "sunum.pptx")), silent_progress()
    )
    return Presentation(result.output)


def _runs(presentation) -> dict:
    return {
        run.text.strip(): run
        for shape in presentation.slides[0].shapes
        if shape.has_text_frame
        for run in shape.text_frame.paragraphs[0].runs
        if run.text.strip()
    }


def test_fake_bold_text_comes_out_bold(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    _write(page, (50, 80), "Kalın başlık", render_mode=2, border_width=0.03)

    run = _runs(_presentation(document, tmp_path))["Kalın başlık"]

    assert run.font.bold is True
    assert run._r.find(qn("a:rPr")).find(qn("a:ln")) is None


def test_outline_only_text_stays_hollow(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    _write(page, (50, 80), "Kontur", render_mode=1, color=(1, 0, 0), border_width=0.05)

    presentation = _presentation(document, tmp_path)

    properties = _runs(presentation)["Kontur"]._r.find(qn("a:rPr"))
    line = properties.find(qn("a:ln"))
    ratio = presentation.slide_width / (400 * EMU_PER_POINT)
    assert line.find(qn("a:solidFill")).find(qn("a:srgbClr")).get("val") == "FF0000"
    assert abs(int(line.get("w")) - SIZE * 0.05 * EMU_PER_POINT * ratio) <= 1
    assert properties.find(qn("a:noFill")) is not None
    assert properties.find(qn("a:solidFill")) is None
    assert _runs(presentation)["Kontur"].font.bold is False


def test_only_the_fake_bold_words_of_a_line_turn_bold(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    label = "Örnek:"
    _write(page, (50, 80), label, render_mode=2, border_width=0.03)
    width = pymupdf.Font(fontfile=FONT).text_length(label, fontsize=SIZE)
    _write(page, (50 + width, 80), " elmalar")

    runs = _runs(_presentation(document, tmp_path))

    assert runs[label].font.bold is True
    assert runs["elmalar"].font.bold is False


def test_plain_filled_text_keeps_its_fill(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    _write(page, (50, 80), "Düz yazı", color=(0, 0, 1))

    run = _runs(_presentation(document, tmp_path))["Düz yazı"]

    properties = run._r.find(qn("a:rPr"))
    assert run.font.bold is False
    assert properties.find(qn("a:ln")) is None
    assert properties.find(qn("a:solidFill")).find(qn("a:srgbClr")).get("val") == "0000FF"
