import re
import threading
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from vivepdf.ops import _docx_parallel as parallel
from vivepdf.ops.convert import DocxParams, to_docx
from vivepdf.rpc.errors import ErrorCode, OpError
from vivepdf.rpc.progress import Progress

FONT = str(
    Path(__file__).resolve().parent.parent / "vivepdf" / "assets" / "fonts" / "DejaVuSans.ttf"
)
PAGES = 10


@pytest.fixture
def long_pdf(tmp_path: Path) -> Path:
    document = pymupdf.open()
    for number in range(PAGES):
        page = document.new_page(width=595, height=842)
        page.insert_text(
            (72, 100), f"Bölüm {number + 1} gövde metni", fontname="dejavu", fontfile=FONT
        )
        for x in (72, 272, 472):
            page.draw_line((x, 200), (x, 300))
        for y in (200, 250, 300):
            page.draw_line((72, y), (472, y))
        page.insert_text((80, 230), f"hücre {number}", fontname="dejavu", fontfile=FONT)
        page.draw_circle((300, 500), 40, color=(0, 0, 1))
    path = tmp_path / "uzun belge ş.pdf"
    document.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="gizli", owner_pw="sahip")
    document.close()
    return path


@pytest.fixture
def in_workers(monkeypatch):
    monkeypatch.setattr(parallel, "PARALLEL_MIN_PAGES", 2)
    monkeypatch.setattr(parallel, "PAGES_PER_WORKER", 1)
    monkeypatch.setattr(parallel, "MAX_WORKERS", 2)


def _recording(cancel_after: int | None = None):
    events: list[tuple[float, str | None, dict | None]] = []
    cancel = threading.Event()

    def sink(value, message, detail):
        events.append((value, message, detail))
        pages = [event for event in events if event[1] == "progress.convertingPages"]
        if cancel_after is not None and len(pages) >= cancel_after:
            cancel.set()

    return Progress(sink, cancel), events


def _body(path: str) -> str:
    with zipfile.ZipFile(path) as archive:
        return re.sub(r'r:(id|embed)="rId\d+"', "", archive.read("word/document.xml").decode())


def test_worker_count_stays_sequential_for_short_documents() -> None:
    assert parallel.worker_count(parallel.PARALLEL_MIN_PAGES - 1, cpus=16) == 1
    assert parallel.worker_count(278, cpus=1) == 1
    assert parallel.worker_count(278, cpus=4) == 3
    assert parallel.worker_count(278, cpus=64) == parallel.MAX_WORKERS
    assert parallel.worker_count(parallel.PARALLEL_MIN_PAGES, cpus=64) == 2


def test_batches_keep_every_page_once_and_in_order() -> None:
    indices = list(range(3, 40))
    chunks = parallel.batches(indices, 3)
    assert [index for chunk in chunks for index in chunk] == indices
    assert all(chunks)


def test_workers_produce_the_same_document(
    long_pdf: Path, tmp_path: Path, in_workers, monkeypatch
) -> None:
    progress, events = _recording()
    shared = to_docx(
        DocxParams(path=str(long_pdf), password="gizli", output=str(tmp_path / "par.docx")),
        progress,
    )
    page_events = [event for event in events if event[1] == "progress.convertingPages"]
    assert len(page_events) == 2 * PAGES
    values = [event[0] for event in events]
    assert values == sorted(values)
    monkeypatch.setattr(parallel, "PARALLEL_MIN_PAGES", 10**6)
    single = to_docx(
        DocxParams(path=str(long_pdf), password="gizli", output=str(tmp_path / "seq.docx")),
        _recording()[0],
    )
    assert _body(shared.output) == _body(single.output)
    text = "\n".join(paragraph.text for paragraph in Document(shared.output).paragraphs)
    assert "Bölüm 1 gövde metni" in text
    assert f"Bölüm {PAGES} gövde metni" in text


def test_workers_honour_the_page_range(long_pdf: Path, tmp_path: Path, in_workers) -> None:
    progress, events = _recording()
    result = to_docx(
        DocxParams(
            path=str(long_pdf), password="gizli", output=str(tmp_path / "r.docx"), pages="3-6"
        ),
        progress,
    )
    totals = {event[2]["total"] for event in events if event[1] == "progress.convertingPages"}
    assert totals == {4}
    text = "\n".join(paragraph.text for paragraph in Document(result.output).paragraphs)
    assert "Bölüm 3 gövde" in text
    assert "Bölüm 6 gövde" in text
    assert "Bölüm 7 gövde" not in text


def test_cancel_stops_the_workers_and_writes_nothing(
    long_pdf: Path, tmp_path: Path, in_workers
) -> None:
    progress, _events = _recording(cancel_after=2)
    target = tmp_path / "iptal.docx"
    with pytest.raises(OpError) as caught:
        to_docx(DocxParams(path=str(long_pdf), password="gizli", output=str(target)), progress)
    assert caught.value.code == ErrorCode.CANCELLED
    assert not target.exists()


def test_a_worker_failure_surfaces_as_an_error(tmp_path: Path, in_workers, monkeypatch) -> None:
    source = tmp_path / "kısa.pdf"
    document = pymupdf.open()
    for _ in range(4):
        document.new_page()
    document.save(source)
    document.close()
    monkeypatch.setattr(parallel, "parse_batch", parallel._announce_page)
    with pytest.raises(OpError) as caught:
        to_docx(DocxParams(path=str(source), output=str(tmp_path / "x.docx")), _recording()[0])
    assert caught.value.code == ErrorCode.INTERNAL
