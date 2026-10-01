import io
import queue
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from openpyxl import load_workbook
from pdf2docx.page.Page import Page

from vivepdf.ops import _docx_parallel as parallel
from vivepdf.ops import convert_slides, convert_tables
from vivepdf.ops._cell_text import table_spans
from vivepdf.ops._docx_bidi import mark_right_to_left
from vivepdf.ops._docx_progress import clip_ratio
from vivepdf.ops._pptx_layout import font_family
from vivepdf.ops.convert import DocxParams, to_docx
from vivepdf.ops.convert_slides import PptxParams, slide_dpi, to_pptx
from vivepdf.ops.convert_tables import XlsxParams, sheet_title, to_xlsx
from vivepdf.rpc.errors import ErrorCode, OpError
from vivepdf.rpc.progress import silent_progress

FONT = str(
    Path(__file__).resolve().parent.parent / "vivepdf" / "assets" / "fonts" / "DejaVuSans.ttf"
)


def _grid_page(document: pymupdf.Document, cells: list[tuple[float, float, str]]) -> None:
    page = document.new_page(width=400, height=300)
    page.draw_rect(pymupdf.Rect(40, 40, 340, 160))
    page.draw_line((140, 40), (140, 160))
    page.draw_line((240, 80), (240, 160))
    page.draw_line((40, 80), (340, 80))
    page.draw_line((140, 120), (340, 120))
    for x, y, text in cells:
        page.insert_text((x, y), text, fontname="dejavu", fontfile=FONT, fontsize=10)


MERGED_CELLS = [
    (50, 60, "Kod"),
    (150, 60, "Başlık"),
    (50, 100, "Grup"),
    (150, 100, "1.234,50"),
    (250, 100, "=SUM(A1)"),
    (150, 140, "%12"),
    (250, 140, "31.12.2025"),
]


@pytest.fixture
def merged_pdf(tmp_path: Path) -> Path:
    document = pymupdf.open()
    _grid_page(document, MERGED_CELLS)
    path = tmp_path / "birleşik.pdf"
    document.save(path)
    document.close()
    return path


def test_table_spans_follow_the_ruling(merged_pdf: Path) -> None:
    with pymupdf.open(merged_pdf) as document:
        table = document[0].find_tables(use_layout=False).tables[0]
        assert table_spans(table) == [(0, 1, 0, 2), (1, 0, 2, 0)]


def test_xlsx_keeps_merged_cells_numbers_and_formula_looking_text(
    merged_pdf: Path, tmp_path: Path
) -> None:
    result = to_xlsx(
        XlsxParams(path=str(merged_pdf), output=str(tmp_path / "t.xlsx")), silent_progress()
    )
    sheet = load_workbook(result.output).active
    assert sorted(str(merged) for merged in sheet.merged_cells.ranges) == ["A2:A3", "B1:C1"]
    assert sheet["B1"].value == "Başlık"
    assert sheet["B2"].value == 1234.5
    assert sheet["B2"].number_format == "#,##0.00"
    assert sheet["C2"].value == "=SUM(A1)"
    assert sheet["C2"].data_type == "s"
    assert sheet["C2"].quotePrefix
    assert sheet["C3"].value.date().isoformat() == "2025-12-31"


def test_csv_still_prefixes_formula_looking_text(merged_pdf: Path, tmp_path: Path) -> None:
    result = to_xlsx(
        XlsxParams(path=str(merged_pdf), output=str(tmp_path / "t.csv"), format="csv"),
        silent_progress(),
    )
    raw = Path(result.output).read_text(encoding="utf-8-sig")
    assert "'=SUM(A1)" in raw


def test_pages_without_tables_go_to_a_text_sheet(tmp_path: Path) -> None:
    document = pymupdf.open()
    _grid_page(document, MERGED_CELLS)
    document.new_page().insert_text((72, 72), "Açıklama satırı", fontname="dejavu", fontfile=FONT)
    document.new_page()
    source = tmp_path / "karışık.pdf"
    document.save(source)
    document.close()
    result = to_xlsx(
        XlsxParams(
            path=str(source),
            output=str(tmp_path / "k.xlsx"),
            page_label="Sayfa",
            table_label="Tablo",
            text_label="Metin",
        ),
        silent_progress(),
    )
    workbook = load_workbook(result.output)
    assert workbook.sheetnames == ["Sayfa1-Tablo1", "Metin"]
    assert [list(row) for row in workbook["Metin"].iter_rows(values_only=True)] == [
        [2, "Açıklama satırı"]
    ]
    assert result.text_pages == [2]
    assert result.textless_pages == [3]


def test_sheet_titles_drop_forbidden_characters_and_stay_unique() -> None:
    taken: set[str] = set()
    assert sheet_title("a/b[c]:d*?\\", taken) == "abcd"
    assert sheet_title("ABCD", taken) == "ABCD (2)"
    assert len(sheet_title("x" * 40, taken)) == 31
    assert sheet_title("''", taken) == "1"


def test_a_failed_write_leaves_no_file_behind(
    merged_pdf: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def broken(*_args, **_kwargs):
        raise OSError("disk full")

    monkeypatch.setattr(convert_tables, "_write_workbook", broken)
    target = tmp_path / "out" / "t.xlsx"
    with pytest.raises(OSError):
        to_xlsx(XlsxParams(path=str(merged_pdf), output=str(target)), silent_progress())
    assert list(target.parent.iterdir()) == []


def test_output_names_keep_dotted_stems(merged_pdf: Path, tmp_path: Path) -> None:
    dotted = to_xlsx(
        XlsxParams(path=str(merged_pdf), output=str(tmp_path / "rapor.v2")), silent_progress()
    )
    assert Path(dotted.output).name == "rapor.v2.xlsx"
    swapped = to_xlsx(
        XlsxParams(path=str(merged_pdf), output=str(tmp_path / "tablo.xlsx"), format="csv"),
        silent_progress(),
    )
    assert Path(swapped.output).name == "tablo.csv"


@pytest.fixture
def scanned_pdf(tmp_path: Path) -> Path:
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    picture = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 80), False)
    picture.clear_with(230)
    page.insert_image(page.rect, pixmap=picture)
    page.insert_text(
        (72, 120), "Taranmış sayfa metni", fontname="dejavu", fontfile=FONT, render_mode=3
    )
    document.new_page(width=595, height=842)
    path = tmp_path / "taranmış.pdf"
    document.save(path)
    document.close()
    return path


def test_docx_keeps_the_recognised_text_of_a_scanned_page(
    scanned_pdf: Path, tmp_path: Path
) -> None:
    result = to_docx(
        DocxParams(path=str(scanned_pdf), output=str(tmp_path / "t.docx")), silent_progress()
    )
    text = "\n".join(paragraph.text for paragraph in Document(result.output).paragraphs)
    assert "Taranmış sayfa metni" in text
    assert result.textless_pages == [2]


def test_docx_reports_pages_it_could_not_write(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    document = pymupdf.open()
    for number in range(3):
        document.new_page().insert_text(
            (72, 72), f"Sayfa {number + 1}", fontname="dejavu", fontfile=FONT
        )
    source = tmp_path / "üç.pdf"
    document.save(source)
    document.close()
    plain_make = Page.make_docx

    def failing_second(self, doc):
        if self.id == 1:
            raise RuntimeError("broken page")
        return plain_make(self, doc)

    monkeypatch.setattr(Page, "make_docx", failing_second)
    result = to_docx(
        DocxParams(path=str(source), output=str(tmp_path / "o.docx"), pages="1-3,2"),
        silent_progress(),
    )
    assert result.skipped_pages == [2]
    text = "\n".join(paragraph.text for paragraph in Document(result.output).paragraphs)
    assert "Sayfa 1" in text and "Sayfa 3" in text


def _docx_bytes(paragraphs: list[tuple[str, WD_ALIGN_PARAGRAPH | None]]) -> bytes:
    document = Document()
    for text, alignment in paragraphs:
        paragraph = document.add_paragraph(text)
        paragraph.alignment = alignment
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_right_to_left_paragraphs_are_marked_and_mirrored() -> None:
    data = _docx_bytes(
        [("مرحبا بالعالم", WD_ALIGN_PARAGRAPH.RIGHT), ("Hello world", WD_ALIGN_PARAGRAPH.RIGHT)]
    )
    paragraphs = Document(io.BytesIO(mark_right_to_left(data))).paragraphs
    arabic, english = paragraphs[0], paragraphs[1]
    assert arabic._p.pPr.find(qn("w:bidi")) is not None
    assert arabic.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert all(run.font.rtl for run in arabic.runs)
    assert english._p.pPr.find(qn("w:bidi")) is None
    assert english.alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_left_to_right_documents_are_returned_untouched() -> None:
    data = _docx_bytes([("Merhaba dünya", None)])
    assert mark_right_to_left(data) is data


class _Job:
    def __init__(self, ready: bool):
        self._ready = ready

    def ready(self) -> bool:
        return self._ready

    def successful(self) -> bool:
        return True


class _Worker:
    def __init__(self, exitcode: int | None):
        self.exitcode = exitcode


def _parser() -> parallel.ParallelParse:
    return parallel.ParallelParse("x.pdf", None, lambda: None, lambda: None)


def test_a_worker_that_dies_mid_job_fails_the_conversion() -> None:
    with pytest.raises(OpError) as caught:
        _parser()._wait([_Job(False)], queue.Queue(), [_Worker(None), _Worker(3221225477)])
    assert caught.value.code == ErrorCode.INTERNAL
    assert caught.value.data == {"reason": "workerCrashed"}


def test_workers_leaving_after_the_last_job_are_not_a_crash() -> None:
    _parser()._wait([_Job(True)], queue.Queue(), [_Worker(0), _Worker(0)])


def test_image_clipping_resolution_shrinks_only_for_huge_pages() -> None:
    assert clip_ratio(595 * 842) == 4.0
    huge = clip_ratio(5000 * 5000)
    assert 1.0 <= huge < 4.0
    assert clip_ratio(0) == 4.0


def test_slide_pictures_stay_within_the_pixel_budget() -> None:
    assert slide_dpi(pymupdf.Rect(0, 0, 595, 842), 300) == 300
    capped = slide_dpi(pymupdf.Rect(0, 0, 14400, 14400), 300)
    assert (14400 * capped / 72) ** 2 <= convert_slides.MAX_SLIDE_PIXELS


def test_pptx_image_notes_drop_characters_a_slide_cannot_hold(tmp_path: Path) -> None:
    from pptx import Presentation

    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "not\x01lar", fontname="helv")
    source = tmp_path / "notlar.pdf"
    document.save(source)
    document.close()
    result = to_pptx(
        PptxParams(path=str(source), output=str(tmp_path / "n.pptx"), dpi=50, mode="image"),
        silent_progress(),
    )
    slide = Presentation(result.output).slides[0]
    assert slide.notes_slide.notes_text_frame.text == "notlar"


def test_font_family_keeps_multi_word_names() -> None:
    assert font_family("MS-Mincho") == "MS Mincho"
    assert font_family("Segoe-UI-Semibold") == "Segoe UI"
    assert font_family("TimesNewRomanPS-BoldItalicMT") == "Times New Roman"
    assert font_family("Arial,Bold") == "Arial"
