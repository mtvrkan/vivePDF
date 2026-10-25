from pathlib import Path
from types import SimpleNamespace

import pymupdf
from openpyxl import load_workbook

from vivepdf.ops._cell_text import page_fragments, table_texts
from vivepdf.ops._table_finding import is_page_frame
from vivepdf.ops.convert_tables import XlsxParams, to_xlsx
from vivepdf.rpc.progress import silent_progress

LINE = (0.6, 0.6, 0.6)
HEADER_FILL = (0.9, 0.93, 0.98)
FONT = str(
    Path(__file__).resolve().parent.parent / "vivepdf" / "assets" / "fonts" / "DejaVuSans.ttf"
)
TEXT = {"fontname": "dejavu", "fontfile": FONT}


def _rows(path: str, title: str) -> list[list[object]]:
    sheet = load_workbook(path)[title]
    return [
        [value for value in row if value is not None] for row in sheet.iter_rows(values_only=True)
    ]


def _grid_page(path: Path) -> Path:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 80), "Teklif özeti", fontsize=14, **TEXT)
    for row in range(4):
        page.draw_line((72, 120 + row * 24), (372, 120 + row * 24), color=LINE)
    for column in range(3):
        page.draw_line((72 + column * 150, 120), (72 + column * 150, 192), color=LINE)
    for row, (left, right) in enumerate((("Kalem", "Tutar"), ("Tasarım", "48"), ("Eğitim", "6"))):
        page.insert_text((78, 136 + row * 24), left, fontsize=11, **TEXT)
        page.insert_text((228, 136 + row * 24), right, fontsize=11, **TEXT)
    page.insert_text((72, 240), "Ödeme teslimden sonra yapılır.", fontsize=11, **TEXT)
    document.save(path)
    document.close()
    return path


def test_text_around_a_table_goes_to_the_text_sheet(tmp_path: Path) -> None:
    source = _grid_page(tmp_path / "teklif.pdf")

    result = to_xlsx(
        XlsxParams(path=str(source), output=str(tmp_path / "o.xlsx")), silent_progress()
    )

    assert result.table_count == 1
    assert result.text_pages == []
    assert _rows(result.output, "Text") == [
        [1, "Teklif özeti"],
        [1, "Ödeme teslimden sonra yapılır."],
    ]
    assert ["Tasarım", 48] in _rows(result.output, "S1-T1")


def test_a_filled_header_keeps_the_rows_under_its_horizontal_rules(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page()
    for column, name in enumerate(("Kalem", "Adet", "Tutar")):
        cell = pymupdf.Rect(72 + column * 150, 100, 222 + column * 150, 126)
        page.draw_rect(cell, fill=HEADER_FILL, color=None)
        page.insert_text((cell.x0 + 6, 117), name, fontsize=11, **TEXT)
    body = (("Arayüz", "1", "48.000,00 TL"), ("Eğitim", "2", "6.000,00 TL"))
    for row, values in enumerate(body):
        bottom = 152 + row * 26
        page.draw_line((72, bottom), (522, bottom), color=LINE)
        for column, value in enumerate(values):
            page.insert_text((78 + column * 150, bottom - 9), value, fontsize=11, **TEXT)
    source = tmp_path / "sozlesme.pdf"
    document.save(source)
    document.close()

    result = to_xlsx(
        XlsxParams(path=str(source), output=str(tmp_path / "o.xlsx")), silent_progress()
    )

    rows = _rows(result.output, "S1-T1")
    assert rows[:3] == [["Kalem", "Adet", "Tutar"], ["Arayüz", 1, 48000], ["Eğitim", 2, 6000]]
    sheet = load_workbook(result.output)["S1-T1"]
    assert sheet.cell(2, 3).number_format == '#,##0.00 "TL"'


def test_a_lone_filled_row_without_rules_below_stays_one_row(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page()
    for column, name in enumerate(("Ad", "Soyad")):
        cell = pymupdf.Rect(72 + column * 150, 100, 222 + column * 150, 126)
        page.draw_rect(cell, fill=HEADER_FILL, color=LINE)
        page.insert_text((cell.x0 + 6, 117), name, fontsize=11, **TEXT)
    page.draw_line((72, 400), (372, 400), color=LINE)
    source = tmp_path / "tek.pdf"
    document.save(source)
    document.close()

    result = to_xlsx(
        XlsxParams(path=str(source), output=str(tmp_path / "o.xlsx")), silent_progress()
    )

    assert _rows(result.output, "S1-T1") == [["Ad", "Soyad"]]


def _table(cells: list[tuple[float, float, float, float]]) -> SimpleNamespace:
    return SimpleNamespace(cells=cells)


def test_a_slide_frame_with_overlapping_cells_is_not_a_table() -> None:
    page = SimpleNamespace(rect=pymupdf.Rect(0, 0, 720, 540))
    frame = _table([(0, 0, 61, 20), (0, 20, 720, 520), (61, 20, 656, 240), (0, 520, 720, 540)])
    grid = _table([(72, 100, 222, 126), (222, 100, 372, 126), (72, 126, 222, 152)])

    assert is_page_frame(page, frame) is True
    assert is_page_frame(page, grid) is False
    assert is_page_frame(page, _table([(10, 10, 700, 500)])) is True


def test_borderless_cells_take_whole_words_and_reach_the_page_edges(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((20, 100), "Not", fontsize=11, **TEXT)
    page.insert_text((72, 100), "Bu sözleşme", fontsize=11, **TEXT)
    boundary = 72 + pymupdf.Font(fontfile=FONT).text_length("B", 11) / 2
    row = SimpleNamespace(cells=[(60, 88, boundary, 104), (boundary, 88, 300, 104)])
    table = SimpleNamespace(rows=[row])
    fragments = page_fragments(page)

    ruled = table_texts(table, fragments)
    unruled = table_texts(table, fragments, (page.rect.x0, page.rect.x1))

    assert ruled == [["B", "u sözleşme"]]
    assert unruled == [["Not", "Bu sözleşme"]]
