import base64
import io
import zipfile
from pathlib import Path

import docx
import pytest
from lxml import etree
from PIL import Image

from vivepdf.ops._document_docx import document_docx
from vivepdf.ops._studio_document_models import StudioDocumentDocxParams
from vivepdf.ops.studio_document import export_docx
from vivepdf.rpc.errors import OpError
from vivepdf.rpc.progress import silent_progress

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"
EQUATION = (
    '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 2000 1000">'
    '<path d="M100 900 L1900 900 L1000 100 Z" fill="#1a1a1a"/></svg>'
)


def _picture() -> str:
    buffer = io.BytesIO()
    Image.new("RGB", (200, 100), (30, 90, 160)).save(buffer, "PNG")
    return "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode()


def _params(body: str, **extra) -> StudioDocumentDocxParams:
    return StudioDocumentDocxParams.model_validate(
        {"html": body, "output": "/tmp/out.docx", **extra}
    )


def _parts(payload: bytes) -> dict[str, etree._Element]:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        return {
            name: etree.fromstring(archive.read(name))
            for name in archive.namelist()
            if name.endswith(".xml") and name.startswith("word/")
        }


def _texts(root: etree._Element) -> str:
    return "".join(node.text or "" for node in root.iter(f"{W}t", f"{W}instrText"))


def test_blocks_marks_lists_and_tables_become_word_paragraphs() -> None:
    body = (
        '<h1 id="h-1">1 Intro</h1><p style="text-align:justify;text-indent:36pt">Plain '
        '<strong>bold</strong> <em>slanted</em> <span style="color:#aa0000">red</span> '
        '<a href="https://example.org">site</a> <a href="#e-1">(1)</a></p>'
        "<ul><li><p>One</p><ul><li><p>Inner</p></li></ul></li></ul>"
        '<ol start="3"><li><p>Third</p></li></ol><blockquote><p>Quoted</p></blockquote>'
        "<pre><code>a\n  b</code></pre><hr>"
        '<table><tr><th style="width:100pt"><p>A</p></th><th><p>B</p></th></tr>'
        '<tr><td colspan="2"><p>wide</p></td></tr><tr><td rowspan="2"><p>x</p></td>'
        "<td><p>y</p></td></tr><tr><td><p>z</p></td></tr></table>"
    )

    payload = document_docx(_params(body), lambda: None)
    opened = docx.Document(io.BytesIO(payload))
    paragraphs = opened.paragraphs
    document = _parts(payload)["word/document.xml"]

    assert paragraphs[0].style.name == "Heading 1" and paragraphs[0].text == "1 Intro"
    assert paragraphs[1].text == "Plain bold slanted red site (1)"
    assert paragraphs[1].runs[1].bold and paragraphs[1].runs[3].italic
    assert document.find(f".//{W}hyperlink") is not None
    numbered = [node for node in document.iter(f"{W}numPr")]
    assert [node.find(f"{W}ilvl").get(f"{W}val") for node in numbered] == ["0", "1", "0"]
    starts = _parts(payload)["word/numbering.xml"].findall(f".//{W}startOverride")
    assert [node.get(f"{W}val") for node in starts] == ["3"]
    assert any(paragraph.text == "a\n  b" for paragraph in paragraphs)
    table = opened.tables[0]
    assert len(table.rows) == 4 and table.cell(1, 0).text == "wide"
    assert document.find(f".//{W}gridSpan").get(f"{W}val") == "2"
    assert document.find(f".//{W}vMerge").get(f"{W}val") == "restart"
    assert document.find(f".//{W}tblHeader") is not None


def test_cover_front_matter_breaks_and_furniture_make_sections() -> None:
    body = (
        '<h1 id="h-1">Abstract</h1><p>Short.</p><p id="c-1" class="caption" data-caption="figure">'
        '<strong>Figure 1.</strong> Blue</p><div class="page-break" data-main="true"></div>'
        '<h1 id="h-2">Chapter</h1><div class="section-break" data-columns="2" '
        'data-kind="continuous"></div><p>Columns</p>'
    )
    settings = {
        "cover": True,
        "title": "Thesis",
        "author": "Beril Ertuğrul",
        "toc": True,
        "tocTitle": "Contents",
        "figureList": True,
        "pageNumbers": "outside",
        "footer": "Draft",
        "furnitureOnFirst": False,
        "pageNumberStart": 1,
    }

    parts = _parts(document_docx(_params(body, settings=settings), lambda: None))
    document = parts["word/document.xml"]
    sections = list(document.iter(f"{W}sectPr"))
    numbering = [section.find(f"{W}pgNumType") for section in sections]

    assert len(sections) == 4
    assert sections[0].find(f"{W}vAlign").get(f"{W}val") == "center"
    assert numbering[1].get(f"{W}fmt") == "lowerRoman" and numbering[1].get(f"{W}start") == "2"
    assert sections[1].find(f"{W}titlePg") is not None
    assert numbering[2].get(f"{W}fmt") == "decimal" and numbering[2].get(f"{W}start") == "1"
    assert sections[3].find(f"{W}type").get(f"{W}val") == "continuous"
    assert sections[3].find(f"{W}cols").get(f"{W}num") == "2"
    instructions = _texts(document)
    assert 'TOC \\o "1-2" \\h \\z \\u' in instructions
    assert 'TOC \\h \\z \\c "Figure"' in instructions
    assert "SEQ Figure \\* ARABIC" in instructions
    settings_part = parts["word/settings.xml"]
    assert settings_part.find(f"{W}evenAndOddHeaders") is not None
    assert settings_part.find(f"{W}updateFields").get(f"{W}val") == "true"
    modes = [node.get(f"{W}val") for node in settings_part.iter(f"{W}compatSetting")]
    assert "15" in modes and "14" not in modes
    footers = [_texts(root) for name, root in parts.items() if "footer" in name]
    assert any("Draft" in text and "PAGE" in text for text in footers)


def test_notes_comments_pictures_and_equations_are_embedded() -> None:
    body = (
        '<p>Text <span data-comment="c-one">marked</span><sup class="fn-ref" id="fnref-1">1</sup> '
        '<img src="vpeq-0" alt="x^2" style="width:2em;height:1em;vertical-align:-0.2em"></p>'
        '<p style="text-align:center"><img src="vpimg-0" style="width:100pt" alt="Blue"></p>'
        '<p id="e-1" class="eq" data-equation="0" data-latex="x^2">(1)</p>'
    )
    comments = [
        {
            "id": "c-one",
            "author": "Esma",
            "createdAt": "2026-10-26T10:00:00.000Z",
            "text": "Check this",
            "resolved": True,
            "replies": [{"id": "c-two", "author": "Nesibe Şafak", "createdAt": "", "text": "Done"}],
        },
        {"id": "c-gone", "author": "Esma", "text": "Orphan"},
    ]
    params = _params(
        body,
        images=[_picture()],
        equations=[EQUATION],
        footnotes=["Note"],
        footnoteHtml=['Note <a href="https://example.org">link</a>'],
        comments=comments,
    )

    payload = document_docx(params, lambda: None)
    parts = _parts(payload)
    notes = parts["word/footnotes.xml"]
    written = parts["word/comments.xml"].findall(f"{W}comment")
    extended = parts["word/commentsExtended.xml"].findall(f"{W15}commentEx")

    assert "Note link" in _texts(notes)
    assert parts["word/document.xml"].find(f".//{W}footnoteReference").get(f"{W}id") == "1"
    assert [comment.get(f"{W}author") for comment in written] == ["Esma", "Nesibe Şafak"]
    assert written[0].get(f"{W}date") == "2026-10-26T10:00:00Z"
    assert written[1].get(f"{W}date") is None
    assert extended[1].get(f"{W15}paraIdParent") == extended[0].get(f"{W15}paraId")
    assert all(entry.get(f"{W15}done") == "1" for entry in extended)
    assert len(list(parts["word/document.xml"].iter(f"{W}commentRangeStart"))) == 2
    descriptions = [
        node.get("descr")
        for node in parts["word/document.xml"].iter()
        if node.tag.endswith("docPr")
    ]
    assert descriptions == ["x^2", "Blue", "x^2"]
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        assert len(media) == 2
        assert not any("thumbnail" in name for name in archive.namelist())


def test_empty_and_broken_markup_still_make_a_document() -> None:
    for body in ("", "<p></p>", "<p>&nbsp;</p>", "<div><span>loose</span> text</div>"):
        opened = docx.Document(io.BytesIO(document_docx(_params(body), lambda: None)))
        assert opened.paragraphs


def test_comment_ids_and_counts_are_checked() -> None:
    with pytest.raises(ValueError):
        _params("<p>x</p>", comments=[{"id": "bad id", "text": "x"}])
    with pytest.raises(ValueError):
        _params("<p>x</p>", comments=[{"id": "c-a", "text": "x" * 5000}])


def test_export_writes_docx_and_refuses_to_overwrite(tmp_path: Path) -> None:
    target = tmp_path / "talk.docx"
    params = StudioDocumentDocxParams.model_validate(
        {"html": "<h1>Talk</h1><p>Body</p>", "title": "Talk", "output": str(target)}
    )

    result = export_docx(params, silent_progress())

    assert result.output.endswith("talk.docx")
    assert result.bytes == target.stat().st_size
    assert docx.Document(str(target)).core_properties.title == "Talk"
    with pytest.raises(OpError) as caught:
        export_docx(params, silent_progress())
    assert caught.value.data["exists"] is True
