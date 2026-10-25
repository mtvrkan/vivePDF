from pathlib import Path

import pymupdf
from pptx import Presentation

from vivepdf.ops._hidden_text import hidden_origins
from vivepdf.ops._pptx_layout import line_boxes, text_free_pixmap
from vivepdf.ops.convert_slides import PptxParams, to_pptx
from vivepdf.rpc.progress import silent_progress

SIDECAR = Path(__file__).resolve().parent.parent
FONT = str(SIDECAR / "vivepdf" / "assets" / "fonts" / "DejaVuSans.ttf")


def _write(page: pymupdf.Page, point: tuple[float, float], text: str) -> None:
    page.insert_text(point, text, fontsize=14, fontname="dejavu", fontfile=FONT)


def _grey(width: int = 10, height: int = 10) -> pymupdf.Pixmap:
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, width, height), 0)
    pixmap.clear_with(200)
    return pixmap


def _slide_texts(document: pymupdf.Document, tmp_path: Path, mode: str = "editable") -> list[str]:
    source = tmp_path / "kaynak.pdf"
    document.save(source)
    document.close()
    result = to_pptx(
        PptxParams(path=str(source), output=str(tmp_path / "sunum.pptx"), mode=mode),
        silent_progress(),
    )
    slide = Presentation(result.output).slides[0]
    if mode == "image":
        return [slide.notes_slide.notes_text_frame.text]
    return [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]


def _blacked_out() -> pymupdf.Document:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    _write(page, (50, 80), "Ad: Gizli Kişi")
    _write(page, (50, 120), "Açık satır")
    page.draw_rect(page.search_for("Gizli Kişi")[0] + (-1, -1, 1, 1), color=0, fill=0)
    return document


def test_text_under_a_black_box_does_not_reach_the_slide(tmp_path: Path) -> None:
    texts = _slide_texts(_blacked_out(), tmp_path)

    assert texts == ["Ad: ", "Açık satır"]


def test_text_under_a_black_box_stays_out_of_the_notes(tmp_path: Path) -> None:
    (notes,) = _slide_texts(_blacked_out(), tmp_path, mode="image")

    assert "Açık satır" in notes and "Ad:" in notes
    assert "Gizli" not in notes and "Kişi" not in notes


def test_text_under_a_picture_is_left_out_but_text_over_it_stays(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    _write(page, (50, 100), "Resim altı yazı")
    page.insert_image(pymupdf.Rect(40, 80, 250, 110), pixmap=_grey())
    _write(page, (50, 200), "Resim üstü yazı")
    page.insert_image(pymupdf.Rect(40, 180, 250, 210), pixmap=_grey(), overlay=False)

    texts = _slide_texts(document, tmp_path)

    assert texts == ["Resim üstü yazı"]


def test_a_filled_box_note_hides_the_text_below_it(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    _write(page, (50, 80), "Numara 12345")
    note = page.add_rect_annot(page.search_for("12345")[0])
    note.set_colors(stroke=(0, 0, 0), fill=(0, 0, 0))
    note.update()

    texts = _slide_texts(document, tmp_path)

    assert texts == ["Numara "]


def test_text_under_a_see_through_box_stays(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    _write(page, (50, 80), "Vurgulu satır")
    page.draw_rect(page.search_for("Vurgulu")[0], color=None, fill=(1, 1, 0), fill_opacity=0.3)

    texts = _slide_texts(document, tmp_path)

    assert texts == ["Vurgulu satır"]


def test_a_later_picture_cut_short_by_its_clip_keeps_the_text_below(tmp_path: Path) -> None:
    picture = pymupdf.open()
    picture.new_page(width=200, height=200).insert_image(
        pymupdf.Rect(0, 0, 200, 200), pixmap=_grey()
    )
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    _write(page, (50, 200), "Kırpılmış resmin altında")
    page.show_pdf_page(pymupdf.Rect(40, 20, 340, 95), picture, 0, clip=pymupdf.Rect(0, 0, 200, 50))

    texts = _slide_texts(document, tmp_path)

    assert texts == ["Kırpılmış resmin altında"]


def test_hidden_text_is_found_on_a_cropped_turned_page() -> None:
    document = _blacked_out()
    page = document[0]
    page.set_cropbox(pymupdf.Rect(30, 20, 380, 290))
    page.set_rotation(90)

    boxes = line_boxes(page, hidden_origins(page))

    assert ["".join(span["text"] for span in box.spans) for box in boxes] == ["Ad: ", "Açık satır"]


def test_background_loses_only_the_lines_put_on_the_slide_on_a_turned_page() -> None:
    document = _blacked_out()
    page = document[0]
    page.set_cropbox(pymupdf.Rect(30, 20, 380, 290))
    page.set_rotation(90)
    boxes = line_boxes(page, hidden_origins(page))

    background = text_free_pixmap(document, 0, 72, [box.area for box in boxes])

    with pymupdf.open() as rendered:
        rendered.insert_pdf(document)
        expected = rendered[0].get_pixmap(dpi=72)
    assert background.width == expected.width and background.height == expected.height
    text_area = (boxes[1].area * page.rotation_matrix).irect
    assert all(
        background.pixel(x, y) == (255, 255, 255)
        for x in range(text_area.x0, text_area.x1)
        for y in range(text_area.y0, text_area.y1)
    )


def test_a_page_without_covers_has_no_hidden_text() -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    _write(page, (50, 80), "Sade sayfa")

    assert hidden_origins(page) == set()
