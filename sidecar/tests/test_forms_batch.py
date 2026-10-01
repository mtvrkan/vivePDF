from pathlib import Path

import pymupdf
import pytest
from openpyxl import load_workbook

from vivepdf.ops._form_detect_params import DetectParams
from vivepdf.ops._form_rows import normalize_key
from vivepdf.ops._page_text import NUMBER_FONT
from vivepdf.ops.forms import FieldsParams, list_fields
from vivepdf.ops.forms_batch import DataPreviewParams, MergeParams, data_preview, merge
from vivepdf.ops.forms_detect import detect_fields
from vivepdf.ops.forms_export import ExportParams, export_forms
from vivepdf.rpc.errors import ErrorCode, OpError
from vivepdf.rpc.progress import silent_progress


@pytest.fixture
def blank_template(tmp_path: Path) -> Path:
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    font = {"fontname": "dejavu", "fontfile": str(NUMBER_FONT), "fontsize": 11}
    page.insert_text((72, 100), "Ad Soyad:", **font)
    page.draw_line((150, 104), (400, 104), width=0.8)
    page.insert_text((72, 140), "Tarih:", **font)
    page.draw_line((150, 144), (300, 144), width=0.8)
    page.draw_rect(pymupdf.Rect(72, 180, 84, 192), width=0.8)
    page.insert_text((90, 190), "Onaylıyorum", **font)
    page.draw_rect(pymupdf.Rect(72, 220, 400, 250), width=0.8)
    page.insert_text((72, 215), "Notlar", **font)
    path = tmp_path / "template.pdf"
    document.save(path)
    document.close()
    return path


@pytest.fixture
def data_csv(tmp_path: Path) -> Path:
    path = tmp_path / "people.csv"
    path.write_text(
        "Ad Soyad;Tarih;Onaylıyorum;Notlar\n"
        "Ayşe Yılmaz;05.09.2026;evet;Merhaba\n"
        "Mehmet Kaya;06.09.2026;hayır;\n",
        encoding="utf-8-sig",
    )
    return path


def test_normalize_key_ignores_case_and_punctuation() -> None:
    assert normalize_key("Ad Soyad") == normalize_key("ad_soyad") == "adsoyad"


def test_detect_creates_fields_with_labels_in_reading_order(
    blank_template: Path, tmp_path: Path
) -> None:
    result = detect_fields(
        DetectParams(path=str(blank_template), output=str(tmp_path / "form.pdf")), silent_progress()
    )
    names = [field.name for field in result.fields]
    assert names == ["Ad_Soyad", "Tarih", "Onaylıyorum", "Notlar"]
    assert [field.kind for field in result.fields] == ["text", "text", "checkbox", "text"]
    listed = list_fields(FieldsParams(path=result.output), silent_progress())
    assert listed.is_form
    assert {field.name for field in listed.fields} == set(names)


def test_data_preview_reads_csv(data_csv: Path) -> None:
    preview = data_preview(DataPreviewParams(path=str(data_csv)), silent_progress())
    assert preview.columns == ["Ad Soyad", "Tarih", "Onaylıyorum", "Notlar"]
    assert preview.total_rows == 2
    assert preview.rows[0]["Ad Soyad"] == "Ayşe Yılmaz"


def test_merge_fills_one_file_per_row(blank_template: Path, data_csv: Path, tmp_path: Path) -> None:
    form = detect_fields(
        DetectParams(path=str(blank_template), output=str(tmp_path / "form.pdf")), silent_progress()
    )
    result = merge(
        MergeParams(
            path=form.output,
            data_path=str(data_csv),
            output_dir=str(tmp_path / "merged"),
            pattern="{Ad Soyad}",
        ),
        silent_progress(),
    )
    assert result.rows == 2 and result.skipped == 0
    assert result.unmatched_fields == []
    names = sorted(Path(item.output).name for item in result.outputs)
    assert names == ["Ayşe Yılmaz.pdf", "Mehmet Kaya.pdf"]
    first = list_fields(FieldsParams(path=result.outputs[0].output), silent_progress())
    values = {field.name: field.value for field in first.fields}
    assert values["Ad_Soyad"] == "Ayşe Yılmaz"
    assert values["Onaylıyorum"] is True
    second = list_fields(FieldsParams(path=result.outputs[1].output), silent_progress())
    assert {field.name: field.value for field in second.fields}["Onaylıyorum"] is False


def test_export_collects_values_into_xlsx_and_csv(
    blank_template: Path, data_csv: Path, tmp_path: Path
) -> None:
    form = detect_fields(
        DetectParams(path=str(blank_template), output=str(tmp_path / "form.pdf")), silent_progress()
    )
    merged = merge(
        MergeParams(path=form.output, data_path=str(data_csv), output_dir=str(tmp_path / "merged")),
        silent_progress(),
    )
    paths = [item.output for item in merged.outputs]
    xlsx = export_forms(
        ExportParams(paths=paths, output=str(tmp_path / "table.xlsx")), silent_progress()
    )
    assert xlsx.files == 2 and xlsx.failed == []
    sheet = load_workbook(xlsx.output).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:2] == ("file", "Ad_Soyad")
    assert rows[1][1] == "Ayşe Yılmaz"
    csv_result = export_forms(
        ExportParams(paths=paths, output=str(tmp_path / "table"), format="csv"), silent_progress()
    )
    text = Path(csv_result.output).read_text(encoding="utf-8-sig")
    assert text.splitlines()[0].startswith("file;Ad_Soyad")
    assert ";1" in text and ";0" in text


def test_merge_without_form_fields_fails(
    blank_template: Path, data_csv: Path, tmp_path: Path
) -> None:
    with pytest.raises(OpError) as error:
        merge(
            MergeParams(
                path=str(blank_template), data_path=str(data_csv), output_dir=str(tmp_path / "x")
            ),
            silent_progress(),
        )
    assert error.value.code == ErrorCode.INVALID_PARAMS


def _mixed_page(page: pymupdf.Page) -> None:
    font = {"fontname": "dejavu", "fontfile": str(NUMBER_FONT), "fontsize": 10}
    page.insert_text(
        (72, 60), "Başvuru formu ve açıklamalı bir paragraf metni burada yer alır.", **font
    )
    page.draw_line((72, 63), (380, 63), width=0.6)
    page.insert_text((72, 100), "Ad Soyad:", **font)
    page.draw_line((140, 104), (400, 104), width=0.8)
    page.insert_text((72, 140), "Adres", **font)
    page.draw_line((72, 160), (400, 160), width=0.8)
    page.draw_line((300, 220), (480, 220), width=0.8)
    page.insert_text((300, 232), "İmza", **font)
    page.draw_rect(pymupdf.Rect(72, 250, 82, 260), width=0.8)
    page.insert_text((88, 259), "Kabul ediyorum", **font)
    page.draw_rect(pymupdf.Rect(72, 280, 400, 300), color=None, fill=(1, 1, 1))
    page.insert_text((74, 294), "Beyaz zemin üzerinde düz metin", **font)
    page.draw_rect(pymupdf.Rect(72, 310, 400, 340), width=0.8)
    page.insert_text((76, 328), "Kutunun içinde yazı var", **font)
    for row in range(7):
        top = 380 + row * 20
        page.draw_line((72, top), (480, top), width=0.5)
        page.draw_line((72, top), (72, top + 20), width=0.5)
        page.draw_line((250, top), (250, top + 20), width=0.5)
        page.draw_line((480, top), (480, top + 20), width=0.5)
        page.insert_text((76, top + 14), f"Satır {row + 1}", **font)
        page.insert_text((254, top + 14), f"Değer {row + 1}", **font)
    page.draw_line((72, 520), (480, 520), width=0.5)
    for row in range(2):
        top = 560 + row * 24
        for x in (72, 200, 480):
            page.draw_line((x, top), (x, top + 24), width=0.5)
        page.draw_line((72, top), (480, top), width=0.5)
        page.insert_text((76, top + 16), ["Telefon", "E-posta"][row], **font)
    page.draw_line((72, 608), (480, 608), width=0.5)
    for index in range(8):
        page.draw_line((72, 650 + index * 18), (480, 650 + index * 18), width=0.4)


def test_detect_skips_text_decoration_and_filled_tables_but_keeps_real_fields(
    tmp_path: Path,
) -> None:
    document = pymupdf.open()
    _mixed_page(document.new_page(width=595, height=842))
    source = tmp_path / "karışık form.pdf"
    document.save(source)
    document.close()
    result = detect_fields(
        DetectParams(path=str(source), output=str(tmp_path / "out.pdf")), silent_progress()
    )
    found = {(field.kind, round(field.rect[3])) for field in result.fields}
    assert ("text", 105) in found
    assert ("text", 161) in found
    assert ("text", 221) in found
    assert ("checkbox", 260) in found
    assert ("text", 585) in found and ("text", 609) in found
    bottoms = [field.rect[3] for field in result.fields]
    assert not any(60 <= bottom <= 70 for bottom in bottoms)
    assert not any(280 <= bottom <= 345 for bottom in bottoms)
    assert not any(370 <= bottom <= 530 for bottom in bottoms)
    assert not any(640 <= bottom <= 800 for bottom in bottoms)
    assert len(result.fields) == 6
    labels = {round(field.rect[3]): field.label for field in result.fields}
    assert labels[105] == "Ad Soyad"


def test_detect_ignores_rules_drawn_over_a_picture(tmp_path: Path) -> None:
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 100), False)
    pixmap.set_rect(pixmap.irect, (240, 240, 240))
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    page.insert_image(pymupdf.Rect(72, 100, 472, 300), pixmap=pixmap)
    for index in range(5):
        page.draw_line((72, 130 + index * 30), (472, 130 + index * 30), width=0.5)
    page.insert_text((72, 400), "Ad:", fontsize=11)
    page.draw_line((100, 404), (300, 404), width=0.8)
    source = tmp_path / "picture.pdf"
    document.save(source)
    document.close()
    result = detect_fields(
        DetectParams(path=str(source), output=str(tmp_path / "out.pdf")), silent_progress()
    )
    assert [round(field.rect[3]) for field in result.fields] == [405]
