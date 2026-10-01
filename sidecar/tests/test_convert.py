from pathlib import Path

import pymupdf
import pytest

from vivepdf.external import libreoffice
from vivepdf.ops.convert import DocxParams, EmptyParams, TextSourceParams, to_docx, tools_status
from vivepdf.ops.convert_images import ImagesParams, ImagesToPdfParams, images_to_pdf, to_images
from vivepdf.ops.convert_slides import PptxParams, to_pptx
from vivepdf.ops.convert_tables import XlsxParams, to_xlsx
from vivepdf.ops.convert_text import MarkdownParams, TextParams, to_html, to_markdown, to_text
from vivepdf.ops.convert_to_pdf import FileToPdfParams, file_to_pdf
from vivepdf.rpc.errors import ErrorCode, OpError
from vivepdf.rpc.progress import silent_progress


@pytest.fixture
def table_pdf(tmp_path: Path) -> Path:
    document = pymupdf.open()
    page = document.new_page()
    cells = [["Ad", "Yaş"], ["Ayşe", "31"], ["Mehmet", "45"]]
    top = 100
    for row in cells:
        page.draw_rect(pymupdf.Rect(72, top, 300, top + 24), color=(0, 0, 0))
        page.draw_line((186, top), (186, top + 24), color=(0, 0, 0))
        page.insert_text(
            (80, top + 16),
            row[0],
            fontsize=11,
            fontfile="vivepdf/assets/fonts/DejaVuSans.ttf",
            fontname="dv",
        )
        page.insert_text(
            (194, top + 16),
            row[1],
            fontsize=11,
            fontfile="vivepdf/assets/fonts/DejaVuSans.ttf",
            fontname="dv",
        )
        top += 24
    path = tmp_path / "table.pdf"
    document.save(path)
    document.close()
    return path


@pytest.fixture
def image_folder(tmp_path: Path) -> Path:
    folder = tmp_path / "photos"
    folder.mkdir()
    for index, size in ((10, (300, 200)), (2, (200, 300)), (1, (100, 100))):
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, *size), False)
        pixmap.set_rect(pixmap.irect, (index * 20, 100, 150))
        pixmap.save(folder / f"img{index}.png")
    (folder / "notes.txt").write_text("ignore me", encoding="utf-8")
    return folder


def test_to_text_markdown_html(sample_pdf: Path, tmp_path: Path) -> None:
    text = to_text(
        TextParams(path=str(sample_pdf), output=str(tmp_path / "out.txt")), silent_progress()
    )
    content = Path(text.output).read_text(encoding="utf-8")
    assert "Page 1" in content and content.count("\f") == 2

    markdown = to_markdown(
        MarkdownParams(path=str(sample_pdf), output=str(tmp_path / "out.md"), pages="1"),
        silent_progress(),
    )
    assert "Page 1" in Path(markdown.output).read_text(encoding="utf-8")

    html = to_html(
        TextSourceParams(path=str(sample_pdf), output=str(tmp_path / "out.html")), silent_progress()
    )
    html_text = Path(html.output).read_text(encoding="utf-8")
    assert html_text.count('class="page"') == 3 and "Page 2" in html_text


def test_to_images_and_back(sample_pdf: Path, tmp_path: Path) -> None:
    result = to_images(
        ImagesParams(
            path=str(sample_pdf),
            output_dir=str(tmp_path / "img"),
            format="jpg",
            dpi=72,
            pages="1-2",
        ),
        silent_progress(),
    )
    assert len(result.outputs) == 2 and result.outputs[0].endswith("sample-1.jpg")
    with pytest.raises(OpError) as raised:
        to_images(
            ImagesParams(
                path=str(sample_pdf),
                output_dir=str(tmp_path / "img"),
                format="jpg",
                dpi=72,
                pages="1",
            ),
            silent_progress(),
        )
    assert raised.value.data and raised.value.data.get("exists") is True

    pdf = images_to_pdf(
        ImagesToPdfParams(images=result.outputs, output=str(tmp_path / "back.pdf"), page_size="a4"),
        silent_progress(),
    )
    assert pdf.page_count == 2
    with pymupdf.open(pdf.output) as document:
        assert round(document[0].rect.width) == 595


def test_images_to_pdf_from_folder_sorted_naturally(image_folder: Path, tmp_path: Path) -> None:
    result = images_to_pdf(
        ImagesToPdfParams(
            folders=[str(image_folder)], output=str(tmp_path / "album.pdf"), page_size="image"
        ),
        silent_progress(),
    )
    assert result.page_count == 3
    with pymupdf.open(result.output) as document:
        widths = [round(page.rect.width) for page in document]
        assert widths == [75, 150, 225]
    with pytest.raises(OpError) as raised:
        images_to_pdf(
            ImagesToPdfParams(folders=[str(tmp_path / "missing")], output=str(tmp_path / "x.pdf")),
            silent_progress(),
        )
    assert raised.value.code == ErrorCode.FILE_NOT_FOUND


def test_to_xlsx_extracts_table(table_pdf: Path, tmp_path: Path) -> None:
    from openpyxl import load_workbook

    result = to_xlsx(
        XlsxParams(path=str(table_pdf), output=str(tmp_path / "t.xlsx")), silent_progress()
    )
    workbook = load_workbook(result.output)
    sheet = workbook[workbook.sheetnames[0]]
    rows = [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]
    assert result.table_count >= 1
    assert any("Ayşe" in str(cell) for row in rows for cell in row)


def test_to_pptx_and_docx(sample_pdf: Path, tmp_path: Path) -> None:
    from pptx import Presentation

    pptx = to_pptx(
        PptxParams(path=str(sample_pdf), output=str(tmp_path / "deck.pptx"), dpi=60, pages="1-2"),
        silent_progress(),
    )
    presentation = Presentation(pptx.output)
    assert len(presentation.slides) == 2 and pptx.page_count == 2

    docx = to_docx(
        DocxParams(path=str(sample_pdf), output=str(tmp_path / "doc.docx"), pages="1-2"),
        silent_progress(),
    )
    from docx import Document

    text = "\n".join(paragraph.text for paragraph in Document(docx.output).paragraphs)
    assert "Page 1" in text


def test_file_to_pdf_text_markdown_html(tmp_path: Path) -> None:
    (tmp_path / "notes.txt").write_text("Merhaba dünya\nşişe çığ\n" * 300, encoding="utf-8")
    (tmp_path / "readme.md").write_text(
        "# Başlık\n\n- madde **kalın**\n\n| a | b |\n|---|---|\n| 1 | 2 |\n", encoding="utf-8"
    )
    (tmp_path / "page.html").write_text("<h1>HTML Başlık</h1><p>paragraf</p>", encoding="utf-8")
    text = file_to_pdf(
        FileToPdfParams(path=str(tmp_path / "notes.txt"), output=str(tmp_path / "notes.pdf")),
        silent_progress(),
    )
    assert text.page_count >= 3
    with pymupdf.open(text.output) as document:
        assert "şişe" in document[0].get_text()
    markdown = file_to_pdf(
        FileToPdfParams(path=str(tmp_path / "readme.md"), output=str(tmp_path / "readme.pdf")),
        silent_progress(),
    )
    with pymupdf.open(markdown.output) as document:
        assert "Başlık" in document[0].get_text() and "kalın" in document[0].get_text()
    html = file_to_pdf(
        FileToPdfParams(path=str(tmp_path / "page.html"), output=str(tmp_path / "page.pdf")),
        silent_progress(),
    )
    assert html.page_count == 1


def test_file_to_pdf_image_and_unsupported(image_folder: Path, tmp_path: Path) -> None:
    result = file_to_pdf(
        FileToPdfParams(path=str(image_folder / "img2.png"), output=str(tmp_path / "img.pdf")),
        silent_progress(),
    )
    assert result.page_count == 1
    (tmp_path / "data.bin").write_bytes(b"\x00\x01")
    with pytest.raises(OpError) as raised:
        file_to_pdf(
            FileToPdfParams(path=str(tmp_path / "data.bin"), output=str(tmp_path / "bin.pdf")),
            silent_progress(),
        )
    assert raised.value.code == ErrorCode.INVALID_PARAMS
    with pytest.raises(OpError) as missing:
        file_to_pdf(
            FileToPdfParams(path=str(tmp_path / "nope.txt"), output=str(tmp_path / "nope.pdf")),
            silent_progress(),
        )
    assert missing.value.code == ErrorCode.FILE_NOT_FOUND


def test_office_conversion_requires_libreoffice(tmp_path: Path) -> None:
    (tmp_path / "letter.docx").write_bytes(b"PK\x03\x04")
    if libreoffice.find_soffice() is not None:
        pytest.skip("LibreOffice installed; behaviour exercised manually")
    with pytest.raises(OpError) as raised:
        file_to_pdf(
            FileToPdfParams(
                path=str(tmp_path / "letter.docx"), output=str(tmp_path / "letter.pdf")
            ),
            silent_progress(),
        )
    assert raised.value.code == ErrorCode.LIBREOFFICE_MISSING


def test_tools_status_reports_languages() -> None:
    status = tools_status(EmptyParams(), silent_progress())
    assert "tur" in status.ocr_languages and "eng" in status.ocr_languages


def test_to_images_webp(sample_pdf: Path, tmp_path: Path) -> None:
    result = to_images(
        ImagesParams(
            path=str(sample_pdf),
            output_dir=str(tmp_path / "webp"),
            format="webp",
            dpi=72,
            quality=60,
            pages="1",
        ),
        silent_progress(),
    )
    assert result.outputs[0].endswith("sample-1.webp")
    assert Path(result.outputs[0]).read_bytes()[8:12] == b"WEBP"
