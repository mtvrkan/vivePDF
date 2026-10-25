from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.oxml.ns import qn

from vivepdf.ops.convert_slides import PptxParams, to_pptx
from vivepdf.rpc.progress import silent_progress

SIDECAR = Path(__file__).resolve().parent.parent
FONT = str(SIDECAR / "vivepdf" / "assets" / "fonts" / "DejaVuSans.ttf")
EMU_PER_POINT = 12700


def _write(page: pymupdf.Page, point: tuple[float, float], text: str) -> None:
    page.insert_text(point, text, fontsize=14, fontname="dejavu", fontfile=FONT)


def _presentation(document: pymupdf.Document, tmp_path: Path, **options):
    source = tmp_path / "kaynak.pdf"
    document.save(source)
    document.close()
    result = to_pptx(
        PptxParams(path=str(source), output=str(tmp_path / "sunum.pptx"), **options),
        silent_progress(),
    )
    return Presentation(result.output)


def _linked_runs(slide) -> list[tuple[str, str | None]]:
    return [
        (run.text, run.hyperlink.address)
        for shape in slide.shapes
        if shape.has_text_frame
        for run in shape.text_frame.paragraphs[0].runs
        if run._r.find(qn("a:rPr")) is not None
        and run._r.find(qn("a:rPr")).find(qn("a:hlinkClick")) is not None
    ]


def _link_areas(slide) -> list:
    return [shape for shape in slide.shapes if shape.name == "Link"]


def test_a_web_link_stays_on_its_words_only(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    _write(page, (50, 80), "Bakınız site adresi burada")
    page.insert_link(
        {
            "kind": pymupdf.LINK_URI,
            "from": page.search_for("site adresi")[0],
            "uri": "https://x.org/a b",
        }
    )

    slide = _presentation(document, tmp_path).slides[0]

    assert _linked_runs(slide) == [("site adresi", "https://x.org/a%20b")]
    assert _link_areas(slide) == []


def test_a_link_to_another_page_jumps_to_its_slide(tmp_path: Path) -> None:
    document = pymupdf.open()
    document.new_page(width=400, height=300)
    document.new_page(width=400, height=300)
    first = document[0]
    _write(first, (50, 80), "İkinci sayfaya git")
    first.insert_link(
        {"kind": pymupdf.LINK_GOTO, "from": first.search_for("İkinci sayfaya")[0], "page": 1}
    )

    presentation = _presentation(document, tmp_path)

    slide = presentation.slides[0]
    (run,) = [
        run
        for shape in slide.shapes
        if shape.has_text_frame
        for run in shape.text_frame.paragraphs[0].runs
        if run.text == "İkinci sayfaya"
    ]
    click = run._r.find(qn("a:rPr")).find(qn("a:hlinkClick"))
    assert click.get("action") == "ppaction://hlinksldjump"
    assert slide.part.related_part(click.get(qn("r:id"))) is presentation.slides[1].part


def test_a_link_over_a_picture_becomes_a_clear_click_area(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 10, 10), 0)
    pixmap.clear_with(90)
    page.insert_image(pymupdf.Rect(100, 100, 200, 200), pixmap=pixmap)
    page.insert_link(
        {"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(100, 100, 200, 200), "uri": "https://x.org"}
    )

    slide = _presentation(document, tmp_path).slides[0]

    (area,) = _link_areas(slide)
    assert area.click_action.hyperlink.address == "https://x.org"
    assert (
        area._element.spPr.find(qn("a:solidFill"))
        .find(qn("a:srgbClr"))
        .find(qn("a:alpha"))
        .get("val")
        == "0"
    )
    assert area._element.spPr.find(qn("a:ln")).find(qn("a:noFill")) is not None


def test_picture_slides_keep_their_links_as_click_areas(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    _write(page, (50, 80), "Bakınız site adresi")
    page.insert_link(
        {
            "kind": pymupdf.LINK_URI,
            "from": page.search_for("site adresi")[0],
            "uri": "https://x.org",
        }
    )

    slide = _presentation(document, tmp_path, mode="image").slides[0]

    assert [area.click_action.hyperlink.address for area in _link_areas(slide)] == ["https://x.org"]


def test_script_and_file_links_are_dropped(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    _write(page, (50, 80), "Tehlikeli bağlantı")
    area = page.search_for("Tehlikeli")[0]
    page.insert_link({"kind": pymupdf.LINK_URI, "from": area, "uri": "javascript:alert(1)"})
    page.insert_link({"kind": pymupdf.LINK_LAUNCH, "from": area, "file": "calc.exe"})

    slide = _presentation(document, tmp_path).slides[0]

    assert _linked_runs(slide) == [] and _link_areas(slide) == []


def test_a_link_to_a_page_left_out_is_dropped(tmp_path: Path) -> None:
    document = pymupdf.open()
    document.new_page(width=400, height=300)
    document.new_page(width=400, height=300)
    first = document[0]
    _write(first, (50, 80), "İkinci sayfaya git")
    first.insert_link(
        {"kind": pymupdf.LINK_GOTO, "from": first.search_for("İkinci sayfaya")[0], "page": 1}
    )

    slide = _presentation(document, tmp_path, pages="1").slides[0]

    assert _linked_runs(slide) == []


def test_click_area_sits_on_the_link_of_a_cropped_turned_page(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    page.insert_link(
        {"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(100, 100, 200, 150), "uri": "https://x.org"}
    )
    page.set_cropbox(pymupdf.Rect(40, 20, 380, 290))
    page.set_rotation(90)
    shown = pymupdf.Rect(document.reload_page(page).get_links()[0]["from"])

    presentation = _presentation(document, tmp_path)

    (area,) = _link_areas(presentation.slides[0])
    ratio = presentation.slide_width / (270 * EMU_PER_POINT)
    assert abs(area.left / EMU_PER_POINT / ratio - shown.x0) < 0.5
    assert abs(area.top / EMU_PER_POINT / ratio - shown.y0) < 0.5
    assert abs(area.width / EMU_PER_POINT / ratio - shown.width) < 0.5
    assert abs(area.height / EMU_PER_POINT / ratio - shown.height) < 0.5
