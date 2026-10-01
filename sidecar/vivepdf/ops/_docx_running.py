import io
import math
import re
import secrets
import statistics
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

import pymupdf

from vivepdf.ops._docx_bidi import mark_paragraph, mostly_right_to_left

Band = Literal["header", "footer"]
Slot = Literal["left", "center", "right"]

BAND_SHARE = 0.1
MIN_PAGES = 3
PRESENCE_SHARE = 0.6
CENTRE_SHARE = 0.1
ROW_TOLERANCE = 4.0
MAX_LINE_CHARACTERS = 200
NUMBER = re.compile(r"\d+")
PAGE_FIELD = "PAGE"


@dataclass(slots=True)
class _Occurrence:
    page: int
    position: int
    rect: pymupdf.Rect
    numbers: list[str]
    size: float
    slot: Slot


@dataclass(slots=True)
class RunningLine:
    band: Band
    pattern: str
    parts: list[str | None]
    size: float
    slot: Slot
    top: float
    left: float


@dataclass(slots=True)
class RunningText:
    lines: list[RunningLine] = field(default_factory=list)
    areas: dict[int, list[pymupdf.Rect]] = field(default_factory=dict)
    first_page_number: int = 1

    def count(self, band: Band) -> int:
        return sum(1 for line in self.lines if line.band == band)


def _slot(rect: pymupdf.Rect, page_rect: pymupdf.Rect) -> Slot:
    middle = (rect.x0 + rect.x1) / 2
    centre = (page_rect.x0 + page_rect.x1) / 2
    if abs(middle - centre) <= page_rect.width * CENTRE_SHARE:
        return "center"
    return "left" if middle < centre else "right"


def _band_lines(page: pymupdf.Page) -> list[tuple[Band, str, pymupdf.Rect, float]]:
    area = page.rect
    top_limit = area.y0 + area.height * BAND_SHARE
    bottom_limit = area.y1 - area.height * BAND_SHARE
    found: list[tuple[Band, str, pymupdf.Rect, float]] = []
    has_body = any(
        pymupdf.Rect(image["bbox"]).y1 > top_limit and pymupdf.Rect(image["bbox"]).y0 < bottom_limit
        for image in page.get_image_info()
    )
    for block in page.get_text("dict", flags=pymupdf.TEXTFLAGS_TEXT)["blocks"]:
        for line in block.get("lines", []):
            if tuple(line["dir"]) != (1.0, 0.0):
                continue
            text = " ".join("".join(span["text"] for span in line["spans"]).split())
            if not text or len(text) > MAX_LINE_CHARACTERS:
                continue
            rect = pymupdf.Rect(line["bbox"])
            if rect.y1 <= top_limit:
                band: Band = "header"
            elif rect.y0 >= bottom_limit:
                band = "footer"
            else:
                has_body = True
                continue
            size = max(span["size"] for span in line["spans"])
            found.append((band, text, rect, size))
    return found if has_body else []


def _page_start(occurrences: list[_Occurrence], slot: int) -> int | None:
    starts = {int(occurrence.numbers[slot]) - occurrence.position for occurrence in occurrences}
    return starts.pop() if len(starts) == 1 else None


def _parts(pattern: str, occurrences: list[_Occurrence]) -> tuple[list[str | None], int | None]:
    literals = NUMBER.split(pattern)
    parts: list[str | None] = [literals[0]] if literals[0] else []
    start: int | None = None
    for slot, literal in enumerate(literals[1:]):
        values = {occurrence.numbers[slot] for occurrence in occurrences}
        if len(values) == 1:
            parts.append(values.pop())
        else:
            tracked = _page_start(occurrences, slot)
            if tracked is None or tracked < 0 or start not in (None, tracked):
                raise ValueError("numbers do not follow the pages")
            start = tracked
            parts.append(None)
        if literal:
            parts.append(literal)
    return parts, start


def find_running_text(document: pymupdf.Document, indices: list[int]) -> RunningText:
    running = RunningText()
    if len(indices) < MIN_PAGES:
        return running
    seen: dict[tuple[Band, str], list[_Occurrence]] = defaultdict(list)
    for position, index in enumerate(indices):
        page = document[index]
        for band, text, rect, size in _band_lines(page):
            pattern = NUMBER.sub("0", text)
            occurrence = _Occurrence(
                index, position, rect, NUMBER.findall(text), size, _slot(rect, page.rect)
            )
            bucket = seen[(band, pattern)]
            if not bucket or bucket[-1].page != index:
                bucket.append(occurrence)
    needed = max(MIN_PAGES, math.ceil(len(indices) * PRESENCE_SHARE))
    for (band, pattern), occurrences in seen.items():
        if len(occurrences) < needed:
            continue
        try:
            parts, start = _parts(pattern, occurrences)
        except ValueError:
            continue
        if start is not None:
            running.first_page_number = start
        slots = [occurrence.slot for occurrence in occurrences]
        running.lines.append(
            RunningLine(
                band=band,
                pattern=pattern,
                parts=parts,
                size=statistics.median(occurrence.size for occurrence in occurrences),
                slot=max(("left", "center", "right"), key=slots.count),
                top=statistics.median(occurrence.rect.y0 for occurrence in occurrences),
                left=statistics.median(occurrence.rect.x0 for occurrence in occurrences),
            )
        )
        for occurrence in occurrences:
            running.areas.setdefault(occurrence.page, []).append(occurrence.rect)
    running.lines.sort(key=lambda line: (line.band, line.top, line.left))
    return running


def clear_running_text(document: pymupdf.Document, running: RunningText) -> None:
    for index, areas in running.areas.items():
        page = document[index]
        for area in areas:
            page.add_redact_annot(area)
        page.apply_redactions(
            images=pymupdf.PDF_REDACT_IMAGE_NONE,
            graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
            text=pymupdf.PDF_REDACT_TEXT_REMOVE,
        )


def cleared_copy(
    document: pymupdf.Document, running: RunningText, folder: Path, locked: bool
) -> tuple[str, str | None]:
    clear_running_text(document, running)
    target = folder / "source.pdf"
    if not locked:
        document.save(target, encryption=pymupdf.PDF_ENCRYPT_NONE)
        return str(target), None
    secret = secrets.token_urlsafe(24)
    document.save(target, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw=secret, user_pw=secret)
    return str(target), secret


def _rows(lines: list[RunningLine]) -> list[list[RunningLine]]:
    rows: list[list[RunningLine]] = []
    for line in lines:
        if rows and abs(rows[-1][0].top - line.top) <= ROW_TOLERANCE:
            rows[-1].append(line)
        else:
            rows.append([line])
    return rows


def _add_field(paragraph, instruction: str, size: float) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    simple = OxmlElement("w:fldSimple")
    simple.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    half_points = OxmlElement("w:sz")
    half_points.set(qn("w:val"), str(round(size * 2)))
    properties.append(half_points)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    simple.append(run)
    paragraph._p.append(simple)


def _write_line(paragraph, line: RunningLine) -> None:
    from docx.shared import Pt

    for part in line.parts:
        if part is None:
            _add_field(paragraph, PAGE_FIELD, line.size)
        else:
            paragraph.add_run(part).font.size = Pt(line.size)


def _write_row(paragraph, row: list[RunningLine], width) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Emu

    if len(row) == 1:
        paragraph.alignment = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }[row[0].slot]
        _write_line(paragraph, row[0])
    else:
        stops = paragraph.paragraph_format.tab_stops
        stops.add_tab_stop(Emu(width // 2), WD_TAB_ALIGNMENT.CENTER)
        stops.add_tab_stop(Emu(width), WD_TAB_ALIGNMENT.RIGHT)
        by_slot: dict[Slot, list[RunningLine]] = {}
        for line in sorted(row, key=lambda line: line.left):
            by_slot.setdefault(line.slot, []).append(line)
        for position, slot in enumerate(("left", "center", "right")):
            if position:
                paragraph.add_run("\t")
            for order, line in enumerate(by_slot.get(slot, [])):
                if order:
                    paragraph.add_run(" ")
                _write_line(paragraph, line)
    if mostly_right_to_left(paragraph.text):
        mark_paragraph(paragraph)


def _number_pages_from(section, start: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = section._sectPr
    numbering = properties.find(qn("w:pgNumType"))
    if numbering is None:
        numbering = OxmlElement("w:pgNumType")
        properties.append(numbering)
    numbering.set(qn("w:start"), str(start))


def add_running_text(data: bytes, running: RunningText) -> bytes:
    from docx import Document

    document = Document(io.BytesIO(data))
    first = document.sections[0]
    if running.first_page_number != 1:
        _number_pages_from(first, running.first_page_number)
    width = first.page_width - first.left_margin - first.right_margin
    for band in ("header", "footer"):
        rows = _rows([line for line in running.lines if line.band == band])
        if not rows:
            continue
        part = first.header if band == "header" else first.footer
        part.is_linked_to_previous = False
        paragraphs = part.paragraphs
        for position, row in enumerate(rows):
            paragraph = paragraphs[0] if position == 0 and paragraphs else part.add_paragraph()
            _write_row(paragraph, row, width)
    for section in document.sections[1:]:
        section.header.is_linked_to_previous = True
        section.footer.is_linked_to_previous = True
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
