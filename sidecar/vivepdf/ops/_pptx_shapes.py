from dataclasses import dataclass, field

import pymupdf

from vivepdf.ops._pptx_layout import SlideFrame

MAX_PAGE_SHAPES = 400
MAX_ERASE_PASSES = 16
AREA_MARGIN = 0.75
AXIS_TOLERANCE = 0.01
OVERLAP_TOLERANCE = 0.01
HAIRLINE_POINTS = 0.5
WHITE_LEVEL = 0.99
SOLID_DASHES = ("", "[] 0", "[]0")
PAINT_KINDS = frozenset({"fill-path", "stroke-path", "fill-image", "fill-imgmask", "fill-shade"})
PATH_KINDS = frozenset({"fill-path", "stroke-path"})
TEXT_KINDS = frozenset({"fill-text", "stroke-text"})
PICTURE_KINDS = frozenset({"fill-image", "fill-imgmask"})
BACKDROP_SHARE = 0.5
LINE_CAPS = {0: "flat", 1: "rnd", 2: "sq"}


@dataclass(frozen=True, slots=True)
class SlideShape:
    kind: str
    points: tuple[float, float, float, float]
    fill: tuple[float, float, float] | None
    fill_opacity: float
    stroke: tuple[float, float, float] | None
    stroke_opacity: float
    width: float
    cap: str


@dataclass(slots=True)
class PageShapes:
    shapes: list[SlideShape] = field(default_factory=list)
    passes: list[list[pymupdf.Rect]] = field(default_factory=list)


@dataclass(slots=True)
class _Candidate:
    entries: set[int]
    area: pymupdf.Rect
    box: pymupdf.Rect
    shapes: list[SlideShape]


def _rgb(colour) -> tuple[float, float, float] | None:
    if not colour:
        return None
    if len(colour) == 1:
        return (colour[0],) * 3
    if len(colour) == 4:
        cyan, magenta, yellow, black = colour
        return tuple((1 - part) * (1 - black) for part in (cyan, magenta, yellow))
    return tuple(colour[:3])


def _axis_rect(item) -> pymupdf.Rect | None:
    if item[0] == "re":
        return pymupdf.Rect(item[1])
    quad = item[1]
    xs = sorted({round(point.x / AXIS_TOLERANCE) for point in quad})
    ys = sorted({round(point.y / AXIS_TOLERANCE) for point in quad})
    if len(xs) != 2 or len(ys) != 2:
        return None
    return quad.rect


def _overlapping(rects: list[pymupdf.Rect]) -> bool:
    for position, first in enumerate(rects):
        for second in rects[position + 1 :]:
            common = first & second
            if common.width > OVERLAP_TOLERANCE and common.height > OVERLAP_TOLERANCE:
                return True
    return False


def _style(drawing: dict) -> dict:
    fill = _rgb(drawing.get("fill")) if "f" in drawing["type"] else None
    stroke = _rgb(drawing.get("color")) if "s" in drawing["type"] else None
    caps = drawing.get("lineCap") or (0,)
    return {
        "fill": fill,
        "fill_opacity": drawing.get("fill_opacity") or 1.0,
        "stroke": stroke,
        "stroke_opacity": drawing.get("stroke_opacity") or 1.0,
        "width": (drawing.get("width") or HAIRLINE_POINTS) if stroke else 0.0,
        "cap": LINE_CAPS.get(caps[0], "flat"),
    }


def _drawing_shapes(drawing: dict, matrix: pymupdf.Matrix) -> list[SlideShape] | None:
    items = drawing.get("items") or []
    if not items or (drawing.get("dashes") or "") not in SOLID_DASHES:
        return None
    style = _style(drawing)
    if style["stroke"] is None and (style["fill"] is None or min(style["fill"]) >= WHITE_LEVEL):
        return None
    kinds = {item[0] for item in items}
    if kinds <= {"re", "qu"}:
        rects = [_axis_rect(item) for item in items]
        if any(rect is None for rect in rects):
            return None
        if len(rects) > 1 and style["fill"] is not None and _overlapping(rects):
            return None
        return [
            SlideShape("rect", tuple(rect * matrix), **style)
            for rect in rects
            if not (rect * matrix).is_empty or style["stroke"] is not None
        ]
    if kinds == {"l"} and style["fill"] is None:
        return [
            SlideShape("line", (*(item[1] * matrix), *(item[2] * matrix)), **style)
            for item in items
        ]
    return None


def _entries(drawing: dict) -> set[int]:
    first = drawing["seqno"]
    return {first, first + 1} if drawing["type"] == "fs" else {first}


def _inside(box, frame) -> bool:
    return box[0] >= frame[0] and box[1] >= frame[1] and box[2] <= frame[2] and box[3] <= frame[3]


def _clipped(drawing: dict, scissor: pymupdf.Rect | None) -> bool:
    if scissor is None:
        return False
    padding = (drawing.get("width") or 0) / 2 + AREA_MARGIN
    return not _inside(
        pymupdf.Rect(drawing["rect"]), scissor + (-padding, -padding, padding, padding)
    )


def _annotation_areas(page: pymupdf.Page) -> list[pymupdf.Rect]:
    margin = (-AREA_MARGIN, -AREA_MARGIN, AREA_MARGIN, AREA_MARGIN)
    return [
        pymupdf.Rect(annotation.rect) + margin for annotation in (*page.annots(), *page.widgets())
    ]


def _scanned(page: pymupdf.Page) -> list[_Candidate]:
    matrix = page.rotation_matrix
    annotations = _annotation_areas(page)
    candidates: list[_Candidate] = []
    stack: list[tuple[int, pymupdf.Rect | None, bool]] = []
    for drawing in page.get_drawings(extended=True):
        level = drawing.get("level", 0)
        while stack and stack[-1][0] >= level:
            stack.pop()
        scissor = stack[-1][1] if stack else None
        plain = stack[-1][2] if stack else True
        if drawing["type"] == "clip":
            clip = pymupdf.Rect(drawing["scissor"])
            stack.append((level, clip if scissor is None else clip & scissor, plain))
            continue
        if drawing["type"] == "group":
            ordinary = (drawing.get("opacity") in (None, 1, 1.0)) and (
                drawing.get("blendmode") in (None, "Normal")
            )
            stack.append((level, scissor, plain and ordinary))
            continue
        if "seqno" not in drawing:
            continue
        if not plain or _clipped(drawing, scissor):
            continue
        if any(_inside(pymupdf.Rect(drawing["rect"]), area) for area in annotations):
            continue
        shapes = _drawing_shapes(drawing, matrix)
        if not shapes:
            continue
        padding = (drawing.get("width") or 0) / 2 + AREA_MARGIN
        area = pymupdf.Rect(drawing["rect"]) + (-padding, -padding, padding, padding)
        box = pymupdf.Rect(drawing["rect"])
        candidates.append(_Candidate(_entries(drawing), area, box, shapes))
    return candidates


def _overlap(first: tuple, second: tuple) -> bool:
    return (
        min(first[2], second[2]) - max(first[0], second[0]) > OVERLAP_TOLERANCE
        and min(first[3], second[3]) - max(first[1], second[1]) > OVERLAP_TOLERANCE
    )


def _hides(candidate: _Candidate) -> bool:
    return any(shape.fill is not None and shape.fill_opacity >= 1 for shape in candidate.shapes)


def _backdrop(box: tuple, page_area: float) -> bool:
    return pymupdf.Rect(box).get_area() >= BACKDROP_SHARE * page_area


def _blocked(
    candidate: _Candidate, background: list[tuple[int, str, tuple]], page_area: float
) -> bool:
    first, last = min(candidate.entries), max(candidate.entries)
    area, own = tuple(candidate.area), tuple(candidate.box)
    for index, kind, box in background:
        if kind in TEXT_KINDS:
            if index < first and _overlap(box, own):
                return True
            continue
        if (
            kind in PICTURE_KINDS
            and index < first
            and _overlap(box, own)
            and _hides(candidate)
            and not _backdrop(box, page_area)
        ):
            return True
        if kind in PATH_KINDS and _inside(box, area):
            return True
        if index > last and _overlap(box, area):
            return True
    return False


def _erase_passes(areas: list[pymupdf.Rect]) -> list[list[pymupdf.Rect]]:
    passes: list[list[pymupdf.Rect]] = []
    for area in areas:
        reach = area + (-AREA_MARGIN, -AREA_MARGIN, AREA_MARGIN, AREA_MARGIN)
        free = next(
            (group for group in passes if not any(reach.intersects(other) for other in group)),
            None,
        )
        if free is None:
            passes.append([area])
        else:
            free.append(area)
    return passes


def page_shapes(page: pymupdf.Page) -> PageShapes:
    candidates = _scanned(page)
    if not candidates:
        return PageShapes()
    log = [
        (index, kind, tuple(box))
        for index, (kind, box) in enumerate(page.get_bboxlog())
        if kind in PAINT_KINDS or kind in TEXT_KINDS
    ]
    while True:
        taken = set().union(*(candidate.entries for candidate in candidates))
        background = [entry for entry in log if entry[0] not in taken]
        kept = [
            candidate
            for candidate in candidates
            if not _blocked(candidate, background, page.rect.get_area())
        ]
        if len(kept) == len(candidates):
            break
        candidates = kept
        if not candidates:
            return PageShapes()
    shapes = [shape for candidate in candidates for shape in candidate.shapes]
    if len(shapes) > MAX_PAGE_SHAPES:
        return PageShapes()
    passes = _erase_passes([candidate.area for candidate in candidates])
    if len(passes) > MAX_ERASE_PASSES:
        return PageShapes()
    return PageShapes(shapes, passes)


def _hex(colour: tuple[float, float, float]) -> str:
    return "".join(f"{max(0, min(255, round(part * 255))):02X}" for part in colour)


def _colour(parent, colour: tuple[float, float, float], opacity: float) -> None:
    from pptx.oxml.ns import qn

    solid = parent.makeelement(qn("a:solidFill"), {})
    rgb = solid.makeelement(qn("a:srgbClr"), {"val": _hex(colour)})
    if opacity < 1:
        rgb.append(rgb.makeelement(qn("a:alpha"), {"val": str(round(opacity * 100000))}))
    solid.append(rgb)
    parent.append(solid)


def add_slide_shape(slide, shape: SlideShape, frame: SlideFrame) -> None:
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.oxml.ns import qn
    from pptx.util import Emu

    x0, y0, x1, y1 = (
        round(frame.left + shape.points[0] * frame.ratio),
        round(frame.top + shape.points[1] * frame.ratio),
        round(frame.left + shape.points[2] * frame.ratio),
        round(frame.top + shape.points[3] * frame.ratio),
    )
    if shape.kind == "rect":
        added = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Emu(x0), Emu(y0), Emu(max(1, x1 - x0)), Emu(max(1, y1 - y0))
        )
    else:
        added = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Emu(x0), Emu(y0), Emu(x1), Emu(y1)
        )
    element = added._element
    style = element.find(qn("p:style"))
    if style is not None:
        element.remove(style)
    properties = element.spPr
    for name in ("a:solidFill", "a:noFill", "a:ln"):
        for old in properties.findall(qn(name)):
            properties.remove(old)
    if shape.kind == "rect":
        if shape.fill is None:
            properties.append(properties.makeelement(qn("a:noFill"), {}))
        else:
            _colour(properties, shape.fill, shape.fill_opacity)
    outline = properties.makeelement(qn("a:ln"), {})
    if shape.stroke is None:
        outline.append(outline.makeelement(qn("a:noFill"), {}))
    else:
        outline.set("w", str(max(1, round(shape.width * frame.ratio))))
        outline.set("cap", shape.cap)
        _colour(outline, shape.stroke, shape.stroke_opacity)
    properties.append(outline)
