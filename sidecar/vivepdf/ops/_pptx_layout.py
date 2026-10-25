import math
import re
from dataclasses import dataclass

import pymupdf

from vivepdf.ops._hidden_text import origin_key
from vivepdf.ops._line_gaps import wide_gap_starts
from vivepdf.ops._pptx_links import PageLink, SlideJumps, link_at, link_run

EMU_PER_POINT = 12700
MIN_FONT_POINTS = 1.0
MAX_FONT_POINTS = 4000.0
BOLD_FLAG = 16
ITALIC_FLAG = 2
FILLED_FLAG = 16
STROKED_FLAG = 32
STROKE_TRACE = 1
TEXT_FLAGS = pymupdf.TEXTFLAGS_TEXT | pymupdf.TEXT_COLLECT_STYLES
SLANT = 1e-3
SUBSET_PREFIX = re.compile(r"^[A-Z]{6}\+")
UNSAFE_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
FONT_SUFFIXES = ("PSMT", "MT", "PS")
FONT_STYLE = re.compile(
    r"(?:(?:semi|demi|extra|ultra)?(?:bold|light|black|heavy|thin|medium|regular|roman|book"
    r"|italic|oblique|condensed|narrow|it|bd|bi|lt|md|rg|bk))+(?:psmt|mt|ps)?",
    re.IGNORECASE,
)
FONT_FAMILIES = {
    "Helvetica": "Arial",
    "Times": "Times New Roman",
    "TimesNewRoman": "Times New Roman",
    "Courier": "Courier New",
    "CourierNew": "Courier New",
}
DEFAULT_FAMILY = "Arial"


@dataclass(slots=True)
class SlideFrame:
    left: int
    top: int
    ratio: float


@dataclass(slots=True)
class LineBox:
    left: float
    top: float
    width: float
    height: float
    angle: float
    spans: list[dict]
    area: pymupdf.Rect


def font_family(name: str) -> str:
    parts = re.split(r"[-,]", SUBSET_PREFIX.sub("", name or ""))
    kept = parts[:1]
    for part in parts[1:]:
        if not part or FONT_STYLE.fullmatch(part):
            break
        kept.append(part)
    base = " ".join(kept)
    for suffix in FONT_SUFFIXES:
        if base.endswith(suffix) and len(base) > len(suffix):
            base = base[: -len(suffix)]
            break
    return FONT_FAMILIES.get(base, base) or DEFAULT_FAMILY


def slide_text(text: str) -> str:
    return UNSAFE_CHARACTERS.sub("", text)


def _visible(span: dict) -> bool:
    return span.get("alpha", 255) > 0


def _direction(line: dict, matrix: pymupdf.Matrix) -> tuple[float, float]:
    origin = pymupdf.Point(0, 0) * matrix
    moved = pymupdf.Point(*line["dir"]) * matrix
    return moved.x - origin.x, moved.y - origin.y


def _piece(span: dict, chars: list[dict], link: PageLink | None) -> dict:
    bounds = pymupdf.Rect(chars[0]["bbox"])
    for char in chars[1:]:
        bounds |= char["bbox"]
    piece = {key: value for key, value in span.items() if key != "chars"}
    piece["text"] = "".join(char["c"] for char in chars)
    piece["bbox"] = tuple(bounds)
    piece["link"] = link
    return piece


def _segments(
    line: dict, hidden: set[tuple[float, float]], links: list[PageLink], breaks: set[int]
) -> list[list[dict]]:
    segments: list[list[dict]] = [[]]
    index = 0
    for span in line["spans"]:
        run: list[dict] = []
        target: PageLink | None = None
        for char in span["chars"]:
            covered = origin_key(char["origin"]) in hidden
            if covered or index in breaks:
                if run:
                    segments[-1].append(_piece(span, run, target))
                    run = []
                if segments[-1]:
                    segments.append([])
            index += 1
            if covered:
                continue
            link = link_at(links, char["bbox"]) if links else None
            if run and link is not target:
                segments[-1].append(_piece(span, run, target))
                run = []
            target = link
            run.append(char)
        if run:
            segments[-1].append(_piece(span, run, target))
    return [segment for segment in segments if segment]


def _line_box(line: dict, bounds: pymupdf.Rect, spans: list[dict], matrix) -> LineBox:
    shown = bounds * matrix
    dx, dy = _direction(line, matrix)
    angle = round(math.degrees(math.atan2(dy, dx)), 1) % 360
    if min(abs(part) for part in line["dir"]) > SLANT:
        quad = pymupdf.recover_line_quad(line, spans) * matrix
        length = abs(quad.ur - quad.ul)
        thickness = abs(quad.ll - quad.ul)
        center = (quad.ul + quad.lr) / 2
    else:
        along_x = abs(line["dir"][0]) >= abs(line["dir"][1])
        length = bounds.width if along_x else bounds.height
        thickness = bounds.height if along_x else bounds.width
        center = (shown.tl + shown.br) / 2
    return LineBox(
        left=center.x - length / 2,
        top=center.y - thickness / 2,
        width=length,
        height=thickness,
        angle=angle,
        spans=spans,
        area=bounds,
    )


def _hollow(span: dict) -> bool:
    flags = span.get("char_flags", FILLED_FLAG)
    return bool(flags & STROKED_FLAG) and not flags & FILLED_FLAG


def _mark_outlines(page: pymupdf.Page, blocks: list[dict]) -> None:
    spans = [
        span
        for block in blocks
        for line in block.get("lines", [])
        for span in line["spans"]
        if _hollow(span) and span["chars"]
    ]
    if not spans:
        return
    widths = {
        origin_key(char[2]): trace["linewidth"]
        for trace in page.get_texttrace()
        if trace["type"] == STROKE_TRACE
        for char in trace["chars"]
    }
    for span in spans:
        span["outline"] = widths.get(origin_key(span["chars"][0]["origin"]))


def _span_bounds(spans: list[dict]) -> pymupdf.Rect:
    bounds = pymupdf.Rect(spans[0]["bbox"])
    for span in spans[1:]:
        bounds |= span["bbox"]
    return bounds


def line_boxes(
    page: pymupdf.Page,
    hidden: set[tuple[float, float]] | None = None,
    links: list[PageLink] | None = None,
) -> list[LineBox]:
    hidden = hidden or set()
    matrix = page.rotation_matrix
    boxes: list[LineBox] = []
    blocks = page.get_text("rawdict", flags=TEXT_FLAGS)["blocks"]
    _mark_outlines(page, blocks)
    for block in blocks:
        for line in block.get("lines", []):
            breaks = wide_gap_starts(line)
            whole = not breaks and all(
                _visible(span) and origin_key(char["origin"]) not in hidden
                for span in line["spans"]
                for char in span["chars"]
            )
            for segment in _segments(line, hidden, links or [], breaks):
                spans = [span for span in segment if _visible(span) and span["text"]]
                if not any(slide_text(span["text"]).strip() for span in spans):
                    continue
                bounds = pymupdf.Rect(line["bbox"]) if whole else _span_bounds(spans)
                boxes.append(_line_box(line, bounds, spans, matrix))
    return boxes


def page_notes(page: pymupdf.Page, hidden: set[tuple[float, float]] | None = None) -> str:
    if not hidden:
        return slide_text(page.get_text()).strip()
    lines = []
    for block in page.get_text("rawdict", flags=pymupdf.TEXTFLAGS_TEXT)["blocks"]:
        for line in block.get("lines", []):
            text = "".join(
                char["c"]
                for span in line["spans"]
                for char in span["chars"]
                if origin_key(char["origin"]) not in hidden
            )
            if text.strip():
                lines.append(text)
    return slide_text("\n".join(lines)).strip()


def text_free_pixmap(
    document: pymupdf.Document,
    index: int,
    dpi: int,
    areas: list[pymupdf.Rect],
) -> pymupdf.Pixmap:
    with pymupdf.open() as single:
        single.insert_pdf(document, from_page=index, to_page=index, annots=True)
        page = single[0]
        for area in areas:
            page.add_redact_annot(area, fill=False, cross_out=False)
        page.apply_redactions(
            images=pymupdf.PDF_REDACT_IMAGE_NONE,
            graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
            text=pymupdf.PDF_REDACT_TEXT_REMOVE,
        )
        return page.get_pixmap(dpi=dpi)


def _font_points(size: float, ratio: float) -> float:
    return min(MAX_FONT_POINTS, max(MIN_FONT_POINTS, round(size * ratio / EMU_PER_POINT, 1)))


def _outline_run(run, colour, width: float | None, ratio: float) -> None:
    from lxml import etree
    from pptx.dml.fill import FillFormat
    from pptx.oxml.ns import qn

    properties = run._r.get_or_add_rPr()
    line = etree.SubElement(properties, qn("a:ln"))
    properties.insert(0, line)
    if width:
        line.set("w", str(max(1, round(width * ratio))))
    stroke = FillFormat.from_fill_parent(line)
    stroke.solid()
    stroke.fore_color.rgb = colour
    FillFormat.from_fill_parent(properties).background()


def add_line_box(slide, box: LineBox, frame: SlideFrame, jumps: SlideJumps | None = None) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.util import Emu, Pt

    shape = slide.shapes.add_textbox(
        Emu(round(frame.left + box.left * frame.ratio)),
        Emu(round(frame.top + box.top * frame.ratio)),
        Emu(max(1, round(box.width * frame.ratio))),
        Emu(max(1, round(box.height * frame.ratio))),
    )
    text_frame = shape.text_frame
    text_frame.word_wrap = False
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.margin_left = text_frame.margin_right = 0
    text_frame.margin_top = text_frame.margin_bottom = 0
    text_frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = text_frame.paragraphs[0]
    for span in box.spans:
        text = slide_text(span["text"])
        if not text:
            continue
        run = paragraph.add_run()
        run.text = text
        font = run.font
        font.size = Pt(_font_points(span["size"], frame.ratio))
        font.name = font_family(span["font"])
        strokes = span.get("char_flags", FILLED_FLAG)
        fake_bold = strokes & FILLED_FLAG and strokes & STROKED_FLAG
        font.bold = bool(span["flags"] & BOLD_FLAG or fake_bold)
        font.italic = bool(span["flags"] & ITALIC_FLAG)
        colour = RGBColor.from_string(f"{span['color'] & 0xFFFFFF:06X}")
        if _hollow(span):
            _outline_run(run, colour, span.get("outline"), frame.ratio)
        else:
            font.color.rgb = colour
        if span.get("link") is not None and jumps is not None:
            link_run(run, span["link"], jumps)
    if box.angle:
        shape.rotation = box.angle
