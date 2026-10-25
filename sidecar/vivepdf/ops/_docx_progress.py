import io
import math
import re
from collections.abc import Callable
from pathlib import Path

import pymupdf
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table, _Cell
from pdf2docx import Converter
from pdf2docx.common import docx as docx_helpers
from pdf2docx.common.Collection import ElementCollection
from pdf2docx.common.share import RectType
from pdf2docx.font.Fonts import Fonts
from pdf2docx.image import ImagesExtractor as images_extractor
from pdf2docx.page.Page import Page
from pdf2docx.page.Pages import Pages
from pdf2docx.page.RawPageFitz import RawPageFitz
from pdf2docx.shape.Shapes import Shapes
from pdf2docx.table.Cell import Cell
from pdf2docx.table.TableBlock import TableBlock
from pdf2docx.text.Line import Line
from pdf2docx.text.Lines import Lines
from pdf2docx.text.TextSpan import TextSpan

from vivepdf.ops._docx_bidi import mark_right_to_left
from vivepdf.ops._docx_contours import inner_contours
from vivepdf.ops._docx_fonts import embed_fonts
from vivepdf.ops._docx_graphics import add_unused_graphics
from vivepdf.ops._docx_lines import (
    make_line_keeping_word_gap,
    parse_line_break_before_list_items,
)
from vivepdf.ops._docx_parallel import ParallelParse, worker_count
from vivepdf.ops._docx_running import RunningText, add_running_text
from vivepdf.ops._docx_visible import visible_images
from vivepdf.ops._output import write_atomically
from vivepdf.rpc.errors import ErrorCode, OpError
from vivepdf.rpc.progress import Progress

ANALYSIS_BAND = (0.05, 0.35)
PARSING_BAND = (0.35, 0.8)
MAKING_BAND = (0.82, 0.97)
HIDDEN_TEXT_SHARE = 0.9
HIDDEN_RENDER_MODE = 3
CLIP_PIXEL_BUDGET = 40_000_000
MAX_CLIP_RATIO = 4.0
MIN_IMAGE_AREA = 4
CLIPPED_MARK = "vivepdfClipped"
COVER_TOLERANCE = 1.0
WORD_GAP_EM = 0.2
TILT_TOLERANCE = 0.035
STYLE_SUFFIX = re.compile(r"^(.+),\s*(Bold|Italic|BoldItalic|Oblique|BoldOblique)$", re.IGNORECASE)
BOLD_FLAG = 16
ITALIC_FLAG = 2


class ConversionCancelled(BaseException):
    pass


class _MergeAwareCell(_Cell):
    def merge(self, other_cell):
        table = self._parent
        try:
            merged = super().merge(other_cell)
        except BaseException:
            table.cell_cache = None
            raise
        if isinstance(table, _CachedCellsTable):
            table.cover_span(merged._tc)
        return merged


class _CachedCellsTable(Table):
    cell_cache: list | None = None

    @property
    def _cells(self):
        if self.cell_cache is None:
            cells = Table._cells.fget(self)
            for cell in cells:
                cell.__class__ = _MergeAwareCell
            self.cell_cache = cells
        return self.cell_cache

    def cover_span(self, tc) -> None:
        if self.cell_cache is None:
            return
        columns = self._column_count
        cell = _MergeAwareCell(tc, self)
        for row in range(tc.top, tc.bottom):
            for column in range(tc.left, tc.right):
                self.cell_cache[row * columns + column] = cell


_plain_table_make_docx = TableBlock.make_docx
_unordered_text_style_shapes = Shapes.text_style_shapes.fget
_plain_extract_raw_dict = RawPageFitz.extract_raw_dict


def text_is_only_hidden(page: pymupdf.Page) -> bool:
    hidden = shown = 0
    for span in page.get_texttrace():
        count = len(span.get("chars", ()))
        if span.get("type") == HIDDEN_RENDER_MODE:
            hidden += count
        else:
            shown += count
    return hidden > 0 and hidden >= HIDDEN_TEXT_SHARE * (hidden + shown)


def _is_fill(shape: dict) -> bool:
    return set(shape) == {"bbox", "color"}


def drop_fills_drawn_in_pictures(raw: dict) -> dict:
    clipped = [
        _widened(pymupdf.Rect(block["bbox"]))
        for block in raw.get("blocks", [])
        if block.pop(CLIPPED_MARK, False)
    ]
    if clipped and "shapes" in raw:
        raw["shapes"] = [
            shape
            for shape in raw["shapes"]
            if not (
                _is_fill(shape)
                and any(area.contains(pymupdf.Rect(shape["bbox"])) for area in clipped)
            )
        ]
    return raw


def _is_blank(span: dict) -> bool:
    chars = span.get("chars")
    return bool(chars) and all(not char["c"].strip() for char in chars)


def level_slight_tilts(raw: dict) -> dict:
    for block in raw.get("blocks", []):
        for line in block.get("lines", []):
            x, y = line.get("dir", (1.0, 0.0))
            if x > 0 and abs(y) <= TILT_TOLERANCE:
                line["dir"] = (1.0, 0.0)
            elif y < 0 and abs(x) <= TILT_TOLERANCE:
                line["dir"] = (0.0, -1.0)
    return raw


def split_style_suffixes(raw: dict) -> dict:
    for block in raw.get("blocks", []):
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                match = STYLE_SUFFIX.match(span.get("font") or "")
                if match is None:
                    continue
                style = match.group(2).lower()
                span["font"] = match.group(1).strip()
                if "bold" in style:
                    span["flags"] = span.get("flags", 0) | BOLD_FLAG
                if "italic" in style or "oblique" in style:
                    span["flags"] = span.get("flags", 0) | ITALIC_FLAG
    return raw


def keep_word_spaces(raw: dict) -> dict:
    for block in raw.get("blocks", []):
        for line in block.get("lines", []):
            spans = line.get("spans", [])
            kept: list[dict] = []
            for index, span in enumerate(spans):
                words_follow = any(not _is_blank(later) for later in spans[index + 1 :])
                if _is_blank(span) and kept and "chars" in kept[-1] and words_follow:
                    kept[-1]["chars"].extend(span["chars"])
                    kept[-1]["bbox"] = tuple(pymupdf.Rect(kept[-1]["bbox"]) | span["bbox"])
                else:
                    kept.append(span)
            if tuple(line.get("dir", (1, 0))) == (1, 0):
                _space_apart(kept)
            line["spans"] = kept
    return raw


def _space_apart(spans: list[dict]) -> None:
    for before, after in zip(spans, spans[1:], strict=False):
        left, right = before.get("chars"), after.get("chars")
        if not left or not right or left[-1]["c"].isspace() or right[0]["c"].isspace():
            continue
        x0, y0, x1, y1 = before["bbox"]
        if after["bbox"][0] - x1 < WORD_GAP_EM * before.get("size", 0):
            continue
        left.append(
            {
                "c": " ",
                "origin": (x1, left[-1]["origin"][1]),
                "bbox": (x1, y0, after["bbox"][0], y1),
            }
        )
        before["bbox"] = (x0, y0, after["bbox"][0], y1)


def _extract_with_recognised_text(raw_page: RawPageFitz, **settings):
    engine = raw_page.page_engine
    if engine is not None and settings.get("ocr") == 0 and text_is_only_hidden(engine):
        settings = {**settings, "ocr": 2}
    raw = _plain_extract_raw_dict(raw_page, **settings)
    return drop_fills_drawn_in_pictures(
        keep_word_spaces(level_slight_tilts(split_style_suffixes(raw)))
    )


def clip_ratio(largest_page_area: float) -> float:
    if largest_page_area <= 0:
        return MAX_CLIP_RATIO
    return max(1.0, min(MAX_CLIP_RATIO, math.sqrt(CLIP_PIXEL_BUDGET / largest_page_area)))


def _make_table_with_cached_cells(block: TableBlock, table: Table) -> None:
    table.__class__ = _CachedCellsTable
    try:
        _plain_table_make_docx(block, table)
    finally:
        table.__class__ = Table


_plain_cell_make_docx = Cell.make_docx


def _make_cell_sizing_its_column(cell: Cell, table: Table, indexes: tuple[int, int]) -> None:
    _plain_cell_make_docx(cell, table, indexes)
    if cell and cell.merged_cells[1] == 1:
        table.columns[indexes[1]].width = Pt(cell.bbox.width)


def _ordered_text_style_shapes(shapes: Shapes) -> ElementCollection:
    chosen = {id(shape) for shape in _unordered_text_style_shapes(shapes)}
    return ElementCollection([shape for shape in shapes._instances if id(shape) in chosen])


def hide_page_text_and_images(
    page: pymupdf.Page, rm_text: bool, rm_image: bool
) -> dict[int, bytes]:
    xref_list = [xref for (xref, _name, _invoker, _bbox) in page.get_xobjects()]
    xref_list.extend(page.get_contents())
    image_names = [item[7] for item in page.get_images(full=True)] if rm_image else []

    def hide_text(stream: bytes) -> tuple[bytes, bool]:
        result = stream
        found = False
        for key in ["BT", "Tm", "Td", "2 Tr"]:
            marker = key.encode()
            if marker in stream:
                found = True
                result = result.replace(marker, f"{key} 3 Tr".encode())
        return result, found

    def hide_images(stream: bytes) -> tuple[bytes, bool]:
        result = stream
        found = False
        for name in image_names:
            marker = f"/{name} Do".encode()
            if marker in stream:
                found = True
                result = result.replace(marker, b"")
        return result, found

    document = page.parent
    source: dict[int, bytes] = {}
    for xref in xref_list:
        original = document.xref_stream(xref)
        stream, found_text = hide_text(original) if rm_text else (original, False)
        stream, found_images = hide_images(stream) if rm_image else (stream, False)
        if found_text or found_images:
            document.update_stream(xref, stream)
            source[xref] = original
    return source


class _PageInItsOwnFrame:
    def __init__(self, page: pymupdf.Page) -> None:
        object.__setattr__(self, "_page", page)

    def __getattr__(self, name: str) -> object:
        return getattr(object.__getattribute__(self, "_page"), name)

    @property
    def cropbox(self) -> pymupdf.Rect:
        box = object.__getattribute__(self, "_page").cropbox
        return pymupdf.Rect(0, 0, box.width, box.height)


_plain_extract_images = images_extractor.ImagesExtractor.extract_images


def _inline_image_areas(page: pymupdf.Page, taken: list[pymupdf.Rect]) -> list[pymupdf.Rect]:
    areas: list[pymupdf.Rect] = []
    for info in page.get_image_info(xrefs=True):
        bbox = pymupdf.Rect(info["bbox"]) & page.cropbox
        if info["xref"] or bbox.get_area() <= MIN_IMAGE_AREA:
            continue
        if any(bbox.intersects(other) for other in taken):
            continue
        for area in [area for area in areas if area.intersects(bbox)]:
            bbox |= area
            areas.remove(area)
        areas.append(bbox)
    return areas


def _widened(area: pymupdf.Rect) -> pymupdf.Rect:
    return area + (-COVER_TOLERANCE, -COVER_TOLERANCE, COVER_TOLERANCE, COVER_TOLERANCE)


def shown_area(
    bbox: pymupdf.Rect,
    placements: list[tuple[pymupdf.Rect, pymupdf.Rect]],
    frame: pymupdf.Rect,
) -> pymupdf.Rect:
    inside = [shown for placed, shown in placements if _widened(bbox).contains(placed)]
    if not inside:
        return pymupdf.Rect(bbox) & frame
    area = pymupdf.Rect()
    for shown in inside:
        if not shown.is_empty:
            area = pymupdf.Rect(shown) if area.is_empty else area | shown
    return area & frame


def _extract_images_inside_the_page(extractor, clip_image_res_ratio: float = 3.0):
    page = extractor._page
    extractor._page = _PageInItsOwnFrame(page)
    plain_clip = extractor.clip_page_to_dict

    def clip(bbox=None, rm_image=False, ratio=clip_image_res_ratio):
        image = plain_clip(bbox, rm_image, ratio)
        image[CLIPPED_MARK] = True
        return image

    extractor.clip_page_to_dict = clip
    try:
        frame = extractor._page.cropbox
        placements = visible_images(page)
        images = []
        for image in _plain_extract_images(extractor, clip_image_res_ratio):
            bbox = pymupdf.Rect(image["bbox"])
            shown = shown_area(bbox, placements, frame)
            if shown.is_empty or shown.get_area() <= MIN_IMAGE_AREA:
                continue
            if _widened(shown).contains(bbox):
                images.append(image)
            else:
                images.append(clip(shown, False, clip_image_res_ratio))
        taken = [pymupdf.Rect(image["bbox"]) for image in images]
        for area in _inline_image_areas(extractor._page, taken):
            images.append(clip(area, False, clip_image_res_ratio))
        return images
    finally:
        del extractor.clip_page_to_dict
        extractor._page = page


_plain_to_raw_dict = images_extractor.ImagesExtractor._to_raw_dict


def _to_raw_dict_even_when_empty(image: pymupdf.Pixmap, bbox: pymupdf.Rect) -> dict:
    if image.width < 1 or image.height < 1:
        image = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 1, 1), True)
        image.clear_with(255)
        image.set_alpha(bytes(1))
    return _plain_to_raw_dict(image, bbox)


def _link_target(span: TextSpan) -> str | None:
    if not span.text.strip():
        return None
    for style in span.style:
        if style["type"] == RectType.HYPERLINK.value:
            return style["uri"]
    return None


def _make_span_with_valid_link(span: TextSpan, paragraph) -> None:
    run = paragraph.add_run(span.text)
    span._set_text_format(run)
    if span.char_spacing:
        docx_helpers.set_char_spacing(run, span.char_spacing)
    target = _link_target(span)
    if target is None:
        return
    link = OxmlElement("w:hyperlink")
    link.set(
        qn("r:id"), paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    )
    link.set(qn("w:history"), "1")
    run._r.addprevious(link)
    link.append(run._r)


_plain_parse_pages = Converter.parse_pages


def _parse_pages_keeping_graphics(converter: Converter, **settings) -> Converter:
    _plain_parse_pages(converter, **settings)
    for page in converter.pages:
        if page.finalized:
            add_unused_graphics(page, converter.fitz_doc[page.id], settings["clip_image_res_ratio"])
    return converter


def _family_without_style(name: str) -> str:
    plain = name.split("+")[-1].split("-")[0]
    match = STYLE_SUFFIX.match(plain)
    return match.group(1).strip() if match else plain


def stabilise_pdf2docx() -> None:
    Fonts._normalized_font_name = staticmethod(_family_without_style)
    Converter.parse_pages = _parse_pages_keeping_graphics
    Lines.parse_line_break = parse_line_break_before_list_items
    Cell.make_docx = _make_cell_sizing_its_column
    Line.make_docx = make_line_keeping_word_gap
    TextSpan.make_docx = _make_span_with_valid_link
    TableBlock.make_docx = _make_table_with_cached_cells
    Shapes.text_style_shapes = property(_ordered_text_style_shapes)
    images_extractor.inner_contours = inner_contours
    images_extractor.ImagesExtractor.extract_images = _extract_images_inside_the_page
    images_extractor.ImagesExtractor._to_raw_dict = staticmethod(_to_raw_dict_even_when_empty)
    images_extractor.ImagesExtractor._hide_page_text_and_images = staticmethod(
        hide_page_text_and_images
    )
    RawPageFitz.extract_raw_dict = _extract_with_recognised_text


class ObservedPages(Pages):
    observer: Callable[[], None] | None = None

    def __iter__(self):
        for page in super().__iter__():
            if self.observer is not None and not page.skip_parsing:
                self.observer()
            yield page


class ObservedPage(Page):
    observer: Callable[[str], None] | None = None
    on_made: Callable[[int], None] | None = None

    def parse(self, **settings):
        if self.observer is not None:
            self.observer("parse")
        return super().parse(**settings)

    def make_docx(self, doc):
        if self.observer is not None:
            self.observer("make")
        made = super().make_docx(doc)
        if self.on_made is not None:
            self.on_made(self.id)
        return made


class DocxRun:
    def __init__(self, converter: Converter, progress: Progress, total: int):
        self.converter = converter
        self.progress = progress
        self.total = max(1, total)
        self.analysed = 0
        self.parsed = 0
        self.made = 0
        self.made_pages: set[int] = set()

    def _report(self, band: tuple[float, float], done: int) -> None:
        start, end = band
        self.progress.report(
            start + (end - start) * done / self.total,
            "progress.convertingPages",
            {"current": min(done + 1, self.total), "total": self.total},
        )

    def _stop_if_cancelled(self) -> None:
        if self.progress.cancelled:
            raise ConversionCancelled

    def on_analysed(self) -> None:
        self._stop_if_cancelled()
        self._report(ANALYSIS_BAND, self.analysed)
        self.analysed += 1

    def on_page(self, phase: str) -> None:
        self._stop_if_cancelled()
        if phase == "parse":
            start, end = PARSING_BAND
            self.progress.report(
                start + (end - start) * self.parsed / self.total, "progress.converting"
            )
            self.parsed += 1
            return
        self._report(MAKING_BAND, self.made)
        self.made += 1

    def _observe_pages(self) -> None:
        for page in self.converter.pages:
            page.__class__ = ObservedPage
            page.observer = self.on_page
            page.on_made = self.made_pages.add

    def _parse_here(self, indices: list[int], settings: dict) -> None:
        self.converter.load_pages(pages=indices)
        pages = self.converter.pages
        pages.__class__ = ObservedPages
        pages.observer = self.on_analysed
        try:
            self.converter.parse_document(**settings)
        finally:
            pages.observer = None
        self._stop_if_cancelled()
        self._observe_pages()
        self.converter.parse_pages(**settings)

    def _parse_in_workers(self, indices: list[int], settings: dict, workers: int) -> None:
        parser = ParallelParse(
            self.converter.filename_pdf,
            self.converter.password or None,
            self.on_analysed,
            self._stop_if_cancelled,
        )
        parsed = {page.id: page for page in parser.run(indices, settings, workers)}
        self._stop_if_cancelled()
        self.progress.report(PARSING_BAND[1], "progress.converting")
        self.converter.load_pages(pages=indices)
        pages = self.converter.pages
        pages.reset([parsed.get(page.id, page) for page in list(pages)])

    def run(
        self,
        indices: list[int],
        target: Path,
        image_ratio: float = MAX_CLIP_RATIO,
        right_to_left: bool = False,
        running: RunningText | None = None,
    ) -> list[int]:
        stabilise_pdf2docx()
        settings = {**self.converter.default_settings, "clip_image_res_ratio": image_ratio}
        workers = worker_count(len(indices))
        try:
            if workers > 1:
                self._parse_in_workers(indices, settings, workers)
            else:
                self._parse_here(indices, settings)
            self._stop_if_cancelled()
            self._observe_pages()
            buffer = io.BytesIO()
            self.converter.make_docx(buffer, **settings)
            self._stop_if_cancelled()
        except ConversionCancelled:
            raise OpError(ErrorCode.CANCELLED, "operation cancelled") from None
        data = buffer.getvalue()
        if running is not None and running.lines:
            data = add_running_text(data, running)
        if right_to_left:
            data = mark_right_to_left(data)
        data = embed_fonts(data, self.converter.fitz_doc, indices, self.progress.check_cancelled)
        write_atomically(target, lambda partial: partial.write_bytes(data))
        return [index + 1 for index in indices if index not in self.made_pages]
