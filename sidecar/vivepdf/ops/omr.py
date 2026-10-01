import contextlib
import csv
import html
import io
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal

import numpy as np
import pymupdf
import zxingcpp
from PIL import ImageOps
from pydantic import Field

from vivepdf.ops import _opencv_subset as cv
from vivepdf.ops._document import open_document
from vivepdf.ops._font_unicode import subset_fonts
from vivepdf.ops._image_files import eight_bit, open_picture
from vivepdf.ops._orientation import capped_dpi
from vivepdf.ops._output import prepare_data_output, prepare_output, save_document, write_atomically
from vivepdf.ops._spreadsheet import clean_text, inert_cell, needs_quote_prefix
from vivepdf.ops._story import BASE_CSS, FONT_DIR
from vivepdf.ops.ocr import HIDDEN_TEXT_FONT
from vivepdf.ops.omr_layout import (
    FIDUCIAL,
    MAX_BOOKLETS,
    MAX_DIGITS,
    MAX_OPTIONS,
    MAX_QUESTIONS,
    LetterCase,
    Paper,
    SheetLayout,
    SheetSpec,
    parse_payload,
    sheet_layout,
)
from vivepdf.rpc.errors import ErrorCode, OpError
from vivepdf.rpc.progress import Progress
from vivepdf.rpc.protocol import RpcModel
from vivepdf.rpc.registry import op

FONT_NAME = "VpOmr"
INK = (0.0, 0.0, 0.0)
OUTLINE = (0.35, 0.35, 0.35)
LETTER = (0.6, 0.6, 0.6)
HINT = (0.4, 0.4, 0.4)
BUBBLE_LINE = 0.7
LETTER_SIZE = 6.0
NUMBER_SIZE = 8.0
LABEL_SIZE = 8.5
TITLE_SIZE = 15.0
HINT_SIZE = 8.0
LABEL_LINE_HEIGHT = 1.3
NAME_LINE_MIN = 60.0

ANALYSIS_DPI = 200
PICTURE_MAX_SIDE = 3600
PICTURE_DPI = 150
FIDUCIAL_SEARCH = 40.0
FIDUCIAL_DARK_SHARE = 0.5
FIDUCIAL_MIN_FILL = 0.6
FIDUCIAL_SIZE_RANGE = (0.5, 1.8)
QR_TOLERANCE = 0.2
SAMPLE_SHARE = 0.6
PAPER_WINDOW = 2.4
PAPER_PERCENTILE = 90
MIN_INK_CONTRAST = 60.0
DEFAULT_THRESHOLD = 0.4
UNCLEAR_SHARE = 0.55
STRONG_SHARE = 0.6
PICTURE_EXTENSIONS = frozenset(
    {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".heic", ".heif"}
)

REVIEW_COLORS = {
    "correct": (0.1, 0.6, 0.25),
    "wrong": (0.85, 0.15, 0.15),
    "unclear": (0.95, 0.55, 0.0),
    "blank": (0.1, 0.6, 0.25),
    "void": (0.45, 0.45, 0.45),
}
REVIEW_RING = 1.45
REVIEW_LINE = 1.6
REVIEW_KEY_DASHES = "[2 2] 0"
REVIEW_JPEG_QUALITY = 75

MarkState = Literal["blank", "single", "multiple", "unclear"]
FailureReason = Literal["noSheetCode", "cornersNotFound", "unreadable"]
Verdict = Literal["correct", "wrong", "blank", "void", "unclear"]


class OmrLabels(RpcModel):
    name: str = Field(default="Name", max_length=60)
    student_id: str = Field(default="Student number", max_length=60)
    booklet: str = Field(default="Booklet", max_length=30)
    hint: str = Field(default="", max_length=400)


class OmrSheetParams(RpcModel):
    output: str
    overwrite: bool = False
    paper: Paper = "a4"
    questions: int = Field(ge=1, le=MAX_QUESTIONS)
    options: int = Field(default=5, ge=2, le=MAX_OPTIONS)
    letter_case: LetterCase = "upper"
    id_digits: int = Field(default=0, ge=0, le=MAX_DIGITS)
    booklets: int = Field(default=1, ge=1, le=MAX_BOOKLETS)
    copies: int = Field(default=1, ge=1, le=500)
    title: str = Field(default="", max_length=120)
    labels: OmrLabels = Field(default_factory=OmrLabels)


class OmrSheetResult(RpcModel):
    output: str
    page_count: int
    bytes: int
    capacity: int


class OmrMark(RpcModel):
    state: MarkState
    chosen: list[int]
    fills: list[float]


class OmrScan(RpcModel):
    source: str
    page: int | None
    layout: str
    questions: int
    options: int
    id_digits: int
    booklets: int
    letter_case: LetterCase
    student_id: str
    id_marks: list[OmrMark]
    booklet: OmrMark | None
    marks: list[OmrMark]
    transform: list[float]
    pixel_scale: float


class OmrFailure(RpcModel):
    source: str
    page: int | None
    reason: FailureReason


class OmrReadParams(RpcModel):
    paths: list[str] = Field(min_length=1, max_length=500)
    password: str | None = None
    threshold: float = Field(default=DEFAULT_THRESHOLD, ge=0.15, le=0.85)


class OmrReadResult(RpcModel):
    sheets: list[OmrScan]
    failures: list[OmrFailure]


class OmrReviewQuestion(RpcModel):
    chosen: list[int] = Field(default_factory=list, max_length=MAX_OPTIONS)
    key: list[int] = Field(default_factory=list, max_length=MAX_OPTIONS)
    verdict: Verdict


class OmrReviewItem(RpcModel):
    source: str
    page: int | None
    layout: str
    transform: list[float] = Field(min_length=9, max_length=9)
    pixel_scale: float = Field(gt=0)
    header: str = Field(default="", max_length=300)
    questions: list[OmrReviewQuestion] = Field(max_length=MAX_QUESTIONS)


class OmrReviewParams(RpcModel):
    output: str
    overwrite: bool = False
    password: str | None = None
    items: list[OmrReviewItem] = Field(min_length=1, max_length=2000)


class OmrTable(RpcModel):
    title: str = Field(min_length=1, max_length=31)
    header: list[str] = Field(min_length=1, max_length=300)
    rows: list[list[str | float | None]] = Field(max_length=5000)


class OmrExportParams(RpcModel):
    output: str
    overwrite: bool = False
    format: Literal["csv", "xlsx"] = "xlsx"
    delimiter: Literal[",", ";", "\t"] = ";"
    tables: list[OmrTable] = Field(min_length=1, max_length=5)


class OmrExportResult(RpcModel):
    output: str
    rows: int


def _text_width(font: pymupdf.Font, text: str, size: float) -> float:
    return font.text_length(text, fontsize=size)


def _centred_text(
    page: pymupdf.Page,
    font: pymupdf.Font,
    centre: tuple[float, float],
    text: str,
    size: float,
    color,
) -> None:
    x = centre[0] - _text_width(font, text, size) / 2
    page.insert_text(
        (x, centre[1] + size * 0.36), text, fontname=FONT_NAME, fontsize=size, color=color
    )


def _label_width(font: pymupdf.Font, text: str, size: float) -> float:
    return sum(
        font.text_length(character, fontsize=size) if font.has_glyph(ord(character)) else size
        for character in text
    )


def _css_color(color: tuple[float, float, float]) -> str:
    return "#" + "".join(f"{round(channel * 255):02x}" for channel in color)


def _html_text(
    page: pymupdf.Page,
    box: pymupdf.Rect,
    text: str,
    size: float,
    color: tuple[float, float, float],
    archive: pymupdf.Archive,
) -> None:
    css = BASE_CSS + (
        f"body{{font-size:{size}pt;line-height:{LABEL_LINE_HEIGHT};color:{_css_color(color)};margin:0;}}"
        "p{margin:0;}"
    )
    body = f'<p dir="auto">{html.escape(text)}</p>'
    page.insert_htmlbox(box, body, css=css, archive=archive, scale_low=0)


def _draw_sheet(page: pymupdf.Page, layout: SheetLayout, params: OmrSheetParams) -> None:
    font = pymupdf.Font(fontfile=str(HIDDEN_TEXT_FONT))
    page.insert_font(fontname=FONT_NAME, fontfile=str(HIDDEN_TEXT_FONT))
    shape = page.new_shape()
    half = FIDUCIAL / 2
    for x, y in layout.fiducials:
        shape.draw_rect(pymupdf.Rect(x - half, y - half, x + half, y + half))
    module = layout.qr_area.width / layout.qr_modules.shape[1]
    for row, column in zip(*np.nonzero(layout.qr_modules), strict=True):
        x = layout.qr_area.x0 + float(column) * module
        y = layout.qr_area.y0 + float(row) * module
        shape.draw_rect(pymupdf.Rect(x, y, x + module, y + module))
    shape.finish(color=None, fill=INK, width=0)
    radius = layout.radius
    bubbles = [point for group in layout.question_bubbles for point in group]
    bubbles += [point for group in layout.digit_bubbles for point in group]
    bubbles += layout.booklet_bubbles
    for point in bubbles:
        shape.draw_circle(point, radius)
    shape.finish(color=OUTLINE, fill=None, width=BUBBLE_LINE)
    for box in layout.id_boxes:
        shape.draw_rect(box)
    if layout.id_boxes:
        shape.finish(color=OUTLINE, fill=None, width=BUBBLE_LINE)
    name_line_y = layout.name_box.y1 - 4
    name_width = _label_width(font, params.labels.name, LABEL_SIZE + 1)
    name_start = min(layout.name_box.x0 + name_width + 6, layout.name_box.x1 - NAME_LINE_MIN)
    shape.draw_line((name_start, name_line_y), (layout.name_box.x1, name_line_y))
    shape.finish(color=OUTLINE, width=BUBBLE_LINE)
    shape.commit()

    archive = pymupdf.Archive(str(FONT_DIR))
    if params.title.strip():
        _html_text(page, layout.title_box, params.title.strip(), TITLE_SIZE, INK, archive)
    name_label = pymupdf.Rect(
        layout.name_box.x0,
        name_line_y - (LABEL_SIZE + 1) * LABEL_LINE_HEIGHT - 1,
        name_start - 2,
        name_line_y + 1,
    )
    _html_text(page, name_label, params.labels.name, LABEL_SIZE + 1, INK, archive)
    if layout.id_label is not None:
        _html_text(page, layout.id_label, params.labels.student_id, LABEL_SIZE, INK, archive)
    if layout.booklet_label is not None:
        _html_text(page, layout.booklet_label, params.labels.booklet, LABEL_SIZE, INK, archive)
    if params.labels.hint.strip():
        _html_text(page, layout.hint_box, params.labels.hint.strip(), HINT_SIZE, HINT, archive)
    for group in layout.digit_bubbles:
        for value, point in enumerate(group):
            _centred_text(page, font, point, str(value), LETTER_SIZE, LETTER)
    booklet_letters = [chr(ord("A") + index) for index in range(layout.spec.booklets)]
    for letter, point in zip(booklet_letters, layout.booklet_bubbles, strict=False):
        _centred_text(page, font, point, letter, LETTER_SIZE, LETTER)
    for number, (box, group) in enumerate(
        zip(layout.number_boxes, layout.question_bubbles, strict=True), start=1
    ):
        label = str(number)
        page.insert_text(
            (
                box.x1 - _text_width(font, label, NUMBER_SIZE),
                (box.y0 + box.y1) / 2 + NUMBER_SIZE * 0.36,
            ),
            label,
            fontname=FONT_NAME,
            fontsize=NUMBER_SIZE,
            color=INK,
        )
        for letter, point in zip(layout.letters, group, strict=True):
            _centred_text(page, font, point, letter, LETTER_SIZE, LETTER)


@op("omr.sheet", OmrSheetParams)
def sheet(params: OmrSheetParams, progress: Progress) -> OmrSheetResult:
    target = prepare_output(params.output, [], params.overwrite)
    spec = SheetSpec(
        questions=params.questions,
        options=params.options,
        digits=params.id_digits,
        booklets=params.booklets,
        letter_case=params.letter_case,
        paper=params.paper,
    )
    layout = sheet_layout(spec)
    document = pymupdf.open()
    try:
        page = document.new_page(width=layout.width, height=layout.height)
        _draw_sheet(page, layout, params)
        for copy in range(1, params.copies):
            progress.check_cancelled()
            document.copy_page(0)
            progress.report(
                copy / params.copies,
                "progress.assembling",
                {"current": copy, "total": params.copies},
            )
        with contextlib.suppress(Exception):
            subset_fonts(document, fallback=False)
        saved = save_document(document, target)
        return OmrSheetResult(
            output=saved.output,
            page_count=saved.page_count,
            bytes=saved.bytes,
            capacity=layout.capacity,
        )
    finally:
        document.close()


@dataclass
class _Picture:
    gray: np.ndarray
    pixel_scale: float


def _page_picture(page: pymupdf.Page) -> _Picture:
    pixmap = page.get_pixmap(
        dpi=capped_dpi(page.rect, ANALYSIS_DPI), colorspace=pymupdf.csGRAY, alpha=False
    )
    values = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(pixmap.height, pixmap.stride)[
        :, : pixmap.width
    ]
    return _Picture(gray=np.array(values), pixel_scale=page.rect.width / pixmap.width)


def _file_picture(path: str) -> _Picture:
    with open_picture(path) as opened:
        image = eight_bit(ImageOps.exif_transpose(opened) or opened).convert("L")
    if max(image.size) > PICTURE_MAX_SIDE:
        image.thumbnail((PICTURE_MAX_SIDE, PICTURE_MAX_SIDE))
    return _Picture(gray=np.asarray(image, dtype=np.uint8).copy(), pixel_scale=72.0 / PICTURE_DPI)


@dataclass
class _Code:
    spec: SheetSpec
    payload: str
    corners: np.ndarray


def _find_code(gray: np.ndarray) -> _Code | None:
    for step in (1, 2):
        sample = gray if step == 1 else np.ascontiguousarray(gray[::step, ::step])
        for found in zxingcpp.read_barcodes(sample, formats=zxingcpp.BarcodeFormat.QRCode):
            spec = parse_payload(found.text)
            if spec is None:
                continue
            position = found.position
            corners = np.array(
                [
                    [point.x, point.y]
                    for point in (
                        position.top_left,
                        position.top_right,
                        position.bottom_right,
                        position.bottom_left,
                    )
                ],
                dtype=float,
            )
            return _Code(spec=spec, payload=found.text, corners=corners * step)
    return None


def _affine(source: np.ndarray, target: np.ndarray) -> np.ndarray:
    design = np.hstack([source, np.ones((len(source), 1))])
    solution, *_ = np.linalg.lstsq(design, target, rcond=None)
    matrix = np.eye(3)
    matrix[:2, :] = solution.T
    return matrix


def _homography(source: np.ndarray, target: np.ndarray) -> np.ndarray:
    rows = []
    values = []
    for (x, y), (u, v) in zip(source, target, strict=True):
        rows.append([x, y, 1, 0, 0, 0, -u * x, -u * y])
        rows.append([0, 0, 0, x, y, 1, -v * x, -v * y])
        values.extend([u, v])
    solution = np.linalg.solve(np.array(rows, dtype=float), np.array(values, dtype=float))
    return np.append(solution, 1.0).reshape(3, 3)


def _map(matrix: np.ndarray, points: np.ndarray) -> np.ndarray:
    points = np.atleast_2d(np.asarray(points, dtype=float))
    lifted = np.hstack([points, np.ones((len(points), 1))]) @ matrix.T
    return lifted[:, :2] / lifted[:, 2:3]


def _scale_at(matrix: np.ndarray, point: tuple[float, float], length: float) -> float:
    ends = _map(
        matrix, np.array([point, (point[0] + length, point[1]), (point[0], point[1] + length)])
    )
    return float(
        (np.linalg.norm(ends[1] - ends[0]) + np.linalg.norm(ends[2] - ends[0])) / (2 * length)
    )


def _corner_points(rect: pymupdf.Rect) -> np.ndarray:
    return np.array(
        [[rect.x0, rect.y0], [rect.x1, rect.y0], [rect.x1, rect.y1], [rect.x0, rect.y1]],
        dtype=float,
    )


def _find_fiducial(
    gray: np.ndarray, guess: np.ndarray, pixels_per_point: float
) -> tuple[np.ndarray, float] | None:
    reach = int(FIDUCIAL_SEARCH * pixels_per_point)
    height, width = gray.shape
    x0, y0 = max(0, int(guess[0]) - reach), max(0, int(guess[1]) - reach)
    x1, y1 = min(width, int(guess[0]) + reach), min(height, int(guess[1]) + reach)
    if x1 - x0 < 4 or y1 - y0 < 4:
        return None
    window = gray[y0:y1, x0:x1].astype(float)
    paper = float(np.percentile(window, PAPER_PERCENTILE))
    dark = window < paper * FIDUCIAL_DARK_SHARE
    contours, hierarchy = cv.findContours(
        dark.astype(np.uint8), cv.RETR_TREE, cv.CHAIN_APPROX_SIMPLE
    )
    expected = FIDUCIAL * pixels_per_point
    best: tuple[float, np.ndarray, float] | None = None
    for index, contour in enumerate(contours):
        if hierarchy is None or hierarchy[0, index, 3] != -1:
            continue
        (left, top), (right, bottom) = contour[0][0], contour[1][0]
        box_width, box_height = right - left + 1, bottom - top + 1
        low, high = FIDUCIAL_SIZE_RANGE
        if not (low <= box_width / expected <= high and low <= box_height / expected <= high):
            continue
        block = dark[top : bottom + 1, left : right + 1]
        if block.mean() < FIDUCIAL_MIN_FILL:
            continue
        ys, xs = np.nonzero(block)
        centre = np.array([x0 + left + xs.mean(), y0 + top + ys.mean()])
        distance = float(np.linalg.norm(centre - guess))
        ink = float(np.median(window[top : bottom + 1, left : right + 1][block]))
        if best is None or distance < best[0]:
            best = (distance, centre, ink)
    return (best[1], best[2]) if best else None


@dataclass
class _Alignment:
    matrix: np.ndarray
    ink: float


def _align(gray: np.ndarray, code: _Code, layout: SheetLayout) -> _Alignment | None:
    symbol = _corner_points(layout.qr_symbol)
    rough = _affine(symbol, code.corners)
    pixels_per_point = _scale_at(
        rough, (layout.qr_symbol.x0, layout.qr_symbol.y0), layout.qr_symbol.width
    )
    found: list[tuple[np.ndarray, np.ndarray, float]] = []
    for point in layout.fiducials:
        guess = _map(rough, np.array([point]))[0]
        located = _find_fiducial(gray, guess, pixels_per_point)
        if located is not None:
            found.append((np.array(point, dtype=float), located[0], located[1]))
    if len(found) < 3:
        return None
    source = np.array([item[0] for item in found])
    target = np.array([item[1] for item in found])
    matrix = _homography(source, target) if len(found) == 4 else _affine(source, target)
    error = np.linalg.norm(_map(matrix, symbol) - code.corners, axis=1).max()
    if error > QR_TOLERANCE * layout.qr_symbol.width * pixels_per_point:
        return None
    return _Alignment(matrix=matrix, ink=float(np.median([item[2] for item in found])))


def _fill(
    gray: np.ndarray, alignment: _Alignment, point: tuple[float, float], radius: float
) -> float:
    centre = _map(alignment.matrix, np.array([point]))[0]
    scale = _scale_at(alignment.matrix, point, radius)
    outer = PAPER_WINDOW * radius * scale
    inner = SAMPLE_SHARE * radius * scale
    height, width = gray.shape
    x0, x1 = max(0, int(centre[0] - outer)), min(width, int(centre[0] + outer) + 1)
    y0, y1 = max(0, int(centre[1] - outer)), min(height, int(centre[1] + outer) + 1)
    if x1 <= x0 or y1 <= y0:
        return 0.0
    window = gray[y0:y1, x0:x1].astype(float)
    ys, xs = np.mgrid[y0:y1, x0:x1]
    inside = (xs - centre[0]) ** 2 + (ys - centre[1]) ** 2 <= inner * inner
    if not inside.any():
        return 0.0
    paper = float(np.percentile(window, PAPER_PERCENTILE))
    contrast = max(paper - alignment.ink, MIN_INK_CONTRAST)
    return float(np.clip((paper - window[inside].mean()) / contrast, 0.0, 1.0))


def classify(fills: list[float], threshold: float) -> OmrMark:
    strong = max(fills, default=0.0)
    marked = [
        index
        for index, value in enumerate(fills)
        if value >= threshold and value >= strong * STRONG_SHARE
    ]
    faint = [
        index
        for index, value in enumerate(fills)
        if index not in marked and value >= threshold * UNCLEAR_SHARE
    ]
    rounded = [round(value, 3) for value in fills]
    if not marked and not faint:
        return OmrMark(state="blank", chosen=[], fills=rounded)
    if len(marked) == 1 and not faint:
        return OmrMark(state="single", chosen=marked, fills=rounded)
    if len(marked) >= 2:
        return OmrMark(state="multiple", chosen=marked, fills=rounded)
    candidates = sorted(marked + faint, key=lambda index: -fills[index])
    return OmrMark(state="unclear", chosen=candidates, fills=rounded)


def _read_picture(
    picture: _Picture, source: str, page: int | None, threshold: float
) -> OmrScan | OmrFailure:
    code = _find_code(picture.gray)
    if code is None:
        return OmrFailure(source=source, page=page, reason="noSheetCode")
    layout = sheet_layout(code.spec)
    alignment = _align(picture.gray, code, layout)
    if alignment is None:
        return OmrFailure(source=source, page=page, reason="cornersNotFound")

    def group(points: list[tuple[float, float]]) -> OmrMark:
        return classify(
            [_fill(picture.gray, alignment, point, layout.radius) for point in points], threshold
        )

    id_marks = [group(column) for column in layout.digit_bubbles]
    student_id = "".join(
        str(mark.chosen[0]) if mark.state == "single" else "?" for mark in id_marks
    )
    booklet = group(layout.booklet_bubbles) if layout.booklet_bubbles else None
    return OmrScan(
        source=source,
        page=page,
        layout=code.payload,
        questions=code.spec.questions,
        options=code.spec.options,
        id_digits=code.spec.digits,
        booklets=code.spec.booklets,
        letter_case=code.spec.letter_case,
        student_id=student_id,
        id_marks=id_marks,
        booklet=booklet,
        marks=[group(points) for points in layout.question_bubbles],
        transform=[round(float(value), 9) for value in alignment.matrix.flatten()],
        pixel_scale=picture.pixel_scale,
    )


def _is_picture(path: str) -> bool:
    return Path(path).suffix.lower() in PICTURE_EXTENSIONS


@op("omr.read", OmrReadParams)
def read(params: OmrReadParams, progress: Progress) -> OmrReadResult:
    sheets: list[OmrScan] = []
    failures: list[OmrFailure] = []
    total = len(params.paths)

    def keep(outcome: OmrScan | OmrFailure) -> None:
        (sheets if isinstance(outcome, OmrScan) else failures).append(outcome)

    for index, path in enumerate(params.paths):
        progress.check_cancelled()
        if _is_picture(path):
            try:
                picture = _file_picture(path)
            except (OSError, ValueError, OpError):
                failures.append(OmrFailure(source=path, page=None, reason="unreadable"))
            else:
                keep(_read_picture(picture, path, None, params.threshold))
        else:
            try:
                document = open_document(path, params.password)
            except OpError as error:
                if error.code == ErrorCode.CANCELLED:
                    raise
                failures.append(OmrFailure(source=path, page=None, reason="unreadable"))
                continue
            try:
                for number, page in enumerate(document, start=1):
                    progress.check_cancelled()
                    keep(_read_picture(_page_picture(page), path, number, params.threshold))
                    progress.report(
                        (index + number / max(1, document.page_count)) / total,
                        "progress.analyzing",
                        {"current": index + 1, "total": total},
                    )
            finally:
                document.close()
        progress.report(
            (index + 1) / total, "progress.analyzing", {"current": index + 1, "total": total}
        )
    return OmrReadResult(sheets=sheets, failures=failures)


def _review_page(
    output: pymupdf.Document, item: OmrReviewItem, password: str | None
) -> pymupdf.Page:
    if item.page is not None:
        with open_document(item.source, password) as source:
            if not 1 <= item.page <= source.page_count:
                raise OpError(
                    ErrorCode.INVALID_PARAMS,
                    "page is outside the file",
                    {"reason": "pageOutOfRange"},
                )
            rect = source[item.page - 1].rect
            page = output.new_page(width=rect.width, height=rect.height)
            page.show_pdf_page(page.rect, source, item.page - 1)
            return page
    picture = _file_picture(item.source)
    height, width = picture.gray.shape
    with open_picture(item.source) as opened:
        image = (ImageOps.exif_transpose(opened) or opened).convert("RGB")
    image = image.resize((width, height))
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=REVIEW_JPEG_QUALITY)
    page = output.new_page(width=width * item.pixel_scale, height=height * item.pixel_scale)
    page.insert_image(page.rect, stream=buffer.getvalue())
    return page


def _draw_review(page: pymupdf.Page, item: OmrReviewItem) -> None:
    spec = parse_payload(item.layout)
    if spec is None:
        raise OpError(
            ErrorCode.INVALID_PARAMS, "unknown sheet layout", {"reason": "omrLayoutUnknown"}
        )
    layout = sheet_layout(spec)
    matrix = np.array(item.transform, dtype=float).reshape(3, 3)
    shape = page.new_shape()
    for question, points in zip(item.questions, layout.question_bubbles, strict=False):
        color = REVIEW_COLORS[question.verdict]
        for option in question.chosen:
            if 0 <= option < len(points):
                centre = _map(matrix, np.array([points[option]]))[0] * item.pixel_scale
                radius = (
                    layout.radius
                    * REVIEW_RING
                    * _scale_at(matrix, points[option], layout.radius)
                    * item.pixel_scale
                )
                shape.draw_circle(pymupdf.Point(*centre), radius)
                shape.finish(color=color, width=REVIEW_LINE)
        if question.verdict in ("wrong", "blank", "unclear"):
            for option in question.key:
                if 0 <= option < len(points) and option not in question.chosen:
                    centre = _map(matrix, np.array([points[option]]))[0] * item.pixel_scale
                    radius = (
                        layout.radius
                        * REVIEW_RING
                        * _scale_at(matrix, points[option], layout.radius)
                        * item.pixel_scale
                    )
                    shape.draw_circle(pymupdf.Point(*centre), radius)
                    shape.finish(
                        color=REVIEW_COLORS["correct"],
                        width=REVIEW_LINE / 2,
                        dashes=REVIEW_KEY_DASHES,
                    )
    shape.commit()
    if item.header.strip():
        box = pymupdf.Rect(8, 6, page.rect.width - 8, 24)
        page.draw_rect(box, color=None, fill=(1, 1, 1), fill_opacity=0.85)
        _html_text(
            page,
            box + (4, 2, -4, 0),
            item.header.strip(),
            10,
            REVIEW_COLORS["wrong"],
            pymupdf.Archive(str(FONT_DIR)),
        )


@op("omr.review", OmrReviewParams)
def review(params: OmrReviewParams, progress: Progress) -> OmrSheetResult:
    target = prepare_output(params.output, [item.source for item in params.items], params.overwrite)
    output = pymupdf.open()
    try:
        for index, item in enumerate(params.items):
            progress.check_cancelled()
            try:
                page = _review_page(output, item, params.password)
            except (OSError, RuntimeError, ValueError) as error:
                raise OpError(
                    ErrorCode.INVALID_PARAMS,
                    f"cannot open {Path(item.source).name}",
                    {"reason": "imageUnreadable", "path": item.source},
                ) from error
            _draw_review(page, item)
            progress.report(
                (index + 1) / len(params.items),
                "progress.assembling",
                {"current": index + 1, "total": len(params.items)},
            )
        with contextlib.suppress(Exception):
            subset_fonts(output, fallback=False)
        saved = save_document(output, target)
        return OmrSheetResult(
            output=saved.output, page_count=saved.page_count, bytes=saved.bytes, capacity=0
        )
    finally:
        output.close()


def _cell_text(value: str | float | None) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _write_csv(partial: Path, table: OmrTable, delimiter: str) -> None:
    with partial.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=delimiter)
        writer.writerow([inert_cell(clean_text(text)) for text in table.header])
        for row in table.rows:
            writer.writerow([inert_cell(clean_text(_cell_text(value))) for value in row])


def _xlsx_cell(sheet: Any, value: str | float | None) -> Any:
    from openpyxl.cell import WriteOnlyCell

    if isinstance(value, (int, float)):
        return WriteOnlyCell(sheet, value=int(value) if float(value).is_integer() else value)
    text = clean_text(value or "")
    cell = WriteOnlyCell(sheet, value=text)
    if needs_quote_prefix(text):
        cell.data_type = "s"
        cell.quotePrefix = True
    return cell


def _write_xlsx(partial: Path, tables: list[OmrTable]) -> None:
    from openpyxl import Workbook

    workbook = Workbook(write_only=True)
    used: set[str] = set()
    for table in tables:
        title = clean_text(table.title).translate(str.maketrans("", "", "[]:*?/\\"))[:31] or "Sheet"
        while title in used:
            title = f"{title[:28]} {len(used)}"
        used.add(title)
        sheet = workbook.create_sheet(title)
        sheet.freeze_panes = "A2"
        sheet.append([_xlsx_cell(sheet, text) for text in table.header])
        for row in table.rows:
            sheet.append([_xlsx_cell(sheet, value) for value in row])
    workbook.save(str(partial))


@op("omr.export", OmrExportParams)
def export(params: OmrExportParams, _progress: Progress) -> OmrExportResult:
    target = prepare_data_output(params.output, params.format, params.overwrite)
    if params.format == "csv":
        write_atomically(
            target, lambda partial: _write_csv(partial, params.tables[0], params.delimiter)
        )
    else:
        write_atomically(target, lambda partial: _write_xlsx(partial, params.tables))
    return OmrExportResult(output=str(target), rows=len(params.tables[0].rows))
