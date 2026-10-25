import io
import math
import tempfile
from pathlib import Path
from typing import Literal

import pymupdf
from pydantic import Field

from vivepdf.ops._document import open_document
from vivepdf.ops._hidden_text import hidden_origins, widget_origins
from vivepdf.ops._output import write_atomically
from vivepdf.ops._pptx_embed import embed_slide_fonts
from vivepdf.ops._pptx_fonts import slide_font_names
from vivepdf.ops._pptx_layout import (
    SlideFrame,
    add_line_box,
    line_boxes,
    page_notes,
    text_free_pixmap,
)
from vivepdf.ops._pptx_links import SlideJumps, add_link_area, apply_jumps, page_links
from vivepdf.ops._ranges import parse_page_ranges
from vivepdf.ops.convert import PagedFileResult, PdfSourceParams, _prepare_file
from vivepdf.rpc.progress import Progress
from vivepdf.rpc.registry import op


class PptxParams(PdfSourceParams):
    dpi: int = Field(default=150, ge=50, le=300)
    mode: Literal["editable", "image"] = "editable"


EMU_PER_POINT = 12700
MAX_SLIDE_EMU = 51206400
MIN_SLIDE_EMU = 914400
MAX_SLIDE_PIXELS = 50_000_000
NOTES_LIMIT = 20000


def slide_dpi(rect: pymupdf.Rect, dpi: int) -> int:
    area = abs(rect.width * rect.height)
    if area <= 0:
        return dpi
    ceiling = int(72 * math.sqrt(MAX_SLIDE_PIXELS / area))
    return max(1, min(dpi, ceiling))


def slide_size(rect: pymupdf.Rect) -> tuple[int, int]:
    width = rect.width * EMU_PER_POINT
    height = rect.height * EMU_PER_POINT
    scale = max(1.0, MIN_SLIDE_EMU / max(1.0, min(width, height)))
    scale = min(scale, MAX_SLIDE_EMU / max(width, height))
    return round(width * scale), round(height * scale)


def fitted_box(
    page: pymupdf.Rect, slide_width: int, slide_height: int
) -> tuple[int, int, int, int]:
    ratio = min(slide_width / page.width, slide_height / page.height)
    width = round(page.width * ratio)
    height = round(page.height * ratio)
    return (slide_width - width) // 2, (slide_height - height) // 2, width, height


@op("convert.to_pptx", PptxParams)
def to_pptx(params: PptxParams, progress: Progress) -> PagedFileResult:
    from pptx import Presentation

    target = _prepare_file(params.output, [params.path], params.overwrite, ".pptx")
    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]
    with open_document(params.path, params.password) as document:
        indices = parse_page_ranges(params.pages, document.page_count)
        slide_width, slide_height = slide_size(document[indices[0]].rect)
        presentation.slide_width = slide_width
        presentation.slide_height = slide_height
        editable = params.mode == "editable"
        jumps = SlideJumps([])
        slides = {}
        names = slide_font_names(document) if editable else None
        with tempfile.TemporaryDirectory(prefix="vivepdf-pptx-") as temp_dir:
            for position, index in enumerate(indices):
                progress.check_cancelled()
                page = document[index]
                image_path = Path(temp_dir) / f"page-{index + 1}.png"
                dpi = slide_dpi(page.rect, params.dpi)
                hidden = hidden_origins(page) | widget_origins(page)
                links = page_links(page)
                boxes = line_boxes(page, hidden, links) if editable else []
                if editable:
                    areas = [box.area for box in boxes]
                    text_free_pixmap(document, index, dpi, areas).save(str(image_path))
                else:
                    page.get_pixmap(dpi=dpi).save(str(image_path))
                slide = presentation.slides.add_slide(blank_layout)
                slides[index] = slide
                left, top, width, height = fitted_box(page.rect, slide_width, slide_height)
                slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)
                frame = SlideFrame(left=left, top=top, ratio=width / page.rect.width)
                for box in boxes:
                    add_line_box(slide, box, frame, jumps, names)
                linked = {span["link"] for box in boxes for span in box.spans if span.get("link")}
                for link in links:
                    if link not in linked:
                        add_link_area(slide, link, frame, jumps)
                if not editable:
                    text = page_notes(page, hidden)
                    if text:
                        slide.notes_slide.notes_text_frame.text = text[:NOTES_LIMIT]
                progress.report(
                    position / len(indices),
                    "progress.convertingPages",
                    {"current": position + 1, "total": len(indices)},
                )
            apply_jumps(jumps, slides)
            buffer = io.BytesIO()
            presentation.save(buffer)
            data = buffer.getvalue()
            if editable:
                data = embed_slide_fonts(data, document, indices, names, progress.check_cancelled)
            write_atomically(target, lambda partial: partial.write_bytes(data))
        return PagedFileResult(
            output=str(target), bytes=target.stat().st_size, page_count=len(indices)
        )
