import csv
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

import pymupdf
from pydantic import Field

from vivepdf.ops._cell_text import page_fragments, reads_upright, table_spans, table_texts
from vivepdf.ops._document import open_document
from vivepdf.ops._ocr_layer import Rule
from vivepdf.ops._output import write_atomically
from vivepdf.ops._ranges import parse_page_ranges
from vivepdf.ops._split_tables import joined_across_pages
from vivepdf.ops._spreadsheet import CellTyping, clean_text, inert_cell, needs_quote_prefix
from vivepdf.ops.convert import FileResult, TextSourceParams, _prepare_file, recognise_first
from vivepdf.rpc.progress import Progress
from vivepdf.rpc.registry import op


class TablesResult(FileResult):
    table_count: int
    text_pages: list[int] = Field(default_factory=list)
    textless_pages: list[int] = Field(default_factory=list)
    ocr_pages: list[int] = Field(default_factory=list)


class XlsxParams(TextSourceParams):
    sheets: Literal["table", "page", "single"] = "table"
    format: Literal["xlsx", "csv"] = "xlsx"
    delimiter: Literal[",", ";", "\t"] = ","
    borderless: bool = False
    page_label: str = Field(default="S", max_length=24)
    table_label: str = Field(default="T", max_length=24)
    text_label: str = Field(default="Text", max_length=24)


@dataclass(slots=True)
class _Table:
    page: int
    number: int
    rows: list[list[str | None]]
    spans: list[tuple[int, int, int, int]] = field(default_factory=list)
    area: pymupdf.Rect = field(default_factory=pymupdf.Rect)
    last_page: int = -1
    edges: list[float] = field(default_factory=list)
    bold_head: bool = False
    top: bool = False
    bottom: bool = False
    lines_above: int = 0
    lines_below: int = 0


@dataclass(frozen=True, slots=True)
class SheetLabels:
    page: str = "S"
    table: str = "T"
    text: str = "Text"


Placed = list[tuple[int, _Table]]

SHEET_NAME_LIMIT = 31
SHEET_NAME_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")
CELL_TEXT_LIMIT = 32767
COLUMN_WIDTH_RANGE = (8, 60)


def spreadsheet_cell(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, str):
        return inert_cell(clean_text(value))
    return value


def _page_lines(page: pymupdf.Page) -> list[str]:
    return [line for line in page.get_text().splitlines() if line.strip()]


def _inside_any(box: pymupdf.Rect, areas: list[pymupdf.Rect]) -> bool:
    middle = pymupdf.Point((box.x0 + box.x1) / 2, (box.y0 + box.y1) / 2)
    return any(middle in area for area in areas)


def _lines_outside(page: pymupdf.Page, areas: list[pymupdf.Rect]) -> list[tuple[pymupdf.Rect, str]]:
    lines: list[tuple[pymupdf.Rect, str]] = []
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            text = "".join(span["text"] for span in line["spans"])
            box = pymupdf.Rect(line["bbox"]) * page.rotation_matrix
            if text.strip() and not _inside_any(box, areas):
                lines.append((box, text))
    return lines


PAGE_END_SHARE = 0.25


def _mark_page_ends(
    tables: list[_Table], outside: list[tuple[pymupdf.Rect, str]], height: float
) -> None:
    first = min(tables, key=lambda table: table.area.y0)
    last = max(tables, key=lambda table: table.area.y1)
    first.top = first.area.y0 <= height * PAGE_END_SHARE
    last.bottom = last.area.y1 >= height * (1 - PAGE_END_SHARE)
    first.lines_above = sum(1 for box, _text in outside if box.y1 <= first.area.y0 + RULE_SLACK)
    last.lines_below = sum(1 for box, _text in outside if box.y0 >= last.area.y1 - RULE_SLACK)


def _page_span(first: int, last: int, label: str) -> str:
    return f"{label}{first}" if first == last else f"{label}{first}-{label}{last}"


def _placed(tables: list[_Table]) -> Placed:
    placed: Placed = []
    offset = 0
    for position, table in enumerate(tables):
        if position:
            offset += 1
        placed.append((offset, table))
        offset += len(table.rows)
    return placed


def sheet_plan(tables: list[_Table], sheets: str, labels: SheetLabels) -> list[tuple[str, Placed]]:
    if sheets == "single":
        title = _page_span(tables[0].page + 1, tables[-1].last_page + 1, labels.page)
        return [(title, _placed(tables))]
    if sheets == "page":
        pages: dict[int, list[_Table]] = {}
        for table in tables:
            pages.setdefault(table.page, []).append(table)
        return [(f"{labels.page}{page + 1}", _placed(group)) for page, group in pages.items()]
    return [
        (
            f"{_page_span(table.page + 1, table.last_page + 1, labels.page)}"
            f"-{labels.table}{table.number + 1}",
            [(0, table)],
        )
        for table in tables
    ]


def sheet_title(title: str, taken: set[str]) -> str:
    base = SHEET_NAME_FORBIDDEN.sub("", clean_text(title)).strip().strip("'").strip() or "1"
    candidate = base[:SHEET_NAME_LIMIT]
    number = 2
    while candidate.casefold() in taken:
        suffix = f" ({number})"
        candidate = base[: SHEET_NAME_LIMIT - len(suffix)] + suffix
        number += 1
    taken.add(candidate.casefold())
    return candidate


def _text_cell(sheet, text: str):
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment

    value = clean_text(text)[:CELL_TEXT_LIMIT]
    cell = WriteOnlyCell(sheet, value=value)
    if needs_quote_prefix(value):
        cell.data_type = "s"
        cell.quotePrefix = True
    if "\n" in value or len(value) > COLUMN_WIDTH_RANGE[1]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    return cell


def _table_cell(sheet, text: str | None, typing: CellTyping):
    from openpyxl.cell import WriteOnlyCell

    if text is None or not text.strip():
        return None
    typed = typing.typed(clean_text(text))
    if typed is None:
        return _text_cell(sheet, text)
    cell = WriteOnlyCell(sheet, value=typed.value)
    if typed.number_format:
        cell.number_format = typed.number_format
    return cell


def _column_widths(rows: list[list[str | None]]) -> list[int]:
    low, high = COLUMN_WIDTH_RANGE
    widths: list[int] = []
    for row in rows:
        for column, text in enumerate(row):
            longest = max((len(line) for line in (text or "").splitlines()), default=0)
            if column == len(widths):
                widths.append(low)
            widths[column] = max(widths[column], min(high, longest + 2))
    return widths


def _set_widths(sheet, rows: list[list[str | None]]) -> None:
    from openpyxl.utils import get_column_letter

    for column, width in enumerate(_column_widths(rows), start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width


HEADER_MIN_CELLS = 2


def has_header(table: _Table, typing: CellTyping) -> bool:
    if len(table.rows) < 2:
        return False
    head = [clean_text(text) for text in table.rows[0] if text and text.strip()]
    if len(head) < HEADER_MIN_CELLS or any(typing.typed(text) for text in head):
        return False
    return table.bold_head or any(
        typing.typed(clean_text(text)) is not None
        for row in table.rows[1:]
        for text in row
        if text and text.strip()
    )


def _header_cells(cells: list) -> list:
    from openpyxl.styles import Font

    for cell in cells:
        if cell is not None:
            cell.font = Font(bold=True)
    return cells


def _write_tables_sheet(sheet, placed: Placed, typing: CellTyping) -> None:
    from openpyxl.worksheet.cell_range import CellRange

    _set_widths(sheet, [row for _offset, table in placed for row in table.rows])
    headed = [has_header(table, typing) for _offset, table in placed]
    if len(placed) == 1 and headed[0]:
        sheet.freeze_panes = "A2"
    written = 0
    for (offset, table), header in zip(placed, headed, strict=True):
        while written < offset:
            sheet.append([])
            written += 1
        for position, row in enumerate(table.rows):
            cells = [_table_cell(sheet, text, typing) for text in row]
            sheet.append(_header_cells(cells) if header and not position else cells)
            written += 1
        for first_row, first_column, last_row, last_column in table.spans:
            sheet.merged_cells.add(
                CellRange(
                    min_row=offset + first_row + 1,
                    min_col=first_column + 1,
                    max_row=offset + last_row + 1,
                    max_col=last_column + 1,
                )
            )


def _write_lines_sheet(sheet, lines: list[tuple[int, str]]) -> None:
    _set_widths(sheet, [[None, text] for _page, text in lines])
    for page, text in lines:
        sheet.append([page, _text_cell(sheet, text)])


def _write_workbook(
    partial: Path,
    tables: list[_Table],
    lines: list[tuple[int, str]],
    sheets: str,
    labels: SheetLabels,
    first_page: int,
) -> None:
    from openpyxl import Workbook

    workbook = Workbook(write_only=True)
    taken: set[str] = set()
    if tables:
        typing = CellTyping(text for table in tables for row in table.rows for text in row if text)
        for title, placed in sheet_plan(tables, sheets, labels):
            _write_tables_sheet(
                workbook.create_sheet(title=sheet_title(title, taken)), placed, typing
            )
        if lines:
            _write_lines_sheet(workbook.create_sheet(title=sheet_title(labels.text, taken)), lines)
    else:
        pages = [page for page, _text in lines] or [first_page]
        title = _page_span(min(pages), max(pages), labels.page)
        _write_lines_sheet(workbook.create_sheet(title=sheet_title(title, taken)), lines)
    workbook.save(str(partial))


def _write_csv(
    partial: Path, tables: list[_Table], lines: list[tuple[int, str]], delimiter: str
) -> None:
    rows: list[list[object]] = []
    if tables:
        for position, table in enumerate(tables):
            if position:
                rows.append([])
            rows.extend([spreadsheet_cell(text) for text in row] for row in table.rows)
    else:
        rows = [[page, spreadsheet_cell(text)] for page, text in lines]
    with partial.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=delimiter)
        for row in rows:
            writer.writerow(row)


BORDERLESS_MIN_ROWS = 2
BORDERLESS_MIN_COLUMNS = 2


def _filled_rows(table: object) -> int:
    return sum(
        1
        for row in table.extract()
        if sum(1 for cell in row if cell and str(cell).strip()) >= BORDERLESS_MIN_COLUMNS
    )


PHRASE_GAP = 0.5
EDGE_SLACK = 0.5


def _row_words(page: pymupdf.Page, table: object) -> list[list[pymupdf.Rect]]:
    area = pymupdf.Rect(table.bbox)
    words = [
        pymupdf.Rect(word[:4]) * page.rotation_matrix
        for word in page.get_text("words", clip=area * page.derotation_matrix)
    ]
    rows = []
    for row in table.rows:
        box = pymupdf.Rect(row.bbox)
        rows.append([word for word in words if box.y0 <= (word.y0 + word.y1) / 2 <= box.y1])
    return rows


def _splits_phrases(rows: list[list[pymupdf.Rect]], left: float, edge: float, right: float) -> bool:
    joined = False
    for words in rows:
        before = [word for word in words if left <= (word.x0 + word.x1) / 2 < edge]
        after = [word for word in words if edge <= (word.x0 + word.x1) / 2 < right]
        if not after:
            continue
        if not before:
            return False
        last = max(before, key=lambda word: word.x1)
        first = min(after, key=lambda word: word.x0)
        if first.x0 - last.x1 > PHRASE_GAP * max(last.height, first.height):
            return False
        joined = True
    return joined


def _with_whole_phrases(page: pymupdf.Page, table: object) -> object:
    edges = _column_edges(table)
    rows = _row_words(page, table)
    kept = [edges[0]] + [
        edge
        for left, edge, right in zip(edges, edges[1:], edges[2:], strict=False)
        if not _splits_phrases(rows, left, edge, right)
    ]
    kept.append(edges[-1])
    if len(kept) == len(edges):
        return table
    clip = pymupdf.Rect(table.bbox) + (-EDGE_SLACK, -EDGE_SLACK, EDGE_SLACK, EDGE_SLACK)
    rebuilt = page.find_tables(
        use_layout=False,
        clip=clip,
        vertical_strategy="explicit",
        vertical_lines=kept,
        horizontal_strategy="text",
    ).tables
    for candidate in rebuilt:
        if candidate.col_count == len(kept) - 1:
            return candidate
    return table


def borderless_tables(page: pymupdf.Page) -> list:
    candidates = page.find_tables(use_layout=False, strategy="text").tables
    return [
        _with_whole_phrases(page, table)
        for table in candidates
        if table.col_count >= BORDERLESS_MIN_COLUMNS and _filled_rows(table) >= BORDERLESS_MIN_ROWS
    ]


RULE_THICKNESS = 2.0
RULE_SLACK = 2.0
RULE_GAP_FACTOR = 3.0
RULE_MIN_GAP = 36.0


def _rule_spans(rect: pymupdf.Rect, area: pymupdf.Rect) -> bool:
    return (
        rect.height <= RULE_THICKNESS
        and rect.x0 <= area.x0 + RULE_SLACK
        and rect.x1 >= area.x1 - RULE_SLACK
    )


def _ruled_bottom(page: pymupdf.Page, area: pymupdf.Rect) -> float | None:
    rules = sorted(
        rect.y1
        for rect in (
            pymupdf.Rect(drawing["rect"]) * page.rotation_matrix for drawing in page.get_drawings()
        )
        if rect.y0 >= area.y1 - RULE_SLACK and _rule_spans(rect, area)
    )
    bottom = area.y1
    gap = max(area.height * RULE_GAP_FACTOR, RULE_MIN_GAP)
    for edge in rules:
        if edge - bottom > gap:
            break
        bottom = max(bottom, edge)
    return bottom if bottom > area.y1 + RULE_SLACK else None


def _with_ruled_body(page: pymupdf.Page, table: object) -> object:
    if table.row_count > 1:
        return table
    area = pymupdf.Rect(table.bbox)
    bottom = _ruled_bottom(page, area)
    if bottom is None:
        return table
    edges = sorted(
        {cell[0] for cell in table.cells if cell} | {cell[2] for cell in table.cells if cell}
    )
    clip = pymupdf.Rect(page.rect.x0, area.y0 - RULE_SLACK, page.rect.x1, bottom + RULE_SLACK)
    extended = page.find_tables(
        use_layout=False,
        clip=clip,
        vertical_strategy="explicit",
        vertical_lines=edges,
        horizontal_strategy="lines",
    ).tables
    for candidate in extended:
        if candidate.col_count == table.col_count and candidate.row_count > table.row_count:
            return candidate
    return table


OVERLAP_TOLERANCE = 1.0
FRAME_SHARE = 0.5


def _overlaps(cells: list[pymupdf.Rect]) -> bool:
    active: list[pymupdf.Rect] = []
    for cell in sorted(cells, key=lambda rect: rect.y0):
        active = [other for other in active if other.y1 - cell.y0 > OVERLAP_TOLERANCE]
        for other in active:
            shared = cell & other
            if shared.width > OVERLAP_TOLERANCE and shared.height > OVERLAP_TOLERANCE:
                return True
        active.append(cell)
    return False


def is_page_frame(page: pymupdf.Page, table: object) -> bool:
    cells = [pymupdf.Rect(cell) for cell in table.cells if cell]
    page_area = abs(page.rect)
    if any(abs(cell) > page_area * FRAME_SHARE for cell in cells):
        return True
    return _overlaps(cells)


def _without_empty_columns(rows: list[list[str | None]]) -> list[list[str | None]]:
    width = max((len(row) for row in rows), default=0)
    kept = [
        column
        for column in range(width)
        if any(column < len(row) and row[column] and row[column].strip() for row in rows)
    ]
    return [[row[column] if column < len(row) else None for column in kept] for row in rows]


def _page_tables(
    page: pymupdf.Page, index: int, borderless: bool, scanned: list[Rule] | None = None
) -> list[_Table]:
    if scanned is None:
        finder = page.find_tables(use_layout=False)
    else:
        finder = page.find_tables(use_layout=False, add_lines=scanned) if scanned else None
    found = (
        [_with_ruled_body(page, table) for table in finder.tables if not is_page_frame(page, table)]
        if finder is not None
        else []
    )
    borderless = borderless or scanned is not None
    unruled = not found and borderless
    if unruled:
        found = borderless_tables(page)
    fragments = page_fragments(page) if found else []
    found = [table for table in found if reads_upright(page, fragments, pymupdf.Rect(table.bbox))]
    tables: list[_Table] = []
    for number, table in enumerate(found):
        reach = (page.rect.x0, page.rect.x1) if unruled else None
        rows = [
            row
            for row in table_texts(table, fragments, reach)
            if not unruled or any(text and text.strip() for text in row)
        ]
        if unruled:
            rows = _without_empty_columns(rows)
        spans = [] if unruled else table_spans(table)
        tables.append(
            _Table(
                page=index,
                number=number,
                rows=rows,
                spans=spans,
                area=pymupdf.Rect(table.bbox),
                last_page=index,
                edges=_column_edges(table),
                bold_head=_bold_head(page, table),
            )
        )
    return tables


def _column_edges(table: object) -> list[float]:
    return sorted({edge for cell in table.cells if cell for edge in (cell[0], cell[2])})


def _bold_head(page: pymupdf.Page, table: object) -> bool:
    if not table.rows:
        return False
    spans = [
        span
        for block in page.get_text(
            "dict", clip=pymupdf.Rect(table.rows[0].bbox) * page.derotation_matrix
        )["blocks"]
        for line in block.get("lines", [])
        for span in line["spans"]
        if span["text"].strip()
    ]
    return bool(spans) and all(
        span["flags"] & pymupdf.TEXT_FONT_BOLD or "bold" in span["font"].lower() for span in spans
    )


@op("convert.to_xlsx", XlsxParams)
def to_xlsx(params: XlsxParams, progress: Progress) -> TablesResult:
    suffix = ".csv" if params.format == "csv" else ".xlsx"
    target = _prepare_file(params.output, [params.path], params.overwrite, suffix)
    tables: list[_Table] = []
    text_by_page: dict[int, list[str]] = {}
    around_tables: dict[int, list[str]] = {}
    textless: list[int] = []
    with open_document(params.path, params.password) as document:
        indices = parse_page_ranges(params.pages, document.page_count)
        rules: dict[int, list[Rule]] = {}
        recognised, progress = recognise_first(document, indices, params, progress, rules)
        for position, index in enumerate(indices):
            progress.check_cancelled()
            page = document[index]
            scanned = rules.get(index) if index + 1 in recognised else None
            found = _page_tables(page, index, params.borderless, scanned)
            tables.extend(found)
            if found:
                outside = _lines_outside(page, [table.area for table in found])
                _mark_page_ends(found, outside, page.cropbox.height)
                if outside:
                    around_tables[index + 1] = [text for _box, text in outside]
            else:
                page_lines = _page_lines(page)
                if page_lines:
                    text_by_page[index + 1] = page_lines
                else:
                    textless.append(index + 1)
            progress.report(
                position / len(indices),
                "progress.convertingPages",
                {"current": position + 1, "total": len(indices)},
            )
    if params.format == "csv" or params.sheets != "page":
        tables = joined_across_pages(tables, indices)
    keeps_text = not tables or params.format == "xlsx"
    lines = [
        (page, text)
        for page, page_lines in sorted({**text_by_page, **around_tables}.items())
        for text in page_lines
        if keeps_text
    ]
    labels = SheetLabels(params.page_label, params.table_label, params.text_label)
    if params.format == "csv":
        write_atomically(
            target, lambda partial: _write_csv(partial, tables, lines, params.delimiter)
        )
    else:
        write_atomically(
            target,
            lambda partial: _write_workbook(
                partial, tables, lines, params.sheets, labels, indices[0] + 1
            ),
        )
    return TablesResult(
        output=str(target),
        bytes=target.stat().st_size,
        table_count=len(tables),
        text_pages=sorted(page for page in text_by_page if keeps_text),
        textless_pages=textless,
        ocr_pages=recognised,
    )
