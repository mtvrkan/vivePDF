import io
import re
from collections.abc import Iterator

RTL_LETTER = re.compile(r"[\u0590-\u08ff\ufb1d-\ufdff\ufe70-\ufefc]")
BIDI_SUCCESSORS = (
    "w:adjustRightInd",
    "w:snapToGrid",
    "w:spacing",
    "w:ind",
    "w:contextualSpacing",
    "w:mirrorIndents",
    "w:suppressOverlap",
    "w:jc",
    "w:textDirection",
    "w:textAlignment",
    "w:textboxTightWrap",
    "w:outlineLvl",
    "w:divId",
    "w:cnfStyle",
    "w:rPr",
    "w:sectPr",
    "w:pPrChange",
)
MIRRORED_ALIGNMENT = {"left": "right", "right": "left", "start": "end", "end": "start"}


def has_right_to_left(text: str) -> bool:
    return RTL_LETTER.search(text) is not None


def mostly_right_to_left(text: str) -> bool:
    right_to_left = len(RTL_LETTER.findall(text))
    letters = sum(1 for character in text if character.isalpha())
    return right_to_left > 0 and right_to_left * 2 >= letters


def _paragraphs(container) -> Iterator:
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _paragraphs(cell)


def mark_paragraph(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = paragraph._p.get_or_add_pPr()
    if properties.find(qn("w:bidi")) is not None:
        return
    properties.insert_element_before(OxmlElement("w:bidi"), *BIDI_SUCCESSORS)
    alignment = properties.find(qn("w:jc"))
    if alignment is not None:
        value = alignment.get(qn("w:val"))
        if value in MIRRORED_ALIGNMENT:
            alignment.set(qn("w:val"), MIRRORED_ALIGNMENT[value])
    for run in paragraph.runs:
        run.font.rtl = True


def mark_right_to_left(data: bytes) -> bytes:
    from docx import Document

    document = Document(io.BytesIO(data))
    marked = False
    for paragraph in _paragraphs(document):
        if not mostly_right_to_left(paragraph.text):
            continue
        mark_paragraph(paragraph)
        marked = True
    if not marked:
        return data
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
