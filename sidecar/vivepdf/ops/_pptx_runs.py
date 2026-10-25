import copy
import re
from collections.abc import Callable

from vivepdf.ops._pptx_layout import (
    FILLED_FLAG,
    STROKED_FLAG,
    LineBox,
    SlideFrame,
    _hollow,
    slide_text,
)
from vivepdf.ops._pptx_links import SlideJumps, link_run

EMU_PER_POINT = 12700
MIN_FONT_POINTS = 1.0
MAX_FONT_POINTS = 4000.0
BOLD_FLAG = 16
ITALIC_FLAG = 2
SPACE_RUNS = re.compile(r"( +)")
SUBSET_PREFIX = re.compile(r"^[A-Z]{6}\+")
FONT_SUFFIXES = ("PSMT", "MT", "PS")
FONT_STYLE = re.compile(
    r"(?:(?:semi|demi|extra|ultra)?(?:bold|light|black|heavy|thin|medium|regular|roman|book"
    r"|italic|ital|oblique|condensed|narrow|it|bd|bi|lt|md|rg|bk))+(?:psmt|mt|ps)?",
    re.IGNORECASE,
)
FONT_FAMILIES = {
    "Helvetica": "Arial",
    "Times": "Times New Roman",
    "TimesNewRoman": "Times New Roman",
    "Courier": "Courier New",
    "CourierNew": "Courier New",
}
DEFAULT_FAMILY = "Arial"


def font_family(name: str) -> str:
    parts = re.split(r"[-,]", SUBSET_PREFIX.sub("", name or ""))
    kept = parts[:1]
    for part in parts[1:]:
        if not part or FONT_STYLE.fullmatch(part):
            break
        kept.append(part)
    base = " ".join(kept)
    for suffix in FONT_SUFFIXES:
        if base.endswith(suffix) and len(base) > len(suffix):
            base = base[: -len(suffix)]
            break
    return FONT_FAMILIES.get(base, base) or DEFAULT_FAMILY


def _font_points(size: float, ratio: float) -> float:
    return min(MAX_FONT_POINTS, max(MIN_FONT_POINTS, round(size * ratio / EMU_PER_POINT, 1)))


def _outline_run(run, colour, width: float | None, ratio: float) -> None:
    from lxml import etree
    from pptx.dml.fill import FillFormat
    from pptx.oxml.ns import qn

    properties = run._r.get_or_add_rPr()
    line = etree.SubElement(properties, qn("a:ln"))
    properties.insert(0, line)
    if width:
        line.set("w", str(max(1, round(width * ratio))))
    stroke = FillFormat.from_fill_parent(line)
    stroke.solid()
    stroke.fore_color.rgb = colour
    FillFormat.from_fill_parent(properties).background()


def add_line_box(
    slide,
    box: LineBox,
    frame: SlideFrame,
    jumps: SlideJumps | None = None,
    names: Callable[[str], str] = font_family,
    language: str | None = None,
) -> None:
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.util import Emu

    shape = slide.shapes.add_textbox(
        Emu(round(frame.left + box.left * frame.ratio)),
        Emu(round(frame.top + box.top * frame.ratio)),
        Emu(max(1, round(box.width * frame.ratio))),
        Emu(max(1, round(box.height * frame.ratio))),
    )
    text_frame = shape.text_frame
    text_frame.word_wrap = False
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.margin_left = text_frame.margin_right = 0
    text_frame.margin_top = text_frame.margin_bottom = 0
    text_frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = text_frame.paragraphs[0]
    for span in box.spans:
        text = slide_text(span["text"])
        if not text:
            continue
        extra = span.get("space_extra", 0.0)
        parts = SPACE_RUNS.split(text) if extra else [text]
        spacing = str(round(extra * frame.ratio / EMU_PER_POINT * 100))
        styled = None
        for part in parts:
            if not part:
                continue
            if styled is None:
                styled = _span_run(paragraph, span, part, frame, jumps, names, language)._r
                element = styled
            else:
                previous, element = element, copy.deepcopy(styled)
                element.text = part
                previous.addnext(element)
            properties = element.get_or_add_rPr()
            if part.startswith(" "):
                properties.set("spc", spacing)
            else:
                properties.attrib.pop("spc", None)
    if box.angle:
        shape.rotation = box.angle


def _span_run(
    paragraph,
    span: dict,
    text: str,
    frame: SlideFrame,
    jumps: SlideJumps | None,
    names: Callable[[str], str],
    language: str | None,
):
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    run = paragraph.add_run()
    run.text = text
    if language:
        run._r.get_or_add_rPr().set("lang", language)
    font = run.font
    font.size = Pt(_font_points(span["size"], frame.ratio))
    font.name = names(span["font"])
    strokes = span.get("char_flags", FILLED_FLAG)
    fake_bold = strokes & FILLED_FLAG and strokes & STROKED_FLAG
    font.bold = bool(span["flags"] & BOLD_FLAG or fake_bold)
    font.italic = bool(span["flags"] & ITALIC_FLAG)
    colour = RGBColor.from_string(f"{span['color'] & 0xFFFFFF:06X}")
    if _hollow(span):
        _outline_run(run, colour, span.get("outline"), frame.ratio)
    else:
        font.color.rgb = colour
    if span.get("link") is not None and jumps is not None:
        link_run(run, span["link"], jumps)
    return run
