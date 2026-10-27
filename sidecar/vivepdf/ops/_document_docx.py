import contextlib
import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field, replace
from datetime import UTC, datetime

import docx
import pymupdf
from docx.document import Document as WordDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.part import Part
from docx.oxml.ns import qn
from docx.oxml.shape import CT_Inline
from lxml import etree
from lxml import html as lxml_html
from PIL import Image

from vivepdf.ops._document_docx_front import cover, front_matter, outline
from vivepdf.ops._document_docx_layout import (
    LIST_STEP,
    POINTS_PER_MM,
    ListNumbers,
    SectionSpec,
    blank_furniture,
    content_width,
    finish_settings,
    font_family,
    furniture,
    section_properties,
    setup_styles,
)
from vivepdf.ops._document_docx_parts import (
    CommentBook,
    CommentEntry,
    anchor_comment,
    comment_parts,
    footnote,
    footnote_reference,
    footnotes_part,
)
from vivepdf.ops._document_docx_style import ParagraphLook, RunLook
from vivepdf.ops._document_docx_tables import table
from vivepdf.ops._document_docx_xml import (
    EMU_PER_POINT,
    break_run,
    element,
    field_runs,
    tab_run,
    text_run,
)
from vivepdf.ops._document_equations import Equations
from vivepdf.ops._document_notes import note_bodies
from vivepdf.ops._studio_document import DEFAULT_FONT, _decoded
from vivepdf.ops._studio_document_models import (
    MAX_COMMENT_AUTHOR,
    StudioDocumentDocxParams,
    StudioDocumentThread,
)

SPACES = re.compile(r"[ \t\r\n\f]+")
STYLE_ENTRY = re.compile(r"([a-z-]+)\s*:\s*([^;]+)")
POINTS = re.compile(r"^(-?\d+(?:\.\d+)?)pt$")
EMS = re.compile(r"^(-?\d+(?:\.\d+)?)em$")
NUMBER = re.compile(r"^\d+(?:\.\d+)?$")
COLOUR = re.compile(r"^#[0-9a-f]{6}$", re.IGNORECASE)
FONT_TOKEN = re.compile(r"^f(\d{1,2})$")
IMAGE_TOKEN = re.compile(r"^vpimg-(\d+)$")
EQUATION_TOKEN = re.compile(r"^vpeq-(\d+)$")
NOTE_ID = re.compile(r"^fnref-(\d+)$")
REVISION_KINDS = frozenset({"ins", "del"})
REVISION_BASE = 100000
CAPTION_LABEL = re.compile(r"^(.*?)(\d+)(\D*)$", re.DOTALL)
HEADINGS = {f"h{level}": level for level in range(1, 7)}
BLOCK_TAGS = frozenset(
    {"p", "ul", "ol", "li", "blockquote", "pre", "hr", "table", "div", "section", *HEADINGS}
)
DROPPED_TAGS = frozenset({"script", "style", "template", "head", "title"})
SEQUENCES = {"figure": "Figure", "table": "Table"}
CODE_FAMILY = "Courier New"
QUOTE_STEP = 10.0
QUOTE_GAP = 8.0
QUOTE_COLOUR = "#444444"
QUOTE_BORDER = "#cccccc"
CODE_SHADE = "#f3f4f6"
CODE_BORDER = "#d1d5db"
RULE_COLOUR = "#bbbbbb"
REFERENCE_HANG = 36.0
NOTE_SCALE = 0.8
NOTE_PICTURE_HEIGHT = 1.2
EQUATION_DPI = 600
EQUATION_MAX_PIXELS = 6000
NUMBER_RESERVE_EM = 3.0
PX_TO_PT = 0.75
TAB_SPACES = "    "
NBSP = " "
NAMESPACE_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@dataclass
class Piece:
    look: RunLook
    text: str | None = None
    runs: list = field(default_factory=list)
    breaks: bool = False


class Inline:
    def __init__(self, preserve: bool = False) -> None:
        self.pieces: list[Piece] = []
        self.preserve = preserve

    def text(self, text: str, look: RunLook) -> None:
        if text:
            self.pieces.append(Piece(look, text))

    def raw(self, runs: list, look: RunLook, breaks: bool = False) -> None:
        if runs:
            self.pieces.append(Piece(look, None, runs, breaks))

    def normalised(self) -> list[Piece]:
        found: list[Piece] = []
        space = True
        for piece in self.pieces:
            if piece.text is None:
                if piece.breaks and found and found[-1].text is not None:
                    found[-1] = replace(found[-1], text=found[-1].text.rstrip(" "))
                found.append(piece)
                space = piece.breaks
                continue
            text = piece.text
            if not self.preserve:
                text = SPACES.sub(" ", text)
                if space and text.startswith(" "):
                    text = text[1:]
                if not text:
                    continue
                space = text.endswith(" ")
            if found and found[-1].text is not None and found[-1].look == piece.look:
                found[-1] = replace(found[-1], text=found[-1].text + text)
            else:
                found.append(Piece(piece.look, text))
        while found and found[-1].text is not None and not self.preserve:
            trimmed = found[-1].text.rstrip(" ")
            if trimmed:
                found[-1] = replace(found[-1], text=trimmed)
                break
            found.pop()
        return [piece for piece in found if piece.text is None or piece.text]

    def blank(self) -> bool:
        pieces = self.normalised()
        return all(piece.text is not None and not piece.text.strip(NBSP + " ") for piece in pieces)


@dataclass(frozen=True)
class Context:
    width: float
    part: Part
    level: int = -1
    tasks: bool = False
    quote: int = 0
    cell: bool = False
    header: bool = False
    note: bool = False
    run: RunLook = field(default_factory=RunLook)
    marker: list = field(default_factory=list, compare=False)

    def with_(self, **changes) -> "Context":
        return replace(self, **changes)


def _styles(node: etree._Element) -> dict[str, str]:
    return {
        key.strip().lower(): value.strip()
        for key, value in STYLE_ENTRY.findall((node.get("style") or "").lower())
    }


def _points(value: str | None) -> float | None:
    match = POINTS.match(value or "")
    return float(match.group(1)) if match else None


def _ems(value: str | None) -> float | None:
    match = EMS.match(value or "")
    return float(match.group(1)) if match else None


def _classes(node: etree._Element) -> set[str]:
    return set((node.get("class") or "").split())


def _tag(node: etree._Element) -> str:
    return node.tag.lower() if isinstance(node.tag, str) else ""


def _xsd_date(value: str) -> str:
    with contextlib.suppress(ValueError):
        moment = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
        if moment.tzinfo is None:
            moment = moment.replace(tzinfo=UTC)
        return moment.astimezone(UTC).replace(microsecond=0).strftime("%Y-%m-%dT%H:%M:%SZ")
    return ""


def _initials(author: str) -> str:
    letters = "".join(word[0] for word in author.split() if word)
    return (letters[:3] or "?").upper()


class DocxWriter:
    def __init__(
        self,
        document: WordDocument,
        content: StudioDocumentDocxParams,
        check_cancelled: Callable[[], None],
    ) -> None:
        self.document = document
        self.content = content
        self.settings = content.settings
        self.size = content.settings.font_size
        self.check_cancelled = check_cancelled
        self.families = [font_family(font_id) for font_id in content.fonts or [DEFAULT_FONT]]
        self.equations = Equations(content.equations)
        self.equation_pictures: dict[tuple[int, int], bytes] = {}
        self.pictures: dict[int, tuple[bytes, int, int]] = {}
        self.numbers = ListNumbers(document)
        self.notes = note_bodies(content.footnotes, content.footnote_html, content.images)
        self.notes_part: Part | None = None
        self.written_notes: set[int] = set()
        self.first: dict[str, etree._Element] = {}
        self.last: dict[str, etree._Element] = {}
        self.revisions = 0
        self.shapes = 0
        self.body: list[etree._Element] = []
        self.section = SectionSpec(columns=max(1, self.settings.columns))
        self.section_used = False
        self.page_break = False

    def column_width(self, columns: int | None = None) -> float:
        count = columns or self.section.columns
        gap = self.settings.column_gap_mm * POINTS_PER_MM
        return max(36.0, (content_width(self.settings) - gap * (count - 1)) / count)

    def add(self, container: etree._Element | None, node: etree._Element) -> None:
        if container is None:
            self.body.append(node)
            self.section_used = True
        else:
            container.append(node)

    def paragraph_node(
        self, look: ParagraphLook, container: etree._Element | None
    ) -> etree._Element:
        if self.page_break and container is None:
            look = look.with_(page_break_before=True)
            self.page_break = False
        node = element("w:p", children=[look.element()])
        self.add(container, node)
        return node

    def flush_page_break(self) -> None:
        if self.page_break:
            self.page_break = False
            self.add(None, element("w:p", children=[ParagraphLook().element(), break_run("page")]))

    def close_section(self, following: SectionSpec) -> None:
        self.flush_page_break()
        if not self.section_used:
            current = self.section
            self.section = replace(
                following,
                headers=following.headers or current.headers,
                title_page=following.title_page or current.title_page,
                number_format=following.number_format or current.number_format,
                number_start=current.number_start
                if following.number_start is None
                else following.number_start,
            )
            return
        holder = self.body[-1] if self.body and self.body[-1].tag == qn("w:p") else None
        if holder is None or holder.find(qn("w:pPr")) is None:
            holder = element("w:p", children=[ParagraphLook().element()])
            self.body.append(holder)
        properties = holder.find(qn("w:pPr"))
        if properties.find(qn("w:sectPr")) is not None:
            holder = element("w:p", children=[ParagraphLook().element()])
            self.body.append(holder)
            properties = holder.find(qn("w:pPr"))
        properties.append(section_properties(self.section, self.settings))
        self.section = following
        self.section_used = False

    def relate_link(self, part: Part, href: str) -> str:
        return part.relate_to(href, RT.HYPERLINK, is_external=True)

    def picture_run(
        self, part: Part, payload: bytes, width: float, height: float, alt: str, look: RunLook
    ) -> etree._Element:
        image_part = part.package.get_or_add_image_part(io.BytesIO(payload))
        relation = part.relate_to(image_part, RT.IMAGE)
        self.shapes += 1
        inline = CT_Inline.new_pic_inline(
            self.shapes,
            relation,
            image_part.filename,
            round(width * EMU_PER_POINT),
            round(height * EMU_PER_POINT),
        )
        if alt:
            inline.docPr.set("descr", alt)
        return element("w:r", children=[look.element(), element("w:drawing", children=[inline])])

    def picture(self, index: int) -> tuple[bytes, int, int] | None:
        if index in self.pictures:
            return self.pictures[index]
        if not 0 <= index < len(self.content.images):
            return None
        payload = _decoded(self.content.images[index], index)
        with Image.open(io.BytesIO(payload)) as opened:
            width, height = opened.size
        self.pictures[index] = (payload, width, height)
        return self.pictures[index]

    def equation_png(self, index: int, size: float) -> bytes | None:
        key = (index, round(size * 10))
        if key in self.equation_pictures:
            return self.equation_pictures[key]
        drawing = self.equations.picture(index)
        found = self.equations.size(index)
        if drawing is None or found is None:
            return None
        zoom = size * EQUATION_DPI / 72 / 100
        page = drawing[0]
        longest = max(page.rect.width, page.rect.height) * zoom
        if longest > EQUATION_MAX_PIXELS:
            zoom *= EQUATION_MAX_PIXELS / longest
        try:
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=True)
            payload = pixmap.tobytes("png")
        except Exception:  # noqa: BLE001
            return None
        self.equation_pictures[key] = payload
        return payload

    def image(self, node: etree._Element, builder: Inline, look: RunLook, context: Context) -> None:
        source = (node.get("src") or "").strip()
        alt = (node.get("alt") or "").strip()
        size = look.size or (self.size * NOTE_SCALE if context.note else self.size)
        equation = EQUATION_TOKEN.match(source)
        if equation:
            index = int(equation.group(1))
            found = self.equations.size(index)
            payload = self.equation_png(index, size) if found else None
            if payload is None:
                builder.text(alt, look)
                return
            depth = _ems(_styles(node).get("vertical-align")) or 0.0
            run = self.picture_run(
                context.part,
                payload,
                found[0] * size,
                found[1] * size,
                alt,
                look.with_(position=depth * size),
            )
            builder.raw([run], look)
            return
        match = IMAGE_TOKEN.match(source)
        found = self.picture(int(match.group(1))) if match else None
        if found is None:
            return
        payload, pixels_wide, pixels_high = found
        ratio = pixels_high / pixels_wide if pixels_wide else 1.0
        if context.note:
            height = size * NOTE_PICTURE_HEIGHT
            width = height / ratio if ratio else height
        else:
            width = _points(_styles(node).get("width")) or pixels_wide * PX_TO_PT
            width = min(width, context.width)
            height = width * ratio
        builder.raw([self.picture_run(context.part, payload, width, height, alt, look)], look)

    def note_reference(self, number: int, builder: Inline, look: RunLook) -> None:
        if not 1 <= number <= len(self.notes) or number in self.written_notes:
            return
        if self.notes_part is None:
            self.notes_part = footnotes_part(self.document.part)
        self.written_notes.add(number)
        paragraph = element("w:p", children=[ParagraphLook(style="FootnoteText").element()])
        holder = Inline()
        context = Context(width=self.column_width(), part=self.notes_part, note=True)
        with contextlib.suppress(etree.ParserError, ValueError):
            root = lxml_html.fragment_fromstring(self.notes[number - 1] or " ", create_parent="div")
            self.inline_children(root, holder, RunLook(), context)
        self.emit(holder, paragraph, context)
        self.notes_part.root.append(footnote(number, paragraph))
        builder.raw([footnote_reference(number)], look)

    def run_look(self, node: etree._Element, tag: str, look: RunLook) -> RunLook:
        if tag in ("strong", "b"):
            return look.with_(bold=True)
        if tag in ("em", "i"):
            return look.with_(italic=True)
        if tag == "u":
            return look.with_(underline=True)
        if tag in ("s", "strike", "del"):
            return look.with_(strike=True)
        if tag == "code":
            return look.with_(family=CODE_FAMILY)
        if tag == "sub":
            return look.with_(script="subscript")
        if tag == "sup":
            return look.with_(script="superscript")
        if tag != "span":
            return look
        styles = _styles(node)
        colour = styles.get("color", "")
        if COLOUR.match(colour):
            look = look.with_(colour=colour)
        background = styles.get("background-color", "")
        if COLOUR.match(background):
            look = look.with_(background=background)
        family = FONT_TOKEN.match(styles.get("font-family", ""))
        if family and int(family.group(1)) < len(self.families):
            look = look.with_(family=self.families[int(family.group(1))])
        size = _points(styles.get("font-size"))
        if size and 1 <= size <= 400:
            look = look.with_(size=size)
        comment = (node.get("data-comment") or "").strip()
        if comment:
            look = look.with_(comments=look.comments | {comment})
        revision = (node.get("data-revision") or "").strip()
        if revision in REVISION_KINDS:
            author = (node.get("data-author") or "").strip()[:MAX_COMMENT_AUTHOR] or "?"
            look = look.with_(revision=(revision, author, _xsd_date(node.get("data-date") or "")))
        return look

    def inline_children(
        self, node: etree._Element, builder: Inline, look: RunLook, context: Context
    ) -> None:
        if node.text:
            builder.text(node.text, look)
        for child in node:
            self.inline(child, builder, look, context)
            if child.tail:
                builder.text(child.tail, look)

    def inline(
        self, node: etree._Element, builder: Inline, look: RunLook, context: Context
    ) -> None:
        tag = _tag(node)
        if not tag or tag in DROPPED_TAGS:
            return
        if tag == "br":
            builder.raw([break_run()], look, breaks=True)
            return
        if tag == "img":
            self.image(node, builder, look, context)
            return
        if tag == "sup" and "fn-ref" in _classes(node):
            match = NOTE_ID.match(node.get("id") or "")
            if match and not context.note:
                self.note_reference(int(match.group(1)), builder, look)
            return
        if tag == "a":
            href = (node.get("href") or "").strip()
            if href and not href.startswith("#"):
                look = look.with_(link=href, colour=self.settings.accent, underline=True)
        else:
            look = self.run_look(node, tag, look)
        self.inline_children(node, builder, look, context)

    def emit(self, builder: Inline, paragraph: etree._Element, context: Context) -> None:
        group: etree._Element | None = None
        link = ""
        for piece in builder.normalised():
            runs = (
                piece.runs if piece.text is None else [text_run(piece.text, piece.look.element())]
            )
            if piece.look.revision:
                runs = [self.revision(piece.look.revision, runs)]
            if piece.look.link:
                if group is None or link != piece.look.link:
                    relation = self.relate_link(context.part, piece.look.link)
                    group = element("w:hyperlink", {"r:id": relation, "w:history": "1"})
                    paragraph.append(group)
                    link = piece.look.link
                for run in runs:
                    group.append(run)
                first, last = group, group
            else:
                group, link = None, ""
                for run in runs:
                    paragraph.append(run)
                first, last = runs[0], runs[-1]
            for comment in piece.look.comments:
                self.first.setdefault(comment, first)
                self.last[comment] = last

    def revision(self, revision: tuple[str, str, str], runs: list) -> etree._Element:
        kind, author, date = revision
        self.revisions += 1
        attributes = {"w:id": str(REVISION_BASE + self.revisions), "w:author": author}
        if date:
            attributes["w:date"] = date
        wrapper = element(f"w:{kind}", attributes)
        for run in runs:
            if kind == "del":
                for node in run.iter(qn("w:t")):
                    node.tag = qn("w:delText")
                for node in run.iter(qn("w:instrText")):
                    node.tag = qn("w:delInstrText")
            wrapper.append(run)
        return wrapper

    def paragraph_look(self, node: etree._Element, base: ParagraphLook) -> ParagraphLook:
        styles = _styles(node)
        look = base
        align = styles.get("text-align", "")
        if align in ("left", "center", "right", "justify"):
            look = look.with_(align=align)
        line = styles.get("line-height", "")
        if NUMBER.match(line):
            look = look.with_(line=min(4.0, max(0.5, float(line))))
        for key, name in (("margin-top", "before"), ("margin-bottom", "after")):
            value = _points(styles.get(key))
            if value is not None and value >= 0:
                look = look.with_(**{name: value})
        left = _points(styles.get("margin-left"))
        if left:
            look = look.with_(left=(look.left or 0) + left)
        right = _points(styles.get("margin-right"))
        if right:
            look = look.with_(right=right)
        indent = _points(styles.get("text-indent"))
        if indent is not None and look.numbering is None:
            look = look.with_(first_line=indent)
        if (node.get("dir") or "").lower() == "rtl":
            look = look.with_(bidi=True)
        return look

    def base_look(self, context: Context) -> ParagraphLook:
        look = ParagraphLook()
        quote_left = QUOTE_STEP * context.quote + (QUOTE_GAP if context.quote else 0)
        if context.quote:
            look = look.with_(
                borders=(("left", 2.0, QUOTE_BORDER, QUOTE_GAP),),
                left=quote_left,
                after=self.size * 0.6,
            )
        if context.level >= 0:
            step = LIST_STEP * (context.level + (0 if context.tasks else 1))
            look = look.with_(left=quote_left + step, after=2.0)
            if context.marker:
                look = look.with_(numbering=context.marker.pop(), first_line=-LIST_STEP)
        if context.cell:
            look = look.with_(after=0, before=0)
        return look

    def text_block(
        self,
        node: etree._Element,
        container: etree._Element | None,
        context: Context,
        look: ParagraphLook,
        run: RunLook | None = None,
    ) -> None:
        look = self.paragraph_look(node, look)
        run_look = run or context.run
        if look.bidi:
            run_look = run_look.with_(rtl=True)
        builder = Inline()
        self.inline_children(node, builder, run_look, context)
        paragraph = self.paragraph_node(look, container)
        self.emit(builder, paragraph, context)

    def caption(
        self, node: etree._Element, container: etree._Element | None, context: Context
    ) -> None:
        look = self.paragraph_look(node, ParagraphLook(style="Caption", keep_next=False))
        run_look = context.run.with_(rtl=look.bidi)
        builder = Inline()
        sequence = SEQUENCES.get(node.get("data-caption") or "")
        label = node[0] if len(node) and _tag(node[0]) == "strong" else None
        if node.text:
            builder.text(node.text, run_look)
        for child in node:
            if child is label and sequence:
                match = CAPTION_LABEL.match(child.text_content())
                bold = run_look.with_(bold=True)
                if match:
                    builder.text(match.group(1), bold)
                    shown = match.group(2)
                    builder.raw(
                        field_runs(f"SEQ {sequence} \\* ARABIC", shown, bold.element()), bold
                    )
                    builder.text(match.group(3), bold)
                else:
                    builder.text(child.text_content(), bold)
            else:
                self.inline(child, builder, run_look, context)
            if child.tail:
                builder.text(child.tail, run_look)
        self.emit(builder, self.paragraph_node(look, container), context)

    def display_equation(
        self, node: etree._Element, container: etree._Element | None, context: Context
    ) -> None:
        index_text = node.get("data-equation") or ""
        found = self.equations.size(int(index_text)) if index_text.isdigit() else None
        payload = self.equation_png(int(index_text), self.size) if found else None
        if payload is None:
            self.text_block(node, container, context, self.base_look(context))
            return
        width = context.width - (self.base_look(context).left or 0)
        wide, high = found[0] * self.size, found[1] * self.size
        room = width - 2 * NUMBER_RESERVE_EM * self.size
        if room > 0 and wide > room:
            wide, high = room, high * room / wide
        look = self.base_look(context).with_(
            tabs=(("center", width / 2), ("right", width)),
            before=self.size * 0.5,
            after=self.size * 0.5,
            line=1.0,
            first_line=0,
        )
        paragraph = self.paragraph_node(look, container)
        builder = Inline(preserve=True)
        builder.raw([tab_run()], context.run)
        latex = node.get("data-latex") or ""
        builder.raw(
            [self.picture_run(context.part, payload, wide, high, latex, context.run)], context.run
        )
        builder.raw([tab_run()], context.run)
        builder.text(SPACES.sub(" ", node.text_content()).strip(), context.run)
        self.emit(builder, paragraph, context)

    def code_block(
        self, node: etree._Element, container: etree._Element | None, context: Context
    ) -> None:
        look = self.base_look(context).with_(
            background=CODE_SHADE,
            borders=(("left", 2.0, CODE_BORDER, 4.0),),
            line=1.0,
            after=8.0,
        )
        run = context.run.with_(family=CODE_FAMILY, size=self.size * 0.88)
        builder = Inline(preserve=True)
        lines = node.text_content().replace("\t", TAB_SPACES).rstrip("\n").split("\n")
        for index, line in enumerate(lines):
            if index:
                builder.raw([break_run()], run)
            builder.text(line, run)
        self.emit(builder, self.paragraph_node(look, container), context)

    def rule(self, container: etree._Element | None, context: Context) -> None:
        look = self.base_look(context).with_(
            borders=(("bottom", 0.75, RULE_COLOUR, 1.0),), before=8.0, after=8.0, line=1.0
        )
        self.paragraph_node(look, container)

    def list_block(
        self, node: etree._Element, container: etree._Element | None, context: Context, tag: str
    ) -> None:
        level = min(context.level + 1, 8)
        tasks = "tasks" in _classes(node)
        number = None
        if not tasks:
            if tag == "ol":
                start = node.get("start") or "1"
                number = self.numbers.ordered(level, int(start) if start.isdigit() else 1)
            else:
                number = self.numbers.bullet
        for item in node:
            if _tag(item) != "li":
                continue
            marker = [(level, number)] if number is not None else []
            self.blocks(item, container, context.with_(level=level, tasks=tasks, marker=marker))

    def table_block(
        self, node: etree._Element, container: etree._Element | None, context: Context
    ) -> None:
        if container is None:
            self.flush_page_break()

        def fill(cell: etree._Element, box: etree._Element, header: bool) -> None:
            inner = context.with_(
                cell=True,
                header=header,
                level=-1,
                quote=0,
                marker=[],
                width=max(36.0, context.width / 2),
                run=context.run.with_(bold=True) if header else context.run,
            )
            self.blocks(cell, box, inner)

        built = table(node, context.width - (self.base_look(context).left or 0), fill)
        if built is None:
            return
        self.add(container, built)
        self.add(container, element("w:p", children=[ParagraphLook(after=0, line=1.0).element()]))

    def section_break(self, node: etree._Element) -> None:
        columns = node.get("data-columns") or "1"
        count = min(3, max(1, int(columns))) if columns.isdigit() else 1
        continuous = node.get("data-kind") == "continuous"
        self.close_section(SectionSpec(columns=count, continuous=continuous))

    def main_break(self) -> None:
        self.close_section(
            SectionSpec(
                columns=self.section.columns,
                number_format="decimal",
                number_start=self.settings.page_number_start,
            )
        )

    def block(
        self, node: etree._Element, container: etree._Element | None, context: Context
    ) -> None:
        tag = _tag(node)
        classes = _classes(node)
        if tag in HEADINGS:
            level = HEADINGS[tag]
            look = self.base_look(context).with_(style=f"Heading{level}")
            self.text_block(node, container, context, look)
        elif tag == "p" and "caption" in classes:
            self.caption(node, container, context)
        elif tag == "p" and "eq" in classes:
            self.display_equation(node, container, context)
        elif tag == "p" and "ref" in classes:
            look = self.base_look(context).with_(left=REFERENCE_HANG, first_line=-REFERENCE_HANG)
            self.text_block(node, container, context, look)
        elif tag == "p":
            self.text_block(node, container, context, self.base_look(context))
        elif tag in ("ul", "ol"):
            self.list_block(node, container, context, tag)
        elif tag == "blockquote":
            inner = context.with_(
                quote=context.quote + 1, run=context.run.with_(colour=QUOTE_COLOUR)
            )
            self.blocks(node, container, inner)
        elif tag == "pre":
            self.code_block(node, container, context)
        elif tag == "hr":
            self.rule(container, context)
        elif tag == "table":
            self.table_block(node, container, context)
        elif tag == "div" and "page-break" in classes and container is None:
            if node.get("data-main") == "true":
                self.main_break()
            else:
                self.flush_page_break()
                self.page_break = True
        elif tag == "div" and "section-break" in classes and container is None:
            self.section_break(node)
        else:
            self.blocks(node, container, context)

    def blocks(
        self, node: etree._Element, container: etree._Element | None, context: Context
    ) -> None:
        loose = Inline()
        loose_used = False

        def flush() -> None:
            nonlocal loose, loose_used
            if loose_used and not loose.blank():
                paragraph = self.paragraph_node(self.base_look(context), container)
                self.emit(loose, paragraph, context)
            loose, loose_used = Inline(), False

        if node.text and node.text.strip():
            loose.text(node.text, context.run)
            loose_used = True
        for child in node:
            tag = _tag(child)
            if tag in BLOCK_TAGS:
                flush()
                if container is None:
                    self.check_cancelled()
                self.block(child, container, context)
            elif tag and tag not in DROPPED_TAGS:
                self.inline(child, loose, context.run, context)
                loose_used = True
            if child.tail and child.tail.strip():
                loose.text(child.tail, context.run)
                loose_used = True
        flush()

    def front(self, root: lxml_html.HtmlElement) -> None:
        parts = front_matter(self.settings, root)
        for index, part in enumerate(parts):
            if index and part:
                _break_before(part[0])
            for paragraph in part:
                self.add(None, paragraph)
        self.page_break = bool(parts)
        if parts and self.settings.columns > 1:
            self.page_break = False
            self.close_section(SectionSpec(columns=self.settings.columns))

    def write(self, root: lxml_html.HtmlElement, title: str) -> None:
        settings = self.settings
        has_main = any(
            node.get("data-main") == "true" and "page-break" in _classes(node)
            for node in root.iter("div")
        )
        offset = 1 if settings.cover else 0
        if settings.cover:
            heading = self.families[1] if len(self.families) > 1 else self.families[0]
            for paragraph in cover(settings, title, heading) or [element("w:p")]:
                self.add(None, paragraph)
            self.section = SectionSpec(centred=True, headers=blank_furniture(self.document))
        has_front = bool(front_matter(settings, root))
        columns = 1 if has_front else max(1, settings.columns)
        first = SectionSpec(
            columns=columns,
            number_format="lowerRoman" if has_main else "decimal",
            number_start=offset + 1 if has_main else settings.page_number_start + offset,
            title_page=not settings.furniture_on_first,
            headers=furniture(self.document, settings, not settings.furniture_on_first),
        )
        if settings.cover:
            self.close_section(first)
        else:
            self.section = first
        self.front(root)
        context = Context(width=self.column_width(), part=self.document.part)
        self.blocks(root, None, context)
        self.flush_page_break()
        if not self.body or self.body[-1].tag != qn("w:p"):
            self.body.append(element("w:p", children=[ParagraphLook().element()]))

    def finish(self) -> etree._Element:
        body = self.document.element.body
        for child in list(body):
            body.remove(child)
        for node in self.body:
            body.append(node)
        body.append(section_properties(self.section, self.settings))
        return body

    def comments(self, threads: list[StudioDocumentThread]) -> None:
        book = CommentBook()
        for thread in threads:
            first, last = self.first.get(thread.id), self.last.get(thread.id)
            if first is None or last is None:
                continue
            author = thread.author or "?"
            root_number = book.add(
                CommentEntry(
                    author,
                    _initials(author),
                    _xsd_date(thread.created_at),
                    thread.text,
                    thread.resolved,
                )
            )
            numbers = [root_number]
            for reply in thread.replies:
                name = reply.author or "?"
                numbers.append(
                    book.add(
                        CommentEntry(
                            name,
                            _initials(name),
                            _xsd_date(reply.created_at),
                            reply.text,
                            thread.resolved,
                            root_number,
                        )
                    )
                )
            for number in reversed(numbers):
                anchor_comment(number, first, last)
        comment_parts(self.document.part, book)

    def close(self) -> None:
        self.equations.close()


def _break_before(paragraph: etree._Element) -> None:
    properties = paragraph.find(qn("w:pPr"))
    if properties is None:
        properties = element("w:pPr")
        paragraph.insert(0, properties)
    style = properties.find(qn("w:pStyle"))
    properties.insert(0 if style is None else 1, element("w:pageBreakBefore"))


def _drop_thumbnail(document: WordDocument) -> None:
    relations = document.part.package.rels
    for key, relation in list(relations.items()):
        if relation.reltype == RT.THUMBNAIL:
            relations.pop(key)


def _properties(document: WordDocument, content: StudioDocumentDocxParams, title: str) -> None:
    settings = content.settings
    properties = document.core_properties
    now = datetime.now(UTC).replace(microsecond=0, tzinfo=None)
    properties.title = (settings.title.strip() or title)[:255]
    properties.subject = settings.subtitle.strip()[:255]
    properties.author = settings.author.strip()[:255]
    properties.last_modified_by = settings.author.strip()[:255]
    properties.comments = ""
    properties.keywords = ""
    properties.category = ""
    properties.revision = 1
    properties.created = now
    properties.modified = now
    properties.language = content.language if content.language else ""


def document_docx(content: StudioDocumentDocxParams, check_cancelled: Callable[[], None]) -> bytes:
    try:
        root = lxml_html.fragment_fromstring(content.html or "<p></p>", create_parent="div")
    except (etree.ParserError, ValueError):
        root = lxml_html.fragment_fromstring("<p></p>", create_parent="div")
    document = docx.Document()
    _drop_thumbnail(document)
    writer = DocxWriter(document, content, check_cancelled)
    try:
        headings = outline(root, 3)
        title = content.title.strip() or (headings[0][1] if headings else "")
        setup_styles(document, content.settings, writer.families, content.language)
        writer.write(root, title)
        writer.finish()
        writer.comments(content.comments)
        finish_settings(
            document, content.settings.page_numbers == "outside", writer.notes_part is not None
        )
        _properties(document, content, title)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()
    finally:
        writer.close()
