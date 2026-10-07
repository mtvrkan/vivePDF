import contextlib
import os
from pathlib import Path
from typing import Literal

import pymupdf
from pydantic import Field

from vivepdf.ops._document import open_document, unwrap_document
from vivepdf.ops._inplace import save_in_place
from vivepdf.ops._ocr_area import read_area
from vivepdf.ops._ocr_parallel import Sheet, recognised_in_order, rendered_sheet, worker_count
from vivepdf.ops._orientation import best_rotation, capped_dpi
from vivepdf.ops._output import prepare_data_output, prepare_output, save_document
from vivepdf.ops._ranges import parse_page_ranges
from vivepdf.ops._toc import normalized_toc
from vivepdf.ops.tessdata import accurate_tessdata_dir, installed_languages, writable_tessdata_dir
from vivepdf.rpc.errors import ErrorCode, OpError
from vivepdf.rpc.progress import Progress
from vivepdf.rpc.protocol import RpcModel
from vivepdf.rpc.registry import op

MIN_TEXT_CHARS = 20
AREA_MIN_SIDE = 4.0
HIDDEN_RENDER_MODE = 3
HIDDEN_TEXT_HEIGHT_RATIO = 1.1
HIDDEN_TEXT_BASELINE_RATIO = 0.18
HIDDEN_TEXT_FONT = Path(__file__).resolve().parent.parent / "assets" / "fonts" / "DejaVuSans.ttf"
HIDDEN_FONT_NAME = "VpOcrText"
HIDDEN_CJK_FONT_NAME = "VpOcrCjk"
SCRIPT_FONTS: dict[str, dict[str, object]] = {
    HIDDEN_CJK_FONT_NAME: {"fontname": "cjk"},
    "VpOcrDeva": {"script": 9},
    "VpOcrBeng": {"script": 10},
    "VpOcrThai": {"script": 19},
}


class HiddenFonts:
    def __init__(self) -> None:
        self._latin = pymupdf.Font(fontfile=str(HIDDEN_TEXT_FONT))
        self._scripts: dict[str, pymupdf.Font] = {}

    def _script_font(self, name: str) -> pymupdf.Font:
        if name not in self._scripts:
            self._scripts[name] = pymupdf.Font(**SCRIPT_FONTS[name])
        return self._scripts[name]

    def pick(self, text: str) -> tuple[str, pymupdf.Font]:
        if all(self._latin.has_glyph(ord(char)) for char in text):
            return HIDDEN_FONT_NAME, self._latin
        best_name, best_font = HIDDEN_FONT_NAME, self._latin
        best_hits = sum(1 for char in text if self._latin.has_glyph(ord(char)))
        for name in SCRIPT_FONTS:
            font = self._script_font(name)
            hits = sum(1 for char in text if font.has_glyph(ord(char)))
            if hits > best_hits:
                best_name, best_font, best_hits = name, font, hits
        return best_name, best_font

    def install(self, page: pymupdf.Page, name: str) -> None:
        if name in SCRIPT_FONTS:
            page.insert_font(fontname=name, fontbuffer=self._script_font(name).buffer)
        else:
            page.insert_font(fontname=name, fontfile=str(HIDDEN_TEXT_FONT))


def available_languages() -> list[str]:
    return installed_languages()


def _language_string(languages: list[str]) -> str:
    available = set(available_languages())
    missing = [language for language in languages if language not in available]
    if missing or not languages:
        raise OpError(
            ErrorCode.TESSDATA_MISSING,
            f"missing OCR language data: {', '.join(missing) or 'none selected'}",
            {"missing": missing, "available": sorted(available)},
        )
    return "+".join(languages)


class OcrParams(RpcModel):
    path: str
    password: str | None = None
    output: str
    overwrite: bool = False
    languages: list[str] = Field(default_factory=lambda: ["tur", "eng"])
    dpi: int = Field(default=200, ge=100, le=400)
    mode: Literal["skip_text", "force", "redo"] = "skip_text"
    orientation: bool = False
    pages: str | None = None
    clean: bool = False
    text_output: str | None = None


class OcrResult(RpcModel):
    output: str
    page_count: int
    bytes: int
    ocr_pages: int
    skipped_pages: int
    rotated_pages: int
    words: int = 0
    redone_pages: int = 0
    text_output: str | None = None


def _has_text(page: pymupdf.Page) -> bool:
    return len(page.get_text().strip()) >= MIN_TEXT_CHARS


def _text_visibility(page: pymupdf.Page) -> tuple[int, int]:
    hidden = 0
    visible = 0
    for span in page.get_texttrace():
        count = sum(1 for char in span["chars"] if chr(char[0]).strip())
        if span["type"] == HIDDEN_RENDER_MODE or not span.get("opacity", 1):
            hidden += count
        else:
            visible += count
    return hidden, visible


def _strip_text(page: pymupdf.Page) -> None:
    links = page.get_links()
    page.add_redact_annot(page.rect, fill=False)
    page.apply_redactions(
        images=pymupdf.PDF_REDACT_IMAGE_NONE,
        graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
        text=pymupdf.PDF_REDACT_TEXT_REMOVE,
    )
    if len(page.get_links()) == len(links):
        return
    for link in page.get_links():
        page.delete_link(link)
    for link in links:
        link.pop("xref", None)
        link.pop("id", None)
        page.insert_link(link)


def _plan(page: pymupdf.Page, mode: str) -> str:
    if mode == "redo":
        hidden, visible = _text_visibility(page)
        if visible >= MIN_TEXT_CHARS:
            return "skip"
        return "redo" if hidden else "overlay"
    if not _has_text(page):
        return "overlay"
    return "replace" if mode == "force" else "skip"


def _page_blocks(page: pymupdf.Page) -> list[str]:
    return [
        block[4].strip()
        for block in page.get_text("blocks", sort=True)
        if block[6] == 0 and block[4].strip()
    ]


def _write_text_export(document: pymupdf.Document, target: Path) -> None:
    staging = target.with_name(f".{target.name}.{os.getpid()}.part")
    try:
        if target.suffix.lower() == ".docx":
            from docx import Document as WordDocument
            from docx.enum.text import WD_BREAK

            word = WordDocument()
            for number, page in enumerate(document):
                if number:
                    word.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                for block in _page_blocks(page):
                    word.add_paragraph(block)
            word.save(str(staging))
        else:
            pages = ["\n\n".join(_page_blocks(page)) for page in document]
            staging.write_text("\n\f".join(pages) + "\n", encoding="utf-8")
        os.replace(staging, target)
    finally:
        staging.unlink(missing_ok=True)


def _text_target(params: OcrParams) -> Path | None:
    if not params.text_output:
        return None
    extension = "docx" if Path(params.text_output).suffix.lower() == ".docx" else "txt"
    return prepare_data_output(params.text_output, extension, params.overwrite)


def _replace_with_sheet(document: pymupdf.Document, index: int, data: bytes) -> int:
    with pymupdf.open("pdf", data) as sheet:
        words = len(sheet[0].get_text("words"))
        document.insert_pdf(sheet, start_at=index)
    document.delete_page(index + 1)
    return words


def _overlay_sheet(page: pymupdf.Page, data: bytes, fonts: HiddenFonts) -> int:
    with pymupdf.open("pdf", data) as sheet:
        return _write_hidden_words(page, sheet[0], fonts)


@op("ocr.run", OcrParams)
def run_ocr(params: OcrParams, progress: Progress) -> OcrResult:
    target = prepare_output(params.output, [params.path], params.overwrite)
    text_target = _text_target(params)
    language = _language_string(params.languages)
    tessdata = str(writable_tessdata_dir())
    fonts = HiddenFonts()
    skipped = 0
    rotated = 0
    redone = 0
    words = 0
    with open_document(params.path, params.password) as document:
        total = document.page_count
        selected = sorted(set(parse_page_ranges(params.pages, total)))
        source_toc = document.get_toc(simple=True)
        planned: dict[int, str] = {}
        for index in selected:
            progress.check_cancelled()
            page = document[index]
            action = _plan(page, params.mode)
            if action == "skip":
                skipped += 1
                continue
            if action == "redo":
                _strip_text(page)
                redone += 1
                action = "overlay"
            if params.orientation:
                turn = best_rotation(page, language)
                if turn:
                    page.set_rotation((page.rotation + turn) % 360)
                    rotated += 1
            planned[index] = action
        indices = list(planned)

        def render(index: int) -> Sheet:
            page = document[index]
            pixmap = page.get_pixmap(dpi=capped_dpi(page.rect, params.dpi), alpha=False)
            return rendered_sheet(pixmap, params.clean and planned[index] == "overlay")

        results = recognised_in_order(
            indices,
            render,
            language,
            tessdata,
            progress.check_cancelled,
            worker_count(len(indices)),
        )
        with contextlib.closing(results):
            for done, (index, data) in enumerate(results, start=1):
                if planned[index] == "replace":
                    words += _replace_with_sheet(document, index, data)
                else:
                    words += _overlay_sheet(document[index], data, fonts)
                progress.report(
                    done / len(indices) * 0.93,
                    "progress.ocr",
                    {"current": done, "total": len(indices)},
                )
        if "replace" in planned.values() and source_toc:
            document.set_toc(normalized_toc(source_toc))
        progress.report(0.95, "progress.saving")
        if planned:
            with contextlib.suppress(Exception):
                document.subset_fonts()
        if text_target is not None:
            progress.check_cancelled()
            _write_text_export(document, text_target)
        saved = save_document(document, target)
    return OcrResult(
        output=saved.output,
        page_count=saved.page_count,
        bytes=saved.bytes,
        ocr_pages=len(planned),
        skipped_pages=skipped,
        rotated_pages=rotated,
        words=words,
        redone_pages=redone,
        text_output=str(text_target) if text_target is not None else None,
    )


class OcrTextParams(RpcModel):
    path: str
    password: str | None = None
    languages: list[str] = Field(default_factory=lambda: ["tur", "eng"])
    dpi: int = Field(default=200, ge=100, le=400)
    pages: str | None = None


class OcrTextResult(RpcModel):
    text: str
    page_count: int


@op("ocr.text", OcrTextParams)
def ocr_text(params: OcrTextParams, progress: Progress) -> OcrTextResult:
    language = _language_string(params.languages)
    chunks: list[str] = []
    with open_document(params.path, params.password) as document:
        indices = parse_page_ranges(params.pages, document.page_count)
        for position, index in enumerate(indices):
            progress.check_cancelled()
            page = document[index]
            textpage = page.get_textpage_ocr(
                flags=0,
                language=language,
                dpi=capped_dpi(page.rect, params.dpi),
                full=True,
                tessdata=str(writable_tessdata_dir()),
            )
            chunks.append(page.get_text(textpage=textpage))
            progress.report(
                (position + 1) / len(indices),
                "progress.ocr",
                {"current": position + 1, "total": len(indices)},
            )
    return OcrTextResult(text="\f".join(chunks), page_count=len(chunks))


class OcrAreaParams(RpcModel):
    path: str
    password: str | None = None
    page: int = Field(ge=0)
    rect: list[float] = Field(min_length=4, max_length=4)
    languages: list[str] = Field(default_factory=lambda: ["tur", "eng"])
    dpi: int = Field(default=300, ge=100, le=400)


class OcrAreaLine(RpcModel):
    text: str
    x0: float
    y0: float
    x1: float
    y1: float


class OcrAreaResult(RpcModel):
    text: str
    lines: list[OcrAreaLine]
    recognized: bool


def _visible_rect(page: pymupdf.Page, rect: pymupdf.Rect) -> pymupdf.Rect:
    shown = rect * page.rotation_matrix if page.rotation else pymupdf.Rect(rect)
    shown.normalize()
    return shown


def _page_bounds(page: pymupdf.Page) -> pymupdf.Rect:
    bounds = pymupdf.Rect(page.rect)
    if page.rotation:
        bounds = bounds * page.derotation_matrix
    bounds.normalize()
    return bounds


def _area_clip(page: pymupdf.Page, rect: list[float]) -> pymupdf.Rect:
    box = pymupdf.Rect(rect[0], rect[1], rect[2], rect[3])
    box.normalize()
    if box.width < AREA_MIN_SIDE or box.height < AREA_MIN_SIDE:
        raise OpError(ErrorCode.INVALID_PARAMS, "area is too small", {"reason": "rect"})
    raw = box * page.derotation_matrix if page.rotation else pymupdf.Rect(box)
    raw.normalize()
    clip = raw & _page_bounds(page)
    if clip.is_empty or clip.width < AREA_MIN_SIDE or clip.height < AREA_MIN_SIDE:
        raise OpError(ErrorCode.INVALID_PARAMS, "area is outside the page", {"reason": "rect"})
    return clip


def _lines_from_words(
    page: pymupdf.Page, words: list, scale: tuple[float, float], origin: tuple[float, float]
) -> list[OcrAreaLine]:
    grouped: dict[tuple[int, int], list] = {}
    for word in words:
        grouped.setdefault((word[5], word[6]), []).append(word)
    lines: list[OcrAreaLine] = []
    for key in sorted(grouped):
        row = sorted(grouped[key], key=lambda item: item[0])
        text = " ".join(item[4] for item in row).strip()
        if not text:
            continue
        box = pymupdf.Rect(row[0][:4])
        for item in row[1:]:
            box |= pymupdf.Rect(item[:4])
        placed = pymupdf.Rect(
            box.x0 * scale[0] + origin[0],
            box.y0 * scale[1] + origin[1],
            box.x1 * scale[0] + origin[0],
            box.y1 * scale[1] + origin[1],
        )
        shown = _visible_rect(page, placed)
        lines.append(
            OcrAreaLine(
                text=text,
                x0=round(shown.x0, 2),
                y0=round(shown.y0, 2),
                x1=round(shown.x1, 2),
                y1=round(shown.y1, 2),
            )
        )
    return lines


@op("ocr.area", OcrAreaParams)
def ocr_area(params: OcrAreaParams, progress: Progress) -> OcrAreaResult:
    with open_document(params.path, params.password, mutable=False) as cached:
        document = unwrap_document(cached)
        if params.page >= document.page_count:
            raise OpError(ErrorCode.INVALID_PARAMS, "page out of range", {"reason": "page"})
        page = document[params.page]
        clip = _area_clip(page, params.rect)
        existing = page.get_text("words", clip=clip)
        if existing:
            lines = _lines_from_words(page, existing, (1.0, 1.0), (0.0, 0.0))
            return OcrAreaResult(
                text="\n".join(line.text for line in lines), lines=lines, recognized=False
            )
        language = _language_string(params.languages)
        progress.check_cancelled()
        progress.report(0.2, "progress.ocr", {"current": 1, "total": 1})
        area = _visible_rect(page, clip)
        tessdata = accurate_tessdata_dir(params.languages) or writable_tessdata_dir()
        rows = read_area(
            page,
            area,
            params.dpi,
            language,
            str(tessdata),
            progress.check_cancelled,
        )
        lines = [
            OcrAreaLine(
                text=row.text,
                x0=round(row.rect.x0, 2),
                y0=round(row.rect.y0, 2),
                x1=round(row.rect.x1, 2),
                y1=round(row.rect.y1, 2),
            )
            for row in rows
        ]
        progress.report(1.0, "progress.ocr", {"current": 1, "total": 1})
        return OcrAreaResult(
            text="\n".join(line.text for line in lines), lines=lines, recognized=True
        )


class OcrSearchableParams(RpcModel):
    path: str
    password: str | None = None
    pages: str | None = None
    languages: list[str] = Field(default_factory=lambda: ["tur", "eng"])
    dpi: int = Field(default=300, ge=100, le=400)
    force: bool = False


class OcrSearchableResult(RpcModel):
    pages: int
    skipped: int
    words: int
    bytes: int


def _write_hidden_words(page: pymupdf.Page, recognised: pymupdf.Page, fonts: HiddenFonts) -> int:
    scale_x = page.rect.width / recognised.rect.width if recognised.rect.width else 1.0
    scale_y = page.rect.height / recognised.rect.height if recognised.rect.height else 1.0
    words = [
        (
            pymupdf.Rect(
                word[0] * scale_x, word[1] * scale_y, word[2] * scale_x, word[3] * scale_y
            ),
            word[4],
        )
        for word in recognised.get_text("words")
    ]
    return write_hidden_words(page, words, fonts)


def write_hidden_words(
    page: pymupdf.Page, words: list[tuple[pymupdf.Rect, str]], fonts: HiddenFonts
) -> int:
    derotate = page.derotation_matrix
    shape = page.new_shape()
    installed: set[str] = set()
    written = 0
    for box, raw in words:
        text = raw.strip()
        if not text:
            continue
        name, font = fonts.pick(text)
        unit = font.text_length(text, fontsize=1)
        if unit <= 0 or box.width <= 0 or box.height <= 0:
            continue
        size = min(box.width / unit, box.height * HIDDEN_TEXT_HEIGHT_RATIO)
        if size <= 0:
            continue
        if name not in installed:
            fonts.install(page, name)
            installed.add(name)
        baseline = pymupdf.Point(box.x0, box.y1 - box.height * HIDDEN_TEXT_BASELINE_RATIO)
        shape.insert_text(
            baseline * derotate,
            text,
            fontname=name,
            fontsize=size,
            render_mode=HIDDEN_RENDER_MODE,
            rotate=page.rotation,
        )
        written += 1
    if written:
        shape.commit()
    return written


@op("ocr.searchable", OcrSearchableParams)
def ocr_searchable(params: OcrSearchableParams, progress: Progress) -> OcrSearchableResult:
    language = _language_string(params.languages)
    tessdata = str(writable_tessdata_dir())
    fonts = HiddenFonts()
    document = open_document(params.path, params.password)
    recognised_pages = 0
    skipped = 0
    words = 0
    try:
        indices = parse_page_ranges(params.pages, document.page_count)
        for position, index in enumerate(indices):
            progress.check_cancelled()
            page = document[index]
            if not params.force and _has_text(page):
                skipped += 1
                continue
            pixmap = page.get_pixmap(dpi=capped_dpi(page.rect, params.dpi), alpha=False)
            try:
                data = pixmap.pdfocr_tobytes(compress=True, language=language, tessdata=tessdata)
            except Exception as error:  # noqa: BLE001
                raise OpError(
                    ErrorCode.INTERNAL, f"OCR failed on page {index + 1}: {error}"
                ) from error
            with pymupdf.open("pdf", data) as sheet:
                written = _write_hidden_words(page, sheet[0], fonts)
            words += written
            if written:
                recognised_pages += 1
            progress.report(
                (position + 1) / len(indices),
                "progress.ocr",
                {"current": position + 1, "total": len(indices)},
            )
        progress.report(0.98, "progress.saving")
        if words:
            with contextlib.suppress(Exception):
                document.subset_fonts()
        saved = save_in_place(document, params.path)
        return OcrSearchableResult(
            pages=recognised_pages, skipped=skipped, words=words, bytes=saved.bytes
        )
    finally:
        if not document.is_closed:
            document.close()
