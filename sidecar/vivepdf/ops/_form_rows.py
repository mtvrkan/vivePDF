import codecs
import csv
import io
import re
import unicodedata
from pathlib import Path
from typing import Any, Literal

from vivepdf.ops._sheet_cells import cell_text
from vivepdf.rpc.errors import ErrorCode, OpError

SPREADSHEET_EXTENSIONS = {".xlsx", ".xlsm"}
MAX_PREVIEW_ROWS = 50
FALLBACK_ENCODINGS = ("cp1254", "cp1252")
PATTERN_BUILTINS = {"name", "file", "n", "total", "date", "time", "year"}
Delimiter = Literal["auto", ",", ";", "\t", "|"]
PLACEHOLDER = re.compile(r"(?<!\{)\{([^{}]+)\}")


def decode_text(data: bytes) -> str:
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    for encoding in FALLBACK_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def unknown_placeholders(pattern: str, columns: list[str]) -> list[str]:
    known = PATTERN_BUILTINS | set(columns)
    unknown: list[str] = []
    for field in PLACEHOLDER.findall(pattern):
        key = re.split(r"[:!]", field, maxsplit=1)[0]
        if key not in known and key not in unknown:
            unknown.append(key)
    return unknown


def normalize_key(value: str) -> str:
    return re.sub(r"[\W_]+", "", unicodedata.normalize("NFKC", value).casefold())


def load_rows(
    path: str, sheet: str | None, delimiter: Delimiter = "auto"
) -> tuple[list[str], list[dict[str, str]], list[str]]:
    source = Path(path)
    if not source.is_file():
        raise OpError(ErrorCode.FILE_NOT_FOUND, f"file not found: {source.name}", {"path": path})
    if source.suffix.lower() in SPREADSHEET_EXTENSIONS:
        from openpyxl import load_workbook

        workbook = load_workbook(source, read_only=True, data_only=True)
        sheets = list(workbook.sheetnames)
        worksheet = workbook[sheet] if sheet and sheet in sheets else workbook[sheets[0]]
        iterator = worksheet.iter_rows()
        header = next(iterator, None) or ()
        columns = [_sheet_text(cell) for cell in header]
        rows: list[dict[str, str]] = []
        for raw in iterator:
            texts = [_sheet_text(cell) for cell in raw]
            if not any(texts):
                continue
            rows.append(
                {column: text for column, text in zip(columns, texts, strict=False) if column}
            )
        workbook.close()
        return [column for column in columns if column], rows, sheets
    text = decode_text(source.read_bytes())
    reader = csv.DictReader(io.StringIO(text, newline=""), dialect=_dialect(text, delimiter))
    columns = [column.strip() for column in (reader.fieldnames or []) if column and column.strip()]
    try:
        rows = [
            {
                key.strip(): (value or "").strip()
                for key, value in row.items()
                if key and key.strip()
            }
            for row in reader
            if any(isinstance(value, str) and value.strip() for value in row.values())
        ]
    except csv.Error as error:
        raise OpError(ErrorCode.INVALID_PARAMS, f"not a readable data file: {error}") from error
    return columns, rows, []


def _sheet_text(cell: Any) -> str:
    return cell_text(getattr(cell, "value", None), getattr(cell, "number_format", None)).strip()


def _dialect(text: str, delimiter: Delimiter) -> type[csv.Dialect] | csv.Dialect:
    if delimiter != "auto":

        class Chosen(csv.excel):
            pass

        Chosen.delimiter = delimiter
        return Chosen
    try:
        return csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
    except csv.Error:
        return csv.excel
